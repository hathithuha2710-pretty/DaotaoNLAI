import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"c:\Users\USER\Desktop\AI-vibe-code\KhungNLĐTAI\03_San_Pham_Chinh_Thuc"

def set_cell_margins(cell, top=100, bottom=100, left=100, right=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="D0D7DE", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def load_grade_data():
    grades_data = {}
    for g in range(1, 6):
        f_full = os.path.join(BASE_DIR, f"SP1_Khung_YCCD_AI_Lop{g}_Day_Du.md")
        f_comp = os.path.join(BASE_DIR, f"SP1_Khung_YCCD_AI_Lop{g}_Rut_Gon.md")
        
        with open(f_full, encoding="utf-8") as f:
            full_lines = [l.strip() for l in f if l.strip().startswith("| **")]
            
        with open(f_comp, encoding="utf-8") as f:
            comp_lines = [l.strip() for l in f if l.strip().startswith("| **")]
            
        full_parsed = []
        for line in full_lines:
            parts = [p.strip() for p in line.split("|")[1:-1]]
            full_parsed.append({
                "stt": parts[0].replace("**", ""),
                "code": parts[1].replace("**", ""),
                "strand_topic": parts[2].replace("**", ""),
                "vn_yccd": parts[3],
                "intl_yccd": parts[4]
            })

        comp_parsed = []
        for line in comp_lines:
            parts = [p.strip() for p in line.split("|")[1:-1]]
            comp_parsed.append({
                "stt": parts[0].replace("**", ""),
                "strand": parts[1].replace("**", ""),
                "subtopic": parts[2],
                "compact": parts[3]
            })

        grades_data[g] = {
            "full": full_parsed,
            "comp": comp_parsed
        }
        print(f"Loaded Grade {g}: Full={len(full_parsed)}, Comp={len(comp_parsed)}")
    return grades_data

def generate_master_full_md(grades_data, output_path):
    roman = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V"}
    md_content = [
        "# TỔNG HỢP KHUNG CHƯƠNG TRÌNH GIÁO DỤC TRÍ TUỆ NHÂN TẠO CẤP TIỂU HỌC (LỚP 1 ĐẾN LỚP 5)",
        "## BẢN ĐẦY ĐỦ: CHUẨN HOÁ THEO QUYẾT ĐỊNH 2422/QĐ-BGDĐT VÀ THAM CHIẾU QUỐC TẾ",
        "",
        "> **CĂN CỨ PHÁP LÝ CAO NHẤT:** Quyết định số 2422/QĐ-BGDĐT ngày 18/8/2026 của Bộ Giáo dục và Đào tạo ban hành *'Khung nội dung giáo dục Trí tuệ nhân tạo cho học sinh phổ thông'*.  ",
        "> **CĂN CỨ THAM CHIẾU QUỐC TẾ:** UNESCO AI Competency Framework for Students, OECD Learning Compass 2030, Singapore MOE NDLP, AI4K12 (5 Big Ideas), Google DeepMind (Experience AI), MIT RAISE (Day of AI).  ",
        "> **QUY MÔ TOÀN CẤP:** Gồm **123 Yêu cầu cần đạt (YCCĐ)** chính thức (Lớp 1: 21, Lớp 2: 29, Lớp 3: 32, Lớp 4: 17, Lớp 5: 24) thuộc 4 mạch năng lực (A. Tư duy lấy con người làm trung tâm; B. Đạo đức AI; C. Các kĩ thuật và ứng dụng AI; D. Thiết kế hệ thống AI).",
        "",
        "---",
        "",
        "## MỤC LỤC CÁC PHẦN",
        "- [Phần I: Khung YCCĐ AI Lớp 1 (21 YCCĐ)](#phan-i-khung-yccd-ai-lop-1-21-yccd)",
        "- [Phần II: Khung YCCĐ AI Lớp 2 (29 YCCĐ)](#phan-ii-khung-yccd-ai-lop-2-29-yccd)",
        "- [Phần III: Khung YCCĐ AI Lớp 3 (32 YCCĐ)](#phan-iii-khung-yccd-ai-lop-3-32-yccd)",
        "- [Phần IV: Khung YCCĐ AI Lớp 4 (17 YCCĐ)](#phan-iv-khung-yccd-ai-lop-4-17-yccd)",
        "- [Phần V: Khung YCCĐ AI Lớp 5 (24 YCCĐ)](#phan-v-khung-yccd-ai-lop-5-24-yccd)",
        "",
        "---"
    ]

    for g in range(1, 6):
        items = grades_data[g]["full"]
        r_num = roman[g]
        md_content.append("")
        md_content.append(f"## PHẦN {r_num}: KHUNG NỘI DUNG GIÁO DỤC AI – LỚP {g} ({len(items)} YCCĐ)")
        md_content.append("")
        md_content.append("| STT | Mã YCCĐ | Mạch nội dung & Chủ đề con (QĐ 2422) | Yêu cầu cần đạt Chuẩn Quốc gia (Bộ GD&ĐT - QĐ 2422) | Chuẩn đầu ra tham chiếu Đa nguồn Quốc tế (UNESCO, OECD, Singapore MOE, AI4K12, Google) |")
        md_content.append("| :---: | :---: | :--- | :--- | :--- |")
        for it in items:
            row = f"| **{it['stt']}** | **{it['code']}** | {it['strand_topic']} | {it['vn_yccd']} | {it['intl_yccd']} |"
            md_content.append(row)
        md_content.append("")
        md_content.append("---")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md_content))
    print(f"Exported Master Full MD: {output_path}")

def generate_master_comp_md(grades_data, output_path):
    roman = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V"}
    md_content = [
        "# TỔNG HỢP KHUNG NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO CẤP TIỂU HỌC (LỚP 1 ĐẾN LỚP 5)",
        "## BẢN RÚT GỌN CHUẨN HOÁ: CHUẨN ĐẦU RA NĂNG LỰC (EXIT COMPETENCIES)",
        "",
        "> **Căn cứ pháp lý:** Quyết định số 2422/QĐ-BGDĐT của Bộ Giáo dục và Đào tạo.  ",
        "> **Đặc điểm:** Tinh gọn các chú giải nguồn quốc tế, tích hợp thành các Chuẩn đầu ra năng lực súc tích, phục vụ trực tiếp việc lập Kế hoạch bài dạy (Giáo án) và Phân phối chương trình giảng dạy thực tế.  ",
        "> **Quy mô toàn cấp:** 123 YCCĐ (Lớp 1: 21, Lớp 2: 29, Lớp 3: 32, Lớp 4: 17, Lớp 5: 24).",
        "",
        "---",
        "",
        "## MỤC LỤC CÁC PHẦN",
        "- [Phần I: Lớp 1 (21 YCCĐ)](#phan-i-lop-1-21-yccd)",
        "- [Phần II: Lớp 2 (29 YCCĐ)](#phan-ii-lop-2-29-yccd)",
        "- [Phần III: Lớp 3 (32 YCCĐ)](#phan-iii-lop-3-32-yccd)",
        "- [Phần IV: Lớp 4 (17 YCCĐ)](#phan-iv-lop-4-17-yccd)",
        "- [Phần V: Lớp 5 (24 YCCĐ)](#phan-v-lop-5-24-yccd)",
        "",
        "---"
    ]

    for g in range(1, 6):
        items = grades_data[g]["comp"]
        r_num = roman[g]
        md_content.append("")
        md_content.append(f"## PHẦN {r_num}: LỚP {g} ({len(items)} YCCĐ)")
        md_content.append("")
        md_content.append("| STT | Mạch năng lực | Chủ đề con | YCCĐ (Chuẩn đầu ra năng lực) |")
        md_content.append("| :---: | :--- | :--- | :--- |")
        for it in items:
            row = f"| **{it['stt']}** | **{it['strand']}** | {it['subtopic']} | {it['compact']} |"
            md_content.append(row)
        md_content.append("")
        md_content.append("---")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md_content))
    print(f"Exported Master Compact MD: {output_path}")

def generate_master_full_docx(grades_data, output_path):
    doc = docx.Document()
    
    # Page setup: Landscape A4
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Document Title Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — CHƯƠNG TRÌNH GIÁO DỤC TRÍ TUỆ NHÂN TẠO PHỔ THÔNG\n")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(10.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(90, 90, 90)

    r_main = p_title.add_run("KHUNG CHƯƠNG TRÌNH GIÁO DỤC TRÍ TUỆ NHÂN TẠO CẤP TIỂU HỌC\n(TỔNG HỢP TOÀN DIỆN TỪ LỚP 1 ĐẾN LỚP 5 – BẢN ĐẦY ĐỦ)")
    r_main.font.name = "Times New Roman"
    r_main.font.size = Pt(15)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(0, 32, 96)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Căn cứ pháp lý: Quyết định số 2422/QĐ-BGDĐT | Tham chiếu: UNESCO, OECD, Singapore MOE, AI4K12, Google DeepMind\nTổng số: 123 Yêu cầu cần đạt (Lớp 1: 21 YCCĐ | Lớp 2: 29 YCCĐ | Lớp 3: 32 YCCĐ | Lớp 4: 17 YCCĐ | Lớp 5: 24 YCCĐ)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(70, 70, 70)
    p_sub.paragraph_format.space_after = Pt(12)

    headers = [
        ("STT", Inches(0.5)),
        ("Mã YCCĐ", Inches(0.9)),
        ("Mạch nội dung & Chủ đề con", Inches(2.3)),
        ("Yêu cầu cần đạt Chuẩn Quốc gia\n(Bộ GD&ĐT - QĐ 2422)", Inches(3.4)),
        ("Chuẩn đầu ra tham chiếu Đa nguồn Quốc tế\n(UNESCO, OECD, Singapore MOE, AI4K12, Google)", Inches(3.5))
    ]

    roman = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V"}

    for g in range(1, 6):
        items = grades_data[g]["full"]
        r_num = roman[g]

        if g > 1:
            doc.add_page_break()

        # Grade Heading
        p_gh = doc.add_paragraph()
        r_gh = p_gh.add_run(f"PHẦN {r_num}: KHUNG NỘI DUNG GIÁO DỤC AI – LỚP {g} ({len(items)} YCCĐ)")
        r_gh.font.name = "Times New Roman"
        r_gh.font.size = Pt(12.5)
        r_gh.font.bold = True
        r_gh.font.color.rgb = RGBColor(0, 51, 102)
        p_gh.paragraph_format.space_before = Pt(6)
        p_gh.paragraph_format.space_after = Pt(6)

        table = doc.add_table(rows=len(items) + 1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        # Header row
        hdr_cells = table.rows[0].cells
        for idx, (h_text, w) in enumerate(headers):
            hdr_cells[idx].width = w
            set_cell_margins(hdr_cells[idx], top=100, bottom=100, left=90, right=90)
            set_cell_shading(hdr_cells[idx], "1B365D")
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        header_tr = table.rows[0]._tr.get_or_add_trPr()
        header_tr.append(OxmlElement('w:tblHeader'))

        # Data rows
        for row_idx, it in enumerate(items, start=1):
            row_cells = table.rows[row_idx].cells
            bg_color = "F7F9FC" if row_idx % 2 == 0 else "FFFFFF"

            for idx, (_, w) in enumerate(headers):
                row_cells[idx].width = w
                set_cell_margins(row_cells[idx], top=70, bottom=70, left=80, right=80)
                set_cell_shading(row_cells[idx], bg_color)

            # Cell 0: STT
            p0 = row_cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(str(it["stt"]))
            r0.font.name = "Times New Roman"
            r0.font.size = Pt(9)
            r0.font.bold = True

            # Cell 1: Code
            p1 = row_cells[1].paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(it["code"])
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(9)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(0, 32, 96)

            # Cell 2: Strand & Topic
            p2 = row_cells[2].paragraphs[0]
            p2.paragraph_format.line_spacing = 1.15
            p2.paragraph_format.space_after = Pt(2)
            st_parts = it["strand_topic"].split('<br>')
            r2_a = p2.add_run(st_parts[0] + "\n")
            r2_a.font.name = "Times New Roman"
            r2_a.font.size = Pt(8.5)
            r2_a.font.bold = True
            if len(st_parts) > 1:
                r2_b = p2.add_run(st_parts[1].replace('*', ''))
                r2_b.font.name = "Times New Roman"
                r2_b.font.size = Pt(8)
                r2_b.font.italic = True
                r2_b.font.color.rgb = RGBColor(60, 60, 60)

            # Cell 3: VN YCCĐ
            p3 = row_cells[3].paragraphs[0]
            p3.paragraph_format.line_spacing = 1.15
            r3 = p3.add_run(it["vn_yccd"])
            r3.font.name = "Times New Roman"
            r3.font.size = Pt(8.5)

            # Cell 4: International Crosswalk
            p4 = row_cells[4].paragraphs[0]
            p4.paragraph_format.line_spacing = 1.15
            lines_intl = it["intl_yccd"].split("<br>")
            for l_idx, line_txt in enumerate(lines_intl):
                if l_idx > 0:
                    p4 = row_cells[4].add_paragraph()
                    p4.paragraph_format.line_spacing = 1.15
                    p4.paragraph_format.space_after = Pt(2)
                c_parts = line_txt.split(":", 1)
                if len(c_parts) == 2:
                    r_bold = p4.add_run(c_parts[0] + ":")
                    r_bold.font.name = "Times New Roman"
                    r_bold.font.size = Pt(8)
                    r_bold.font.bold = True
                    r_bold.font.color.rgb = RGBColor(0, 51, 102)
                    r_norm = p4.add_run(c_parts[1])
                    r_norm.font.name = "Times New Roman"
                    r_norm.font.size = Pt(8)
                else:
                    r_all = p4.add_run(line_txt)
                    r_all.font.name = "Times New Roman"
                    r_all.font.size = Pt(8)

    doc.save(output_path)
    print(f"Exported Master Full DOCX: {output_path}")

def generate_master_comp_docx(grades_data, output_path):
    doc = docx.Document()
    
    # Page setup: Portrait A4
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Document Title Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO\n")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(90, 90, 90)

    r_main = p_title.add_run("KHUNG NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO CẤP TIỂU HỌC\n(TỔNG HỢP LỚP 1 ĐẾN LỚP 5 – BẢN RÚT GỌN CHUẨN HOÁ)")
    r_main.font.name = "Times New Roman"
    r_main.font.size = Pt(14)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(0, 32, 96)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Chuẩn đầu ra năng lực (Exit Competencies) súc tích – Phục vụ lập KHBD và PPCT thực tế\nQuy mô: 123 YCCĐ (Lớp 1: 21 | Lớp 2: 29 | Lớp 3: 32 | Lớp 4: 17 | Lớp 5: 24)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(70, 70, 70)
    p_sub.paragraph_format.space_after = Pt(10)

    headers = [
        ("STT", Inches(0.5)),
        ("Mạch năng lực", Inches(1.8)),
        ("Chủ đề con", Inches(1.9)),
        ("YCCĐ (Chuẩn đầu ra năng lực)", Inches(2.8))
    ]

    roman = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V"}

    for g in range(1, 6):
        items = grades_data[g]["comp"]
        r_num = roman[g]

        if g > 1:
            doc.add_page_break()

        # Grade Heading
        p_gh = doc.add_paragraph()
        r_gh = p_gh.add_run(f"PHẦN {r_num}: LỚP {g} ({len(items)} YCCĐ)")
        r_gh.font.name = "Times New Roman"
        r_gh.font.size = Pt(12)
        r_gh.font.bold = True
        r_gh.font.color.rgb = RGBColor(0, 51, 102)
        p_gh.paragraph_format.space_before = Pt(6)
        p_gh.paragraph_format.space_after = Pt(6)

        table = doc.add_table(rows=len(items) + 1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        # Header row
        hdr_cells = table.rows[0].cells
        for idx, (h_text, w) in enumerate(headers):
            hdr_cells[idx].width = w
            set_cell_margins(hdr_cells[idx], top=100, bottom=100, left=90, right=90)
            set_cell_shading(hdr_cells[idx], "1B365D")
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        header_tr = table.rows[0]._tr.get_or_add_trPr()
        header_tr.append(OxmlElement('w:tblHeader'))

        # Data rows
        for row_idx, it in enumerate(items, start=1):
            row_cells = table.rows[row_idx].cells
            bg_color = "F7F9FC" if row_idx % 2 == 0 else "FFFFFF"

            for idx, (_, w) in enumerate(headers):
                row_cells[idx].width = w
                set_cell_margins(row_cells[idx], top=70, bottom=70, left=80, right=80)
                set_cell_shading(row_cells[idx], bg_color)

            # Cell 0: STT
            p0 = row_cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(str(it["stt"]))
            r0.font.name = "Times New Roman"
            r0.font.size = Pt(9)
            r0.font.bold = True

            # Cell 1: Strand
            p1 = row_cells[1].paragraphs[0]
            p1.paragraph_format.line_spacing = 1.15
            r1 = p1.add_run(it["strand"])
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(9)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(0, 32, 96)

            # Cell 2: Subtopic
            p2 = row_cells[2].paragraphs[0]
            p2.paragraph_format.line_spacing = 1.15
            r2 = p2.add_run(it["subtopic"])
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(8.5)

            # Cell 3: Compact YCCĐ
            p3 = row_cells[3].paragraphs[0]
            p3.paragraph_format.line_spacing = 1.15
            r3 = p3.add_run(it["compact"])
            r3.font.name = "Times New Roman"
            r3.font.size = Pt(8.5)

    doc.save(output_path)
    print(f"Exported Master Compact DOCX: {output_path}")

if __name__ == "__main__":
    out_dir = BASE_DIR
    data = load_grade_data()
    
    # 1. File Word đầy đủ (Lớp 1-5)
    f_master_full_docx = os.path.join(out_dir, "SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Day_Du.docx")
    generate_master_full_docx(data, f_master_full_docx)
    
    # 2. File Word rút gọn (Lớp 1-5)
    f_master_comp_docx = os.path.join(out_dir, "SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.docx")
    generate_master_comp_docx(data, f_master_comp_docx)

    # 3. File Markdown đầy đủ (Lớp 1-5)
    f_master_full_md = os.path.join(out_dir, "SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Day_Du.md")
    generate_master_full_md(data, f_master_full_md)

    # 4. File Markdown rút gọn (Lớp 1-5)
    f_master_comp_md = os.path.join(out_dir, "SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.md")
    generate_master_comp_md(data, f_master_comp_md)

    print("\nSUCCESS: TOÀN BỘ 4 TỆP MASTER TỔNG HỢP CẢ 5 LỚP ĐÃ ĐƯỢC TẠO THÀNH CÔNG VỚI CHẤT LƯỢNG CAO NHẤT!")
