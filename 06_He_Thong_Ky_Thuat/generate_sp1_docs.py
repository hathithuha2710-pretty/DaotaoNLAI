# -*- coding: utf-8 -*-
"""
Script tạo file Word đầy đủ và file Markdown tinh gọn cho Sản phẩm 1:
Khung chương trình chuẩn Cụm Tiểu học khởi đầu (Lớp 1 và Lớp 2)
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

data_items = [
    # LỚP 1
    {
        "stt": 1, "grade": 1, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – Con người có cảm xúc, AI thì không",
        "code": "1.A1.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được ví dụ về một số cảm xúc quen thuộc của con người (vui, buồn, giận dữ, sợ hãi, ngạc nhiên,…) qua các tình huống gần gũi trong cuộc sống và bước đầu nhận ra rằng con người có cảm xúc còn AI thì không.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 1 & UNESCO Human Agency]",
        "int_text": "Biết biểu đạt cảm xúc sinh học qua nét mặt, giọng nói và đối chiếu với phản ứng cơ học/điện tử của máy thông minh; khẳng định máy móc không có cảm xúc nội tại."
    },
    {
        "stt": 2, "grade": 1, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – Con người có cảm xúc, AI thì không",
        "code": "1.A1.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được rằng AI có thể mô phỏng hoặc nhận diện cảm xúc của con người, chứ không trải nghiệm cảm xúc như con người.",
        "tag_int": "[INT-ENRICH: OECD AI Literacy]",
        "int_text": "Phân biệt được sự khác nhau giữa \"cảm xúc thật\" của con người với việc máy móc chỉ \"nhận diện điểm ảnh nụ cười\" để hiển thị biểu tượng vui vẻ."
    },
    {
        "stt": 3, "grade": 1, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – AI thể hiện cảm xúc do con người lập trình",
        "code": "1.A1.3", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được rằng việc AI thể hiện cảm xúc (ví dụ: cười, buồn, ngạc nhiên, tức giận trên màn hình hoặc qua giọng nói) là do con người lập trình hoặc thiết kế trước.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 3 - Learning/Programming]",
        "int_text": "Nhận ra mọi biểu cảm của AI đều là chuỗi chỉ dẫn do con người cài đặt sẵn, không phải là phản xạ tự nhiên của máy."
    },
    {
        "stt": 4, "grade": 1, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – AI thể hiện cảm xúc do con người lập trình",
        "code": "1.A1.4", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được một số ví dụ minh họa (như robot cười khi người dùng kể chuyện vui, hoặc trợ lí ảo nói “Tôi rất vui được giúp bạn”), và giải thích được rằng những biểu hiện này chỉ là phản ứng được lập trình sẵn.",
        "tag_int": "[INT-ENRICH: Singapore MOE - Cognitive Independence]",
        "int_text": "Thể hiện sự tỉnh táo khi tương tác với trợ lí ảo, không đồng nhất lời chào hay nụ cười của AI với sự đồng cảm của một người bạn thật sự."
    },
    {
        "stt": 5, "grade": 1, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – Ý nghĩa cảm xúc mà AI thể hiện",
        "code": "1.A2.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được các cách mà AI thể hiện cảm xúc qua hình ảnh, giọng nói hoặc hành vi (ví dụ: biểu tượng khuôn mặt vui, giọng nói thân thiện, hoặc lời chào dễ thương).",
        "tag_int": "[INT-ENRICH: UNESCO User Interface Interaction]",
        "int_text": "Quan sát và liệt kê được các tín hiệu nghe – nhìn mà hệ thống AI sử dụng để tạo giao diện thân thiện với người dùng."
    },
    {
        "stt": 6, "grade": 1, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – Ý nghĩa cảm xúc mà AI thể hiện",
        "code": "1.A2.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Nêu được rằng những biểu hiện này giúp AI giao tiếp tự nhiên hơn với con người, khiến người dùng cảm thấy thoải mái, gần gũi và dễ hợp tác hơn.",
        "tag_int": "[INT-ENRICH: UNESCO Human-AI Interaction & UX]",
        "int_text": "Giải thích được vì sao loa thông minh hay robot sử dụng giọng nói ấm áp và lời chào lịch sự để tạo trải nghiệm cộng tác dễ chịu cho con người."
    },
    {
        "stt": 7, "grade": 1, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – Nhận diện AI trong cuộc sống",
        "code": "1.A2.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết và kể tên được một số sản phẩm hoặc thiết bị có sử dụng trí tuệ nhân tạo (ví dụ: loa thông minh, điện thoại có trợ lí ảo, robot hút bụi, xe tự lái, camera nhận diện khuôn mặt, ứng dụng dịch ngôn ngữ…).",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 5 - Societal Impact]",
        "int_text": "Chỉ ra sự hiện diện của các thiết bị thông minh trong môi trường xung quanh (nhà ở, trường học, đường phố) và phân biệt chúng với các đồ dùng cơ khí thông thường."
    },
    {
        "stt": 8, "grade": 1, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – Nhận diện AI trong cuộc sống",
        "code": "1.A2.3", "tag_vn": "[VN-CORE]",
        "qd_text": "Mô tả được công dụng chính của từng sản phẩm và cách AI giúp sản phẩm hoạt động thông minh hơn (ví dụ: loa thông minh hiểu lệnh bằng giọng nói; robot hút bụi tự tìm đường và tránh chướng ngại vật).",
        "tag_int": "[INT-ENRICH: OECD AI Applications]",
        "int_text": "Trình bày được tính năng tự thích ứng của thiết bị AI (như tự chuyển hướng khi gặp chướng ngại vật, tự kích hoạt khi nghe đúng từ khóa gọi tên)."
    },
    {
        "stt": 9, "grade": 1, "strand": "B. Đạo đức AI",
        "subtopic": "B1. Các khía cạnh đạo đức của AI – Việc làm tốt và việc làm xấu",
        "code": "1.B1.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được một số hành vi sử dụng AI có thể gây hại cho người khác.",
        "tag_int": "[INT-ENRICH: UNESCO AI Ethics & Safety]",
        "int_text": "Nhận diện các hành vi không an toàn như dùng thiết bị AI để trêu chọc, gây ồn ào làm phiền người khác hoặc chia sẻ hình ảnh riêng tư của bạn bè."
    },
    {
        "stt": 10, "grade": 1, "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Máy thông minh làm việc tốt",
        "code": "1.B3.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết được rằng không được phép sử dụng AI với mục đích làm hại người khác.",
        "tag_int": "[INT-ENRICH: UNESCO Ethics - Non-harm Principle]",
        "int_text": "Nêu được nguyên tắc đạo đức cốt lõi: Công nghệ chỉ được dùng để giúp đỡ, tuyệt đối không dùng AI làm công cụ bắt nạt, nói dối hay gây tổn hại."
    },
    {
        "stt": 11, "grade": 1, "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Máy thông minh làm việc tốt",
        "code": "1.B3.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được ví dụ về việc con người sử dụng AI đúng cách, vì mục đích tốt đẹp.",
        "tag_int": "[INT-ENRICH: AI4K12 AI for Social Good]",
        "int_text": "Kể được các tình huống AI giúp ích cụ thể như cảnh báo nguy hiểm, hỗ trợ người khuyết tật, giúp người già điều khiển thiết bị trong nhà khi gặp khó khăn."
    },
    {
        "stt": 12, "grade": 1, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Nhận biết AI và ứng dụng AI",
        "code": "1.C1.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết được AI trong một số ví dụ cụ thể và đơn giản về AI.",
        "tag_int": "[INT-ENRICH: AI4K12 What is AI]",
        "int_text": "Phân biệt được đồ vật thông minh có AI (có khả năng tự nhận biết và đưa ra phản hồi phù hợp) với thiết bị điện tử thông thường chỉ chạy theo công tắc bật/tắt."
    },
    {
        "stt": 13, "grade": 1, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Nhận biết AI và ứng dụng AI",
        "code": "1.C1.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận diện được một số công cụ AI quen thuộc và phổ biến trên điện thoại hoặc máy tính bảng (ví dụ: trợ lí ảo, nhận diện khuôn mặt, …).",
        "tag_int": "[INT-ENRICH: UNESCO Technology Familiarity]",
        "int_text": "Nhận biết được biểu tượng ứng dụng AI trên màn hình cảm ứng (nút micro trợ lí ảo, khung mở khóa nhận diện khuôn mặt)."
    },
    {
        "stt": 14, "grade": 1, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Chức năng và công cụ AI",
        "code": "1.C1.3", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết được các thiết bị thông minh có các bộ phận giống với bộ phận của con người (ví dụ: camera là “mắt” và micro là “tai” của các thiết bị đó).",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 1 - Perception]",
        "int_text": "Biết thử nghiệm đóng vai che \"mắt\" (camera) hoặc bịt \"tai\" (micro) của thiết bị để nhận ra khi thiếu cảm biến, AI không thể thu nhận thông tin từ môi trường xung quanh."
    },
    {
        "stt": 15, "grade": 1, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Chức năng và công cụ AI",
        "code": "1.C1.4", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết được AI có khả năng hiểu các mệnh lệnh đơn giản của con người, trò chuyện theo kịch bản định sẵn.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 4 - Natural Interaction]",
        "int_text": "Thử nghiệm đọc câu lệnh mẫu ngắn, rõ ràng và quan sát phản hồi theo kịch bản được định hình trước của máy."
    },
    {
        "stt": 16, "grade": 1, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Chức năng và công cụ AI",
        "code": "1.C1.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Nhận biết được AI có khả năng xử lí hình ảnh nhận được để nhận diện các đồ vật.",
        "tag_int": "[INT-ENRICH: MIT RAISE - Computer Vision Basics]",
        "int_text": "Quan sát khung nhận diện đối tượng trên màn hình camera để hiểu cách máy tính \"thấy\" và phân biệt đồ vật trong không gian thực tế."
    },
    {
        "stt": 17, "grade": 1, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Chức năng và công cụ AI",
        "code": "1.C1.MR2", "tag_vn": "[VN-MR]",
        "qd_text": "Nhận biết được AI có khả năng xử lí âm thanh để phân biệt các loại âm thanh khác nhau.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 1 - Audio Perception]",
        "int_text": "Trải nghiệm thử thách phát các âm thanh khác nhau (tiếng vỗ tay, tiếng chuông, tiếng nói) để kiểm tra khả năng phân biệt âm thanh của hệ thống thông minh."
    },
    {
        "stt": 18, "grade": 1, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Máy thông minh học từ ví dụ",
        "code": "1.D1.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được ví dụ về một tình huống mà AI “học” từ hình ảnh hoặc thông tin do con người cung cấp (ví dụ: AI học nhận biết con mèo, quả táo…).",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 3 - Learning from Examples]",
        "int_text": "Trải nghiệm hoạt động Unplugged xếp các tấm thẻ tranh mẫu để mô phỏng cách con người nạp hình ảnh cho máy thông minh học hỏi."
    },
    {
        "stt": 19, "grade": 1, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Máy thông minh học từ ví dụ",
        "code": "1.D1.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Biết được rằng để AI trả lời đúng cần có nhiều ví dụ đúng và khác nhau để AI học.",
        "tag_int": "[INT-ENRICH: OECD AI Data Diversity]",
        "int_text": "Giải thích được vì sao chỉ cho xem 1 bức tranh thì máy chưa nhận diện chính xác; máy cần nhiều hình ảnh với màu sắc, kích thước và góc chụp đa dạng."
    },
    {
        "stt": 20, "grade": 1, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Nhiều loại máy thông minh",
        "code": "1.D2.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết được rằng có loại máy thông minh chỉ làm được một việc (như nhận biết hình ảnh), có loại có thể làm nhiều việc khác nhau (như lắng nghe - trả lời, thực hiện theo lệnh…).",
        "tag_int": "[INT-ENRICH: UNESCO AI Systems Architecture]",
        "int_text": "Nhận diện sự khác biệt giữa hệ thống chuyên biệt (chỉ quét khuôn mặt/mã vạch) và hệ thống đa năng (vừa nghe, vừa nói, vừa điều khiển đồ gia dụng)."
    },
    {
        "stt": 21, "grade": 1, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Nhiều loại máy thông minh",
        "code": "1.D2.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Xếp được một số máy thông minh vào hai nhóm: nhóm chỉ làm được một việc và nhóm làm được nhiều việc khác nhau.",
        "tag_int": "[INT-ENRICH: AI4K12 Computational Thinking]",
        "int_text": "Thực hành phân loại thẻ tranh các thiết bị AI vào hai nhóm chức năng: \"Thiết bị chuyên biệt\" và \"Thiết bị đa nhiệm\"."
    },

    # LỚP 2
    {
        "stt": 22, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – Khi nào nên và không nên dùng AI",
        "code": "2.A1.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Trình bày được một số tình huống trong cuộc sống mà AI có thể hỗ trợ con người hiệu quả, chẳng hạn như: sử dụng AI để dịch ngôn ngữ, tìm kiếm thông tin nhanh, phát hiện lỗi chính tả, hoặc điều khiển thiết bị thông minh trong nhà, …",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 5 & UNESCO Application]",
        "int_text": "Lựa chọn đúng công cụ AI hỗ trợ cho từng tình huống thực tế cụ thể nhằm tiết kiệm thời gian và tăng năng suất cho con người."
    },
    {
        "stt": 23, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – Khi nào nên và không nên dùng AI",
        "code": "2.A1.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Trình bày được một số tình huống trong cuộc sống mà không nên hoặc cần thận trọng khi sử dụng AI, ví dụ: khi AI có thể làm lộ thông tin cá nhân (như khuôn mặt, địa chỉ, dữ liệu giọng nói); khi AI thay thế hoàn toàn con người, dẫn đến sự thiếu trách nhiệm.",
        "tag_int": "[INT-ENRICH: Singapore MOE - Cognitive Independence & Ethics]",
        "int_text": "Xác định các ranh giới cảnh giác: không nhập thông tin bí mật gia đình vào ứng dụng AI; không để máy làm thay bài tập dẫn đến lười suy nghĩ."
    },
    {
        "stt": 24, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – AI làm việc, con người kiểm soát",
        "code": "2.A1.3", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được ví dụ cụ thể về tình huống cần con người giám sát AI (ví dụ: người lái xe vẫn cần quan sát khi xe tự lái hoạt động; bác sĩ cần kiểm tra lại kết quả do AI gợi ý).",
        "tag_int": "[INT-ENRICH: OECD Principles - Human Oversight]",
        "int_text": "Nêu bật nguyên tắc \"Con người trong vòng lặp\" (Human-in-the-loop): AI chỉ là trợ thủ, con người phải luôn theo dõi và giữ quyền quyết định cuối cùng."
    },
    {
        "stt": 25, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – AI làm việc, con người kiểm soát",
        "code": "2.A1.4", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được ví dụ cụ thể về thái độ đúng đắn khi sử dụng AI: biết tin tưởng vào công nghệ ở mức hợp lý, nhưng vẫn có trách nhiệm và sự kiểm soát của con người.",
        "tag_int": "[INT-ENRICH: UNESCO Mindset - Balanced Trust]",
        "int_text": "Hình thành thái độ tự chủ, tỉnh táo trước công nghệ: biết tận dụng tiện ích của AI nhưng luôn giữ tinh thần hoài nghi khoa học lành mạnh và ý thức trách nhiệm cá nhân."
    },
    {
        "stt": 26, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI trong gia đình",
        "code": "2.A2.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được một số thiết bị hoặc ứng dụng trong gia đình có sử dụng AI, chẳng hạn như loa thông minh, robot hút bụi, máy điều hoà tự động điều chỉnh nhiệt độ, camera nhận diện người trong nhà, hoặc tivi thông minh gợi ý chương trình yêu thích.",
        "tag_int": "[INT-ENRICH: UNESCO AI in Everyday Life]",
        "int_text": "Lập sơ đồ các \"điểm chạm AI\" trong ngôi nhà thông minh, mô tả cách các thiết bị kết nối để phục vụ sinh hoạt hàng ngày."
    },
    {
        "stt": 27, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI trong gia đình",
        "code": "2.A2.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Thực hành sử dụng một số thiết bị hoặc ứng dụng có sử dụng AI với sự giám sát, hỗ trợ của giáo viên.",
        "tag_int": "[INT-ENRICH: UNESCO Hands-on Experience]",
        "int_text": "Thực hiện an toàn các thao tác ra lệnh bằng giọng nói hoặc cử chỉ đơn giản dưới sự hướng dẫn trực tiếp của giáo viên trên thiết bị học tập."
    },
    {
        "stt": 28, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI trong gia đình",
        "code": "2.A2.MR2", "tag_vn": "[VN-MR]",
        "qd_text": "Nêu được rằng không được sử dụng thiết bị hoặc ứng dụng có sử dụng AI mà không có sự giám sát của người lớn.",
        "tag_int": "[INT-ENRICH: OECD Digital Safety for Children]",
        "int_text": "Nhận thức và thực thi quy tắc an toàn số: Luôn xin phép và có cha mẹ hoặc thầy cô đồng hành khi tương tác với các ứng dụng công nghệ thông minh."
    },
    {
        "stt": 29, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI hỗ trợ mọi người",
        "code": "2.A2.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết và kể tên được các thành viên trong gia đình mà AI có thể hỗ trợ.",
        "tag_int": "[INT-ENRICH: UNESCO Human-Centric Values]",
        "int_text": "Phân tích nhu cầu thực tế của từng thành viên (ông bà cần hỗ trợ sức khỏe, cha mẹ cần dọn dẹp, bản thân cần học tập) để thấy AI phục vụ con người."
    },
    {
        "stt": 30, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI hỗ trợ mọi người",
        "code": "2.A2.MR3", "tag_vn": "[VN-MR]",
        "qd_text": "Nêu được rằng mục tiêu của AI trong bối cảnh gia đình là giúp ích cho cuộc sống của các thành viên tiện nghi hơn, an toàn hơn và có nhiều thời gian cho nhau hơn.",
        "tag_int": "[INT-ENRICH: AI4K12 Societal Impact - Family Well-being]",
        "int_text": "Khẳng định ý nghĩa nhân văn của AI: Giảm bớt gánh nặng việc nhà để con người có nhiều thời gian trò chuyện, gắn kết tình cảm bên nhau."
    },
    {
        "stt": 31, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A3. Công dân trong kỉ nguyên AI – Con người dạy AI qua tương tác",
        "code": "2.A3.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết được rằng mỗi lần con người nói chuyện, đặt câu hỏi, tìm kiếm, hoặc đưa ra lựa chọn với các hệ thống AI (như trợ lí ảo, chatbot, ứng dụng gợi ý video hoặc nhạc) thì AI có thể ghi nhận dữ liệu để học hỏi cách con người suy nghĩ và phản hồi.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 3 - Feedback Loops]",
        "int_text": "Nhận diện khái niệm \"dấu chân dữ liệu\" sơ khai: Mỗi cú nhấp chuột hay câu hỏi đều trở thành dữ liệu để máy tính học thói quen của người dùng."
    },
    {
        "stt": 32, "grade": 2, "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A3. Công dân trong kỉ nguyên AI – Con người dạy AI qua tương tác",
        "code": "2.A3.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Nêu được rằng con người có vai trò quan trọng trong việc “dạy” AI thông qua cách tương tác: khi con người đưa ra phản hồi đúng, lịch sự, và có trách nhiệm, AI có thể học được các mẫu thông tin tích cực; ngược lại, nếu dữ liệu sai lệch hoặc thiếu chính xác, AI có thể học sai.",
        "tag_int": "[INT-ENRICH: UNESCO Digital Citizenship & Responsible Interaction]",
        "int_text": "Thể hiện ý thức công dân số: Luôn giao tiếp lịch thiệp, cung cấp phản hồi chuẩn xác vì hiểu rằng hành vi của con người định hình hành vi của máy."
    },
    {
        "stt": 33, "grade": 2, "strand": "B. Đạo đức AI",
        "subtopic": "B1. Các khía cạnh đạo đức của AI – Sự đối xử không công bằng",
        "code": "2.B1.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết được rằng AI đôi khi có thể thiên kiến, tức là đối xử không công bằng với một số nhóm người.",
        "tag_int": "[INT-ENRICH: UNESCO AI Ethics - Fairness & Bias]",
        "int_text": "Bước đầu làm quen với khái niệm \"thiên lệch\": AI có thể nhận nhầm hoặc đối xử khác biệt nếu hệ thống không được dạy công bằng."
    },
    {
        "stt": 34, "grade": 2, "strand": "B. Đạo đức AI",
        "subtopic": "B1. Các khía cạnh đạo đức của AI – Sự đối xử không công bằng",
        "code": "2.B1.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Nêu được ví dụ về các tình huống thể hiện sự thiên kiến của AI.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 5 - Fairness Examples]",
        "int_text": "Đưa ra tình huống trực quan (ví dụ: máy nhận diện giọng nói không nghe rõ tiếng bạn nhỏ hoặc giọng vùng miền khác) để chỉ ra biểu hiện thiên lệch của AI."
    },
    {
        "stt": 35, "grade": 2, "strand": "B. Đạo đức AI",
        "subtopic": "B1. Các khía cạnh đạo đức của AI – Sự đối xử không công bằng",
        "code": "2.B1.MR2", "tag_vn": "[VN-MR]",
        "qd_text": "Giải thích được rằng nếu dữ liệu chưa đa dạng hoặc chưa công bằng thì AI có thể học theo cách thiên vị.",
        "tag_int": "[INT-ENRICH: Experience AI - Training Data Bias]",
        "int_text": "Nêu được mối quan hệ bản chất: \"Dữ liệu đưa vào như thế nào thì máy học ra như thế đó\" (Garbage in, Garbage out ở mức độ trực quan)."
    },
    {
        "stt": 36, "grade": 2, "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "code": "2.B3.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Kể được tên một số thứ là của riêng em và một số thứ là của người khác.",
        "tag_int": "[INT-ENRICH: UNESCO Privacy Foundations]",
        "int_text": "Phân định rõ ràng ranh giới tài sản cá nhân và đồ dùng của người khác trong không gian thực tế và môi trường thông tin số."
    },
    {
        "stt": 37, "grade": 2, "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "code": "2.B3.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được ví dụ về quyền sở hữu đối với sản phẩm do con người hoặc AI tạo ra.",
        "tag_int": "[INT-ENRICH: UNESCO Intellectual Property & Attribution]",
        "int_text": "Nhận biết quyền tác giả: Tranh do bạn vẽ thuộc về bạn, ảnh AI sinh ra cần ghi nhận nguồn gốc rõ ràng, không được nhận vơ là của mình."
    },
    {
        "stt": 38, "grade": 2, "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "code": "2.B3.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Phân biệt được thứ của mình với thứ của người khác trong một vài tình huống đơn giản.",
        "tag_int": "[INT-ENRICH: OECD Ethical Behavior]",
        "int_text": "Thực hành phản xạ tôn trọng quyền sở hữu trong các hoạt động nhóm, không tùy tiện sử dụng tài khoản hay bài làm của bạn."
    },
    {
        "stt": 39, "grade": 2, "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "code": "2.B3.MR2", "tag_vn": "[VN-MR]",
        "qd_text": "Nêu được rằng muốn dùng thứ của người khác thì phải hỏi và được người đó đồng ý.",
        "tag_int": "[INT-ENRICH: UNESCO Ethics - Consent Principle]",
        "int_text": "Hình thành thói quen đạo đức nền tảng: Nguyên tắc Đồng thuận (Consent) - luôn xin phép trước khi chạm vào đồ dùng hoặc mượn thông tin/hình ảnh của người khác."
    },
    {
        "stt": 40, "grade": 2, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Cách AI học và học liệu của AI",
        "code": "2.C1.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Giải thích được “dữ liệu” là những ví dụ (hình ảnh, âm thanh) mà con người dùng để dạy cho AI.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 3 - Data Concepts]",
        "int_text": "Chuẩn hóa định nghĩa khoa học máy tính cơ bản: Dữ liệu không phải là khái niệm trừu tượng mà chính là tập hợp các hình ảnh, âm thanh cụ thể mà ta gom lại để huấn luyện máy."
    },
    {
        "stt": 41, "grade": 2, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Cách AI học và học liệu của AI",
        "code": "2.C1.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "So sánh được ở mức độ cơ bản giữa cách học của con người và AI.",
        "tag_int": "[INT-ENRICH: OECD Cognitive Science - Human vs Machine Learning]",
        "int_text": "Chỉ ra sự khác biệt kỳ diệu: Trẻ em học con mèo qua quan sát đời thật và cảm xúc; AI học con mèo bằng cách xem hàng nghìn bức ảnh số hoá."
    },
    {
        "stt": 42, "grade": 2, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C3. Công nghệ AI – Sơ lược cách AI phân loại đồ vật",
        "code": "2.C3.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nhận biết được rằng AI có thể phân loại.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 2 - Classification Tasks]",
        "int_text": "Nhận diện bài toán phân loại (Sorting/Classification) là nhiệm vụ nền tảng nhất của các hệ thống AI xung quanh ta."
    },
    {
        "stt": 43, "grade": 2, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C3. Công nghệ AI – Sơ lược cách AI phân loại đồ vật",
        "code": "2.C3.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được rằng AI có thể phân loại sai.",
        "tag_int": "[INT-ENRICH: UNESCO Critical Thinking - Model Limitations]",
        "int_text": "Khắc sâu bài học: Máy móc không vạn năng; AI có thể nhầm lẫn quả cà chua với quả táo đỏ nếu dữ liệu chưa đủ rõ."
    },
    {
        "stt": 44, "grade": 2, "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C3. Công nghệ AI – Sơ lược cách AI phân loại đồ vật",
        "code": "2.C3.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Quan sát được thao tác phân loại đồ vật do giáo viên thực hiện trên thiết bị điện tử hoặc mô phỏng; từ kết quả quan sát đó, so sánh được cách AI phân loại với cách con người phân loại.",
        "tag_int": "[INT-ENRICH: MIT Teachable Machine Demonstration]",
        "int_text": "Đối chiếu tiêu chí phân loại: Con người dựa trên nhiều đặc điểm tổng hợp, máy tính thường dựa trên các đặc trưng thị giác (màu sắc, hình dạng viền)."
    },
    {
        "stt": 45, "grade": 2, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Máy thông minh giúp giải quyết vấn đề quanh em",
        "code": "2.D1.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được một số vấn đề đơn giản, gần gũi trong đời sống có thể áp dụng AI để giải quyết.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 5 - Problem Identification]",
        "int_text": "Phát hiện các vấn đề quen thuộc trong lớp học hoặc gia đình (phân loại rác, nhắc nhở rửa tay, tìm sách vở) có thể áp dụng giải pháp thông minh."
    },
    {
        "stt": 46, "grade": 2, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Ý tưởng máy thông minh",
        "code": "2.D1.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Nêu được một số ý tưởng máy thông minh.",
        "tag_int": "[INT-ENRICH: UNESCO AI System Design - Creative Ideation]",
        "int_text": "Vẽ tranh hoặc mô tả bằng lời ý tưởng một chú robot/cỗ máy thông minh giúp ích cho bạn bè hoặc người thân."
    },
    {
        "stt": 47, "grade": 2, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Ý tưởng máy thông minh",
        "code": "2.D1.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Nêu một số ví dụ phù hợp để “dạy” AI trong một tình huống cụ thể.",
        "tag_int": "[INT-ENRICH: Experience AI - Curating Training Examples]",
        "int_text": "Đóng vai \"huấn luyện viên máy tính\", chọn đúng tập hợp tranh ảnh mẫu (đủ góc độ, màu sắc) để dạy máy nhận biết một đồ vật cụ thể."
    },
    {
        "stt": 48, "grade": 2, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Vai trò của dữ liệu",
        "code": "2.D2.1", "tag_vn": "[VN-CORE]",
        "qd_text": "Giải thích được ở mức độ cơ bản về vai trò của dữ liệu trong việc “dạy” AI trong một tình huống cụ thể.",
        "tag_int": "[INT-ENRICH: AI4K12 Big Idea 3 - Importance of Data]",
        "int_text": "Giải thích được vì sao dữ liệu là \"thức ăn\" giúp AI lớn lên và thông minh hơn, thiếu dữ liệu thì mô hình không thể hoạt động."
    },
    {
        "stt": 49, "grade": 2, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Vai trò của dữ liệu",
        "code": "2.D2.2", "tag_vn": "[VN-CORE]",
        "qd_text": "Giải thích được ở mức độ cơ bản rằng cần cung cấp cho AI dữ liệu chính xác, rõ ràng để AI đưa ra kết quả đúng.",
        "tag_int": "[INT-ENRICH: OECD Data Quality Assurance]",
        "int_text": "Nhận thức được tầm quan trọng của chất lượng dữ liệu: Ảnh bị mờ, méo mó hoặc nhầm nhãn sẽ khiến AI dự đoán sai lệch."
    },
    {
        "stt": 50, "grade": 2, "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Vai trò của dữ liệu",
        "code": "2.D2.MR1", "tag_vn": "[VN-MR]",
        "qd_text": "Thực hành thu thập một số dữ liệu để “dạy” AI trong một tình huống cụ thể.",
        "tag_int": "[INT-ENRICH: UNESCO Hands-on System Design & AI4K12 Unplugged]",
        "int_text": "Tự tay thu thập và sắp xếp 10-15 tấm thẻ hình ảnh mẫu rõ nét để chuẩn bị bộ dữ liệu huấn luyện cho một bài toán phân loại đơn giản."
    }
]

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def create_full_docx(output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t1 = p_title.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    run_t1.font.name = "Times New Roman"
    run_t1.font.size = Pt(11)
    run_t1.font.bold = True
    
    run_t2 = p_title.add_run("KHUNG CHƯƠNG TRÌNH DẠY HỌC TRÍ TUỆ NHÂN TẠO (AI)\nCỤM TIỂU HỌC KHỞI ĐẦU: LỚP 1 VÀ LỚP 2\n")
    run_t2.font.name = "Times New Roman"
    run_t2.font.size = Pt(14)
    run_t2.font.bold = True
    run_t2.font.color.rgb = RGBColor(0, 32, 96)
    
    run_sub = p_title.add_run("(Mô hình Tích hợp Làm giàu Quốc tế - Hybrid Enriched Framework: QĐ 2422/QĐ-BGDĐT + UNESCO + AI4K12)")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    
    doc.add_paragraph()
    
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(0.5), Inches(0.6), Inches(1.5), Inches(1.6), Inches(2.8)]
    headers = [
        "STT", "Lớp", "Mạch nội dung", "Chủ đề con",
        "Yêu cầu cần đạt chuẩn hóa tích hợp\n(Chuẩn Bộ GD&ĐT + Làm giàu Quốc tế)"
    ]
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        set_cell_shading(hdr_cells[i], "1F4E79")
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    for item in data_items:
        row_cells = table.add_row().cells
        for idx in range(5):
            row_cells[idx].width = col_widths[idx]
            set_cell_margins(row_cells[idx], top=80, bottom=80, left=100, right=100)
            
        if item["stt"] % 2 == 0:
            for c in row_cells:
                set_cell_shading(c, "F2F5F9")
                
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(str(item["stt"]))
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9)
        r0.font.bold = True
        
        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(f"Lớp {item['grade']}")
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(9)
        
        p2 = row_cells[2].paragraphs[0]
        r2 = p2.add_run(item["strand"])
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(9)
        r2.font.bold = True
        
        p3 = row_cells[3].paragraphs[0]
        r3 = p3.add_run(item["subtopic"])
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(9)
        
        p4 = row_cells[4].paragraphs[0]
        r_code = p4.add_run(f"{item['tag_vn']} {item['code']}: ")
        r_code.font.name = "Times New Roman"
        r_code.font.size = Pt(9)
        r_code.font.bold = True
        if "MR" in item["tag_vn"]:
            r_code.font.color.rgb = RGBColor(192, 0, 0)
        else:
            r_code.font.color.rgb = RGBColor(0, 32, 96)
            
        r_qd = p4.add_run(item["qd_text"] + "\n\n")
        r_qd.font.name = "Times New Roman"
        r_qd.font.size = Pt(9)
        
        r_int_hdr = p4.add_run(f"🎯 {item['tag_int']}: ")
        r_int_hdr.font.name = "Times New Roman"
        r_int_hdr.font.size = Pt(8.5)
        r_int_hdr.font.bold = True
        r_int_hdr.font.color.rgb = RGBColor(38, 128, 0)
        
        r_int = p4.add_run(item["int_text"])
        r_int.font.name = "Times New Roman"
        r_int.font.size = Pt(8.5)
        r_int.font.italic = True
        
    doc.save(output_path)

def create_clean_md(output_path):
    lines = []
    lines.append("# KHUNG CHƯƠNG TRÌNH DẠY HỌC TRÍ TUỆ NHÂN TẠO (AI)")
    lines.append("## CỤM TIỂU HỌC KHỞI ĐẦU: LỚP 1 VÀ LỚP 2")
    lines.append("*(Bản tinh gọn – Giữ nguyên toàn bộ nội dung giáo dục cốt lõi và quốc tế, lược bỏ các ký hiệu kỹ thuật)*\n")
    lines.append("| STT | Lớp | Mạch nội dung | Chủ đề con | Yêu cầu cần đạt chuẩn hóa |")
    lines.append("|:---:|:---:|:---|:---|:---|")
    
    for item in data_items:
        stt = item["stt"]
        grade = f"Lớp {item['grade']}"
        strand = item["strand"]
        subtopic = item["subtopic"]
        code = item["code"]
        qd = item["qd_text"]
        enrich = item["int_text"]
        
        yccd_clean = f"**{code}**: {qd}<br><br>*{enrich}*"
        yccd_clean = yccd_clean.replace("|", "/")
        subtopic_clean = subtopic.replace("|", "/")
        strand_clean = strand.replace("|", "/")
        
        lines.append(f"| {stt} | {grade} | {strand_clean} | {subtopic_clean} | {yccd_clean} |")
        
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

def create_clean_docx(output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t1 = p_title.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    run_t1.font.name = "Times New Roman"
    run_t1.font.size = Pt(11)
    run_t1.font.bold = True
    
    run_t2 = p_title.add_run("KHUNG CHƯƠNG TRÌNH DẠY HỌC TRÍ TUỆ NHÂN TẠO (AI)\nCỤM TIỂU HỌC KHỞI ĐẦU: LỚP 1 VÀ LỚP 2\n")
    run_t2.font.name = "Times New Roman"
    run_t2.font.size = Pt(14)
    run_t2.font.bold = True
    run_t2.font.color.rgb = RGBColor(0, 32, 96)
    
    run_sub = p_title.add_run("(Bản trình bày chuẩn mực sư phạm – Tinh gọn ký hiệu kỹ thuật)")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    
    doc.add_paragraph()
    
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(0.5), Inches(0.6), Inches(1.5), Inches(1.6), Inches(2.8)]
    headers = ["STT", "Lớp", "Mạch nội dung", "Chủ đề con", "Yêu cầu cần đạt chuẩn hóa"]
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        set_cell_shading(hdr_cells[i], "1F4E79")
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    for item in data_items:
        row_cells = table.add_row().cells
        for idx in range(5):
            row_cells[idx].width = col_widths[idx]
            set_cell_margins(row_cells[idx], top=80, bottom=80, left=100, right=100)
            
        if item["stt"] % 2 == 0:
            for c in row_cells:
                set_cell_shading(c, "F2F5F9")
                
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(str(item["stt"]))
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9)
        r0.font.bold = True
        
        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(f"Lớp {item['grade']}")
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(9)
        
        p2 = row_cells[2].paragraphs[0]
        r2 = p2.add_run(item["strand"])
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(9)
        r2.font.bold = True
        
        p3 = row_cells[3].paragraphs[0]
        r3 = p3.add_run(item["subtopic"])
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(9)
        
        p4 = row_cells[4].paragraphs[0]
        r_code = p4.add_run(f"{item['code']}: ")
        r_code.font.name = "Times New Roman"
        r_code.font.size = Pt(9)
        r_code.font.bold = True
        r_code.font.color.rgb = RGBColor(0, 32, 96)
        
        r_qd = p4.add_run(item["qd_text"] + "\n\n")
        r_qd.font.name = "Times New Roman"
        r_qd.font.size = Pt(9)
        
        r_int = p4.add_run(item["int_text"])
        r_int.font.name = "Times New Roman"
        r_int.font.size = Pt(8.5)
        r_int.font.italic = True
        r_int.font.color.rgb = RGBColor(50, 50, 50)
        
    doc.save(output_path)

if __name__ == "__main__":
    out_dir = r"c:\Users\USER\Desktop\AI-vibe-code\KhungNLĐTAI\03_San_Pham_Chinh_Thuc"
    full_docx = os.path.join(out_dir, "SP1_Khung_Chuong_Trinh_AI_Lop1_Lop2_Day_Du.docx")
    clean_docx = os.path.join(out_dir, "SP1_Khung_Chuong_Trinh_AI_Lop1_Lop2_Tinh_Gon.docx")
    clean_md = os.path.join(out_dir, "SP1_Khung_Chuong_Trinh_AI_Lop1_Lop2_Tinh_Gon.md")
    
    create_full_docx(full_docx)
    create_clean_docx(clean_docx)
    create_clean_md(clean_md)
    print("All 3 files generated successfully!")
