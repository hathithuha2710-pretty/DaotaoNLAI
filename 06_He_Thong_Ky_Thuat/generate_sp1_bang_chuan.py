# -*- coding: utf-8 -*-
"""
Script tạo KHUNG CHƯƠNG TRÌNH CHUẨN (SẢN PHẨM 1) - CỤM TIỂU HỌC KHỞI ĐẦU (LỚP 1 VÀ LỚP 2)
Theo đúng hình thức bảng cũ (5 cột):
| STT | Lớp | Mạch nội dung | Chủ đề con | Yêu cầu cần đạt chuẩn hóa |

Nội dung: Dựa trên bản nội dung mới đã được tinh chỉnh khoa học, chuẩn hóa lứa tuổi 6-8 tuổi,
khắc phục hoàn toàn 8 điểm hạn chế theo thẩm định của Hội đồng chuyên môn.
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

items = [
    # ==================== LỚP 1 (21 YCCĐ) ====================
    # Mạch A - Lớp 1 (8 YCCĐ: 7 Core + 1 MR)
    {
        "stt": 1, "grade": "Lớp 1", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – Con người có cảm xúc, AI thì không",
        "code": "1.A1.1",
        "yccd": "Nêu được ví dụ về các cảm xúc quen thuộc của bản thân (vui, buồn, tức giận, sợ hãi) và nhận ra rằng con người cảm thấy vui buồn thật sự, còn máy móc thì không.",
        "pedagogy": "Biết biểu đạt cảm xúc qua nét mặt, giọng nói; nhận biết máy móc chỉ bắt chước hoặc phản hồi theo chương trình do con người tạo ra, không có cảm xúc thật sự."
    },
    {
        "stt": 2, "grade": "Lớp 1", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – Con người có cảm xúc, AI thì không",
        "code": "1.A1.2",
        "yccd": "Nêu được rằng máy thông minh chỉ bắt chước hoặc nhận biết nét mặt của con người, chứ không có cảm xúc như con người.",
        "pedagogy": "Phân biệt được nét mặt vui cười thật sự của bản thân và bạn bè với việc máy tính chỉ quét hình ảnh khuôn mặt để hiện biểu tượng mặt cười."
    },
    {
        "stt": 3, "grade": "Lớp 1", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – AI thể hiện cảm xúc do con người lập trình",
        "code": "1.A1.3",
        "yccd": "Nêu được rằng việc AI thể hiện cảm xúc (như hình mặt cười, giọng nói vui vẻ) là do con người tạo ra hoặc lập trình từ trước.",
        "pedagogy": "Nhận ra mọi cử chỉ, lời nói hay biểu tượng cảm xúc của máy đều do các cô chú kỹ sư cài đặt sẵn, không phải máy tự nhiên nghĩ ra."
    },
    {
        "stt": 4, "grade": "Lớp 1", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – AI thể hiện cảm xúc do con người lập trình",
        "code": "1.A1.4",
        "yccd": "Nêu được ví dụ khi máy phát ra câu chào hoặc cười theo mẫu do con người thiết kế (ví dụ: loa thông minh nói 'Chào bạn, tôi có thể giúp gì?').",
        "pedagogy": "Tỉnh táo khi trò chuyện với loa thông minh hay robot đồ chơi; hiểu rằng câu chào thân thiện chỉ là lời thoại được thiết kế sẵn, không nhầm máy là một người bạn thật sự."
    },
    {
        "stt": 5, "grade": "Lớp 1", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – Ý nghĩa cảm xúc mà AI thể hiện",
        "code": "1.A2.1",
        "yccd": "Chỉ ra được các cách máy thông minh thể hiện cảm xúc (như nháy đèn, hiện hình mặt cười, phát giọng nói nhẹ nhàng).",
        "pedagogy": "Quan sát và gọi tên được các tín hiệu nghe – nhìn (hình vẽ, âm thanh) mà người thiết kế tạo ra để máy giao tiếp thân thiện với con người."
    },
    {
        "stt": 6, "grade": "Lớp 1", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – Ý nghĩa cảm xúc mà AI thể hiện",
        "code": "1.A2.MR1",
        "yccd": "Nêu được rằng giọng nói thân thiện và hình ảnh dễ thương giúp con người cảm thấy gần gũi, dễ chịu khi nhờ máy giúp đỡ.",
        "pedagogy": "Hiểu được ý nghĩa của thiết kế thân thiện: Giúp người dùng cảm thấy thoải mái, yên tâm và dễ hợp tác khi sử dụng đồ dùng thông minh."
    },
    {
        "stt": 7, "grade": "Lớp 1", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – Nhận diện AI trong cuộc sống",
        "code": "1.A2.2",
        "yccd": "Kể tên được một số đồ dùng thông minh quen thuộc có thể tự thực hiện việc đơn giản (như robot hút bụi, camera nhận mặt, loa nghe lệnh).",
        "pedagogy": "Nhận biết các đồ dùng thông minh xung quanh em ở nhà hoặc trường học, phân biệt chúng với đồ dùng cơ học thông thường (như chổi quét nhà, quạt cơ)."
    },
    {
        "stt": 8, "grade": "Lớp 1", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – Nhận diện AI trong cuộc sống",
        "code": "1.A2.3",
        "yccd": "Nói được việc làm có ích của từng đồ dùng thông minh (như robot hút bụi giúp quét nhà sạch sẽ, loa giúp mở bài hát theo lời nói).",
        "pedagogy": "Hiểu rằng đồ dùng thông minh được tạo ra để phục vụ, giảm bớt việc vất vả và mang lại tiện ích cho con người trong cuộc sống hàng ngày."
    },

    # Mạch B - Lớp 1 (3 YCCĐ: 3 Core + 0 MR)
    {
        "stt": 9, "grade": "Lớp 1", "strand": "B. Đạo đức AI",
        "subtopic": "B1. Các khía cạnh đạo đức của AI – Việc làm tốt và việc làm xấu",
        "code": "1.B1.1",
        "yccd": "Nêu được một vài việc không nên làm khi dùng thiết bị thông minh (như bật tiếng quá to làm phiền mọi người, dùng máy trêu chọc bạn).",
        "pedagogy": "Nhận biết hành vi sử dụng công nghệ chưa đúng mực; hình thành ý thức giữ gìn trật tự và tôn trọng mọi người xung quanh."
    },
    {
        "stt": 10, "grade": "Lớp 1", "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Máy thông minh làm việc tốt",
        "code": "1.B3.1",
        "yccd": "Biết và nhắc lại được quy tắc: Không được dùng máy móc thông minh để làm hại hoặc làm đau người khác.",
        "pedagogy": "Khắc sâu nguyên tắc an toàn đạo đức nền tảng: Công nghệ luôn dùng để giúp đỡ, tuyệt đối không được dùng để bắt nạt, nói dối hay gây tổn hại."
    },
    {
        "stt": 11, "grade": "Lớp 1", "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Máy thông minh làm việc tốt",
        "code": "1.B3.2",
        "yccd": "Nêu được việc làm tốt đẹp mà con người sử dụng máy thông minh để thực hiện (như gọi giúp đỡ khi khẩn cấp, học nghe phát âm tiếng Anh).",
        "pedagogy": "Khích lệ tinh thần sử dụng công nghệ đúng mục đích, phục vụ việc học tập và giúp đỡ người thân trong gia đình."
    },

    # Mạch C - Lớp 1 (6 YCCĐ: 4 Core + 2 MR)
    {
        "stt": 12, "grade": "Lớp 1", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Nhận biết AI và ứng dụng AI",
        "code": "1.C1.1",
        "yccd": "Chỉ ra được sự khác biệt đơn giản giữa quạt máy thông thường (bật nút là quay) và quạt thông minh (tự quay khi có người đến gần).",
        "pedagogy": "Nhận biết đặc điểm cơ bản của thiết bị có AI: Có khả năng tự nhận biết môi trường xung quanh để đưa ra phản hồi phù hợp mà không cần bấm nút thủ công."
    },
    {
        "stt": 13, "grade": "Lớp 1", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Nhận biết AI và ứng dụng AI",
        "code": "1.C1.2",
        "yccd": "Nhận biết được hình ảnh biểu tượng trợ lí giọng nói (micro) hoặc ô vuông nhận diện khuôn mặt trên màn hình điện thoại do cô giáo chỉ.",
        "pedagogy": "Làm quen với các biểu tượng giao diện số quen thuộc trên thiết bị thông minh dưới sự hướng dẫn trực quan của giáo viên."
    },
    {
        "stt": 14, "grade": "Lớp 1", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Chức năng và công cụ AI",
        "code": "1.C1.3",
        "yccd": "Nhận biết được camera giống như 'mắt nhìn' và micro giống như 'tai nghe' của thiết bị thông minh.",
        "pedagogy": "Thông qua trải nghiệm đóng vai, hiểu rằng máy thông minh cần có cảm biến (camera, micro) để thu nhận hình ảnh và âm thanh từ thế giới bên ngoài."
    },
    {
        "stt": 15, "grade": "Lớp 1", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Chức năng và công cụ AI",
        "code": "1.C1.4",
        "yccd": "Nêu được rằng máy thông minh có thể hiểu được một số câu nói ngắn, rõ ràng của con người và trả lời lại theo cách con người đã dạy.",
        "pedagogy": "Thực hành phát âm câu lệnh ngắn gọn, rành mạch; hiểu rằng máy chỉ có thể phản hồi tốt khi con người nói rõ ràng và đúng câu lệnh."
    },
    {
        "stt": 16, "grade": "Lớp 1", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Chức năng và công cụ AI",
        "code": "1.C1.MR1",
        "yccd": "Quan sát và nhận biết được hình ảnh camera của máy đóng khung xung quanh quả táo hoặc chú mèo trên màn hình.",
        "pedagogy": "Làm quen với khái niệm thị giác máy tính ở dạng trực quan: Máy nhận biết và đóng khung từng đối tượng cụ thể trong không gian thực tế."
    },
    {
        "stt": 17, "grade": "Lớp 1", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Chức năng và công cụ AI",
        "code": "1.C1.MR2",
        "yccd": "Quan sát được máy thông minh phản ứng khác nhau khi nghe tiếng vỗ tay so với khi nghe tiếng nói của con người.",
        "pedagogy": "Trải nghiệm phân biệt âm thanh: Máy thông minh có thể phân loại các dạng âm thanh khác nhau để thực hiện hành động tương ứng."
    },

    # Mạch D - Lớp 1 (4 YCCĐ: 3 Core + 1 MR)
    {
        "stt": 18, "grade": "Lớp 1", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Máy thông minh học từ ví dụ",
        "code": "1.D1.1",
        "yccd": "Nêu được ví dụ về việc con người cho máy xem nhiều bức tranh con mèo để máy biết nhận ra con mèo.",
        "pedagogy": "Hiểu bản chất ban đầu của việc máy học: Máy tính không tự nhiên biết mà cần con người cung cấp nhiều ví dụ mẫu bằng hình ảnh để học nhận diện."
    },
    {
        "stt": 19, "grade": "Lớp 1", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Máy thông minh học từ ví dụ",
        "code": "1.D1.MR1",
        "yccd": "Giải thích được rằng nếu chỉ đưa 1 bức tranh thì máy dễ đoán sai; máy cần xem nhiều bức tranh khác nhau để học đúng.",
        "pedagogy": "Hiểu nguyên lý đa dạng của dữ liệu: Ví dụ mẫu cần có nhiều màu sắc, kích thước và góc nhìn khác nhau thì máy mới nhận diện chính xác."
    },
    {
        "stt": 20, "grade": "Lớp 1", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Nhiều loại máy thông minh",
        "code": "1.D2.1",
        "yccd": "Nhận biết được có đồ dùng chỉ làm được một việc riêng (như cân điện tử), có đồ dùng làm được nhiều việc (như điện thoại vừa nghe, vừa chụp ảnh).",
        "pedagogy": "Phân loại thiết bị theo phạm vi nhiệm vụ: Nhóm thiết bị chỉ chuyên làm một việc và nhóm thiết bị đa năng thực hiện được nhiều công việc khác nhau."
    },
    {
        "stt": 21, "grade": "Lớp 1", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Nhiều loại máy thông minh",
        "code": "1.D2.2",
        "yccd": "Xếp được các thẻ tranh đồ dùng thông minh vào 2 nhóm: nhóm chỉ làm 1 việc và nhóm làm được nhiều việc khác nhau.",
        "pedagogy": "Thực hành kỹ năng phân loại qua hoạt động không dùng máy tính (Unplugged) với bộ thẻ tranh trực quan, rèn luyện tư duy logic."
    },

    # ==================== LỚP 2 (29 YCCĐ) ====================
    # Mạch A - Lớp 2 (11 YCCĐ: 7 Core + 4 MR)
    {
        "stt": 22, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – Khi nào nên và không nên dùng AI",
        "code": "2.A1.1",
        "yccd": "Nêu được một số việc máy thông minh giúp con người làm nhanh hơn (như tìm đường đi, tìm một bài hát yêu thích, dịch một câu chào tiếng nước ngoài).",
        "pedagogy": "Lựa chọn đúng tình huống cần công cụ thông minh trợ giúp để tiết kiệm thời gian và nâng cao hiệu quả công việc cho con người."
    },
    {
        "stt": 23, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – Khi nào nên và không nên dùng AI",
        "code": "2.A1.2",
        "yccd": "Nêu được tình huống không nên dùng máy thông minh (như không nói địa chỉ nhà, tên bố mẹ cho máy; không để máy làm hết bài tập về nhà của mình).",
        "pedagogy": "Nhận diện ranh giới an toàn: Bảo vệ thông tin bí mật cá nhân và gia đình; rèn luyện thói quen tự suy nghĩ, không ỷ lại công nghệ."
    },
    {
        "stt": 24, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – AI làm việc, con người kiểm soát",
        "code": "2.A1.3",
        "yccd": "Nêu được ví dụ con người luôn cần để mắt kiểm tra khi máy làm việc (như người lái xe vẫn phải nhìn đường khi xe tự chạy, bác sĩ kiểm tra lại kết quả gợi ý của máy).",
        "pedagogy": "Hiểu được nguyên tắc con người làm chủ: AI chỉ là công cụ hỗ trợ, con người phải luôn quan sát, kiểm tra và chịu trách nhiệm về quyết định cuối cùng."
    },
    {
        "stt": 25, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A1. Tính chủ động của con người – AI làm việc, con người kiểm soát",
        "code": "2.A1.4",
        "yccd": "Thể hiện được thái độ đúng đắn: Biết dùng máy để học hỏi nhưng luôn tự suy nghĩ và kiểm tra lại chứ không tin tưởng hoàn toàn vào máy.",
        "pedagogy": "Hình thành thói quen kiểm chứng độc lập và tự chủ; hiểu rằng máy móc có thể nhầm lẫn nên cần con người suy xét kỹ càng."
    },
    {
        "stt": 26, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI trong gia đình",
        "code": "2.A2.1",
        "yccd": "Kể tên được một số thiết bị thông minh dùng trong gia đình (như robot quét nhà, điều hòa tự chỉnh nhiệt độ, tivi gợi ý phim hoạt hình em thích).",
        "pedagogy": "Quan sát và liệt kê các ứng dụng công nghệ thông minh trong ngôi nhà, mô tả sự tiện ích mà chúng đem lại cho sinh hoạt hàng ngày."
    },
    {
        "stt": 27, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI trong gia đình",
        "code": "2.A2.MR1",
        "yccd": "Thực hành thử nói một câu lệnh nhờ máy giúp dưới sự cho phép và quan sát của thầy cô giáo.",
        "pedagogy": "Trải nghiệm tương tác an toàn có kiểm soát: Tập phát âm câu lệnh mẫu rõ ràng, đúng ngữ điệu trên thiết bị học tập dưới sự bảo trợ của giáo viên."
    },
    {
        "stt": 28, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI trong gia đình",
        "code": "2.A2.MR2",
        "yccd": "Nhắc lại và làm theo quy tắc: Chỉ sử dụng đồ dùng thông minh khi có sự đồng ý và trông chừng của người lớn.",
        "pedagogy": "Khắc sâu quy tắc an toàn số: Luôn xin phép và có cha mẹ hoặc thầy cô đồng hành khi chạm vào các thiết bị kết nối mạng."
    },
    {
        "stt": 29, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI hỗ trợ mọi người",
        "code": "2.A2.2",
        "yccd": "Kể được các thành viên trong nhà mà máy thông minh có thể giúp đỡ (như nhắc bà uống thuốc, giúp mẹ lau sàn nhà, giúp bố hẹn giờ làm việc).",
        "pedagogy": "Liên hệ công nghệ với sự quan tâm trong gia đình; thấy được máy móc thông minh giúp đỡ từng người thân như thế nào."
    },
    {
        "stt": 30, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A2. AI vì sự tiến bộ của con người – AI hỗ trợ mọi người",
        "code": "2.A2.MR3",
        "yccd": "Nêu được rằng khi máy giúp làm bớt việc nhà, cả nhà sẽ có nhiều thời gian vui chơi, trò chuyện cùng nhau hơn.",
        "pedagogy": "Hiểu giá trị nhân văn cốt lõi: Công nghệ giúp giải phóng sức lao động để mọi người có thêm thời gian gắn kết tình cảm gia đình."
    },
    {
        "stt": 31, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A3. Công dân trong kỉ nguyên AI – Con người dạy AI qua tương tác",
        "code": "2.A3.1",
        "yccd": "Nhận biết được rằng những câu hỏi, lựa chọn của con người có thể được hệ thống lưu lại để làm ví dụ giúp máy học và cải tiến cách trả lời.",
        "pedagogy": "Hiểu rằng thói quen tìm kiếm, lựa chọn video hay câu hỏi của người dùng có thể được hệ thống ghi nhận để gợi ý phù hợp hơn; từ đó cẩn trọng khi tương tác."
    },
    {
        "stt": 32, "grade": "Lớp 2", "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "A3. Công dân trong kỉ nguyên AI – Con người dạy AI qua tương tác",
        "code": "2.A3.MR1",
        "yccd": "Nêu được rằng khi con người dùng lời nói lịch sự, đúng đắn thì sẽ cung cấp những ví dụ tốt cho máy; nếu nói lời thô tục hoặc thông tin sai thì máy có thể bị dạy sai.",
        "pedagogy": "Hình thành ý thức công dân số: Luôn giao tiếp lịch thiệp, cung cấp phản hồi đúng đắn vì dữ liệu con người đưa vào sẽ định hình cách phản hồi của máy."
    },

    # Mạch B - Lớp 2 (7 YCCĐ: 3 Core + 4 MR)
    {
        "stt": 33, "grade": "Lớp 2", "strand": "B. Đạo đức AI",
        "subtopic": "B1. Các khía cạnh đạo đức của AI – Sự đối xử không công bằng",
        "code": "2.B1.1",
        "yccd": "Nhận biết được rằng máy thông minh đôi khi có thể phân biệt hoặc đối xử không công bằng giữa những người khác nhau.",
        "pedagogy": "Làm quen với khái niệm thiên lệch thuật toán ở mức trực quan: Máy móc có thể nhận diện sai hoặc ưu ái một nhóm người nếu việc dạy máy chưa chu đáo."
    },
    {
        "stt": 34, "grade": "Lớp 2", "strand": "B. Đạo đức AI",
        "subtopic": "B1. Các khía cạnh đạo đức của AI – Sự đối xử không công bằng",
        "code": "2.B1.MR1",
        "yccd": "Kể lại một tình huống máy nhận diện không công bằng (như máy nhận diện giọng nói chỉ hiểu giọng bạn miền Bắc mà không hiểu bạn miền Trung hay miền Nam).",
        "pedagogy": "Thông qua câu chuyện tình huống học đường, nhận ra biểu hiện thiên lệch cụ thể của thiết bị thông minh đối với sự đa dạng vùng miền."
    },
    {
        "stt": 35, "grade": "Lớp 2", "strand": "B. Đạo đức AI",
        "subtopic": "B1. Các khía cạnh đạo đức của AI – Sự đối xử không công bằng",
        "code": "2.B1.MR2",
        "yccd": "Giải thích được rằng nếu lúc dạy máy, con người chỉ đưa ví dụ của một nhóm người thì máy sẽ không nhận biết tốt những người khác.",
        "pedagogy": "Hiểu nguyên nhân gốc rễ của thiên lệch: Dữ liệu huấn luyện bị thiên vị hoặc thiếu hụt mẫu sẽ dẫn đến việc máy phân biệt đối xử không công bằng."
    },
    {
        "stt": 36, "grade": "Lớp 2", "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "code": "2.B3.1",
        "yccd": "Phân biệt được đồ dùng, tranh vẽ của riêng em với đồ dùng, tranh vẽ của bạn khác trong lớp học.",
        "pedagogy": "Xác định rõ ràng ranh giới tài sản cá nhân và tôn trọng đồ dùng của người khác trong môi trường học tập hàng ngày."
    },
    {
        "stt": 37, "grade": "Lớp 2", "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "code": "2.B3.2",
        "yccd": "Nêu được ví dụ về việc bức tranh do em tự vẽ là của em, còn bức tranh do máy tạo ra thì phải ghi rõ là máy tạo ra, không nhận là do mình tự vẽ.",
        "pedagogy": "Hình thành đức tính trung thực và liêm chính học thuật ban đầu: Tôn trọng công sức thật sự của bản thân và bạn bè, không gian dối nhận vơ sản phẩm của máy."
    },
    {
        "stt": 38, "grade": "Lớp 2", "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "code": "2.B3.MR1",
        "yccd": "Phân biệt được trong tình huống cụ thể đồ vật nào được phép dùng và đồ vật nào phải hỏi trước khi chạm vào.",
        "pedagogy": "Thực hành phản xạ tôn trọng quyền sở hữu trong các hoạt động nhóm, không tùy tiện cầm đồ dùng hoặc xem bài làm của bạn khi chưa xin phép."
    },
    {
        "stt": 39, "grade": "Lớp 2", "strand": "B. Đạo đức AI",
        "subtopic": "B3. Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "code": "2.B3.MR2",
        "yccd": "Thực hành nói lời xin phép lịch sự trước khi mượn đồ dùng hoặc mượn xem bức tranh của bạn.",
        "pedagogy": "Rèn luyện nguyên tắc Đồng thuận (Consent): Luôn hỏi xin phép và chỉ sử dụng khi nhận được sự đồng ý vui vẻ của chủ sở hữu."
    },

    # Mạch C - Lớp 2 (5 YCCĐ: 3 Core + 2 MR)
    {
        "stt": 40, "grade": "Lớp 2", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Cách AI học và học liệu của AI",
        "code": "2.C1.1",
        "yccd": "Nêu được rằng 'dữ liệu' chính là các hình ảnh, âm thanh mẫu mà con người gom lại để dạy cho máy học.",
        "pedagogy": "Hiểu định nghĩa dữ liệu ở mức độ trực quan, cụ thể: Dữ liệu là tập hợp các bức tranh, âm thanh mẫu do con người chuẩn bị để huấn luyện máy."
    },
    {
        "stt": 41, "grade": "Lớp 2", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C1. Đặc điểm chính của AI – Cách AI học và học liệu của AI",
        "code": "2.C1.MR1",
        "yccd": "Nêu được điểm khác nhau đơn giản: Con người học bằng cách nhìn, chơi đùa và trải nghiệm thật; còn máy học bằng cách xem qua các bức ảnh hoặc âm thanh mẫu.",
        "pedagogy": "So sánh phương thức nhận thức sinh học phong phú của con người với cách học dựa trên phân tích dữ liệu số của máy tính; tự hào về trí tuệ con người."
    },
    {
        "stt": 42, "grade": "Lớp 2", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C3. Công nghệ AI – Sơ lược cách AI phân loại đồ vật",
        "code": "2.C3.1",
        "yccd": "Nhận biết được rằng máy thông minh có thể giúp con người chia đồ vật thành các nhóm khác nhau (như chia nhóm quả cam và quả chuối).",
        "pedagogy": "Làm quen với nhiệm vụ phân loại (Classification) - chức năng cốt lõi của công nghệ học máy và thị giác máy tính trong thực tế."
    },
    {
        "stt": 43, "grade": "Lớp 2", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C3. Công nghệ AI – Sơ lược cách AI phân loại đồ vật",
        "code": "2.C3.2",
        "yccd": "Nêu được ví dụ máy thông minh có thể đoán nhầm khi đồ vật nhìn quá giống nhau (như nhầm quả bóng màu cam với quả cam).",
        "pedagogy": "Nhận ra giới hạn của AI: Máy móc không hoàn hảo và có thể phán đoán sai nếu đồ vật có đặc điểm bên ngoài tương tự nhau; không nên tin máy tuyệt đối."
    },
    {
        "stt": 44, "grade": "Lớp 2", "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "C3. Công nghệ AI – Sơ lược cách AI phân loại đồ vật",
        "code": "2.C3.MR1",
        "yccd": "Quan sát cô giáo làm mẫu trên phần mềm mô phỏng và nói được: Con người nhận ra đồ vật nhanh hơn máy khi đồ vật bị che khuất một phần.",
        "pedagogy": "So sánh khả năng suy luận linh hoạt của mắt người với thuật toán máy tính khi gặp vật cản, nhận ra sự vượt trội của trí tuệ con người."
    },

    # Mạch D - Lớp 2 (6 YCCĐ: 4 Core + 2 MR)
    {
        "stt": 45, "grade": "Lớp 2", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Máy thông minh giúp giải quyết vấn đề quanh em",
        "code": "2.D1.1",
        "yccd": "Nêu được một việc vất vả trong lớp hoặc ở nhà mà em ước có một chú robot thông minh đến giúp đỡ (như lau bảng trên cao, nhặt rác dưới gầm bàn).",
        "pedagogy": "Khởi đầu tư duy giải quyết vấn đề: Xuất phát từ nhu cầu và khó khăn thực tế của con người trước khi nghĩ đến việc sáng tạo công nghệ."
    },
    {
        "stt": 46, "grade": "Lớp 2", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Ý tưởng máy thông minh",
        "code": "2.D1.2",
        "yccd": "Vẽ tranh hoặc kể tên một chú robot mơ ước do em tưởng tượng ra để giúp đỡ mọi người.",
        "pedagogy": "Nuôi dưỡng trí tưởng tượng sáng tạo và tình yêu khoa học: Thể hiện ý tưởng thiết kế một cỗ máy thông minh phục vụ cuộc sống."
    },
    {
        "stt": 47, "grade": "Lớp 2", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D1. Nhận diện và hình thành giải pháp – Ý tưởng máy thông minh",
        "code": "2.D1.MR1",
        "yccd": "Chọn được một số tấm thẻ hình ảnh đúng để chuẩn bị làm 'bài học mẫu' dạy cho chú robot của em nhận biết đồ vật.",
        "pedagogy": "Đóng vai người huấn luyện máy: Tự tay lựa chọn và chuẩn bị bộ ví dụ mẫu phù hợp để dạy cho mô hình AI nhận diện đồ vật."
    },
    {
        "stt": 48, "grade": "Lớp 2", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Vai trò của dữ liệu",
        "code": "2.D2.1",
        "yccd": "Nêu được rằng máy thông minh cần có các ví dụ mẫu (tranh ảnh, âm thanh) do con người chuẩn bị thì mới có thể học và làm việc được.",
        "pedagogy": "Hiểu vai trò của dữ liệu mẫu: Dữ liệu là các bài học mẫu mà con người cung cấp cho máy; nếu không có ví dụ mẫu thì máy không thể học hỏi."
    },
    {
        "stt": 49, "grade": "Lớp 2", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Vai trò của dữ liệu",
        "code": "2.D2.2",
        "yccd": "Giải thích được rằng nếu đưa cho máy những bức ảnh bị mờ, chụp tối hoặc sai tên thì máy sẽ học sai và đưa ra kết quả sai.",
        "pedagogy": "Nhận thức tầm quan trọng của chất lượng dữ liệu: Ảnh mẫu cần rõ ràng, sắc nét và đúng tên thì máy mới học đúng và dự đoán chính xác."
    },
    {
        "stt": 50, "grade": "Lớp 2", "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "D2. Cấu trúc và tương tác, cải tiến hệ thống – Vai trò của dữ liệu",
        "code": "2.D2.MR1",
        "yccd": "Thực hành thu thập và sắp xếp 5-10 bức tranh rõ nét từ sách báo cũ để tạo thành bộ ví dụ mẫu dạy máy nhận diện đồ vật.",
        "pedagogy": "Thực hành Unplugged: Tự tay cắt, dán và chuẩn bị một bộ dữ liệu hình ảnh mẫu rõ nét để mô phỏng quy trình huấn luyện máy trong thực tế."
    }
]

def generate_markdown(output_path):
    lines = []
    lines.append("# KHUNG CHƯƠNG TRÌNH DẠY HỌC TRÍ TUỆ NHÂN TẠO (AI)")
    lines.append("## CỤM TIỂU HỌC KHỞI ĐẦU: LỚP 1 VÀ LỚP 2")
    lines.append("*(Bản chuẩn hóa toàn diện – Chuẩn 5 cột, tích hợp YCCĐ đã tinh chỉnh khoa học và diễn giải sư phạm quốc tế vừa sức 6-8 tuổi)*\n")
    lines.append("| STT | Lớp | Mạch nội dung | Chủ đề con | Yêu cầu cần đạt chuẩn hóa |")
    lines.append("|:---:|:---:|:---|:---|:---|")
    
    for item in items:
        yccd_cell = f"**{item['code']}**: {item['yccd']}<br><br>*{item['pedagogy']}*"
        row_str = f"| {item['stt']} | {item['grade']} | {item['strand']} | {item['subtopic']} | {yccd_cell} |"
        lines.append(row_str)
        
    content = "\n".join(lines)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)

def generate_docx(output_path):
    doc = docx.Document()
    
    # Page setup
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_cell_shading(cell, color_hex):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, size=11, color=None, space_after=4):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if text:
            r = p.add_run(text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            if color:
                r.font.color.rgb = color
        return p

    # Tiêu đề văn bản
    add_p("BỘ GIÁO DỤC VÀ ĐÀO TẠO", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11)
    add_p("DỰ ÁN XÂY DỰNG CHƯƠNG TRÌNH GIÁO DỤC TRÍ TUỆ NHÂN TẠO PHỔ THÔNG", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, color=RGBColor(0, 32, 96))
    add_p("────────────────────────", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, color=RGBColor(128, 128, 128))
    doc.add_paragraph()
    add_p("KHUNG CHƯƠNG TRÌNH GIÁO DỤC TRÍ TUỆ NHÂN TẠO (AI)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, color=RGBColor(192, 0, 0))
    add_p("CỤM TIỂU HỌC KHỞI ĐẦU: LỚP 1 VÀ LỚP 2", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, color=RGBColor(0, 32, 96))
    add_p("(Bản chuẩn hóa theo định dạng 5 cột – Tinh chỉnh nội dung sư phạm vừa sức lứa tuổi 6–8 tuổi\nTuân thủ Quyết định số 2422/QĐ-BGDĐT và Công văn số 5588/BGDĐT-GDPT)", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=10)
    
    doc.add_paragraph()

    # Bảng 5 cột
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["STT", "Lớp", "Mạch nội dung", "Chủ đề con", "Yêu cầu cần đạt chuẩn hóa"]
    col_widths = [Inches(0.5), Inches(0.8), Inches(1.5), Inches(1.8), Inches(3.2)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_shading(hdr_cells[i], "1F4E79")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Đổ dữ liệu 50 dòng
    for idx, item in enumerate(items):
        row_cells = table.add_row().cells
        
        # Zebra striping
        bg_color = "F2F5F9" if idx % 2 == 1 else "FFFFFF"
        for i in range(5):
            set_cell_shading(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=80, bottom=80, left=90, right=90)
            row_cells[i].width = col_widths[i]

        # STT
        row_cells[0].text = str(item["stt"])
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.runs[0].font.name = "Times New Roman"
        p0.runs[0].font.size = Pt(9.5)

        # Lớp
        row_cells[1].text = item["grade"]
        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.runs[0].font.name = "Times New Roman"
        p1.runs[0].font.size = Pt(9.5)
        p1.runs[0].font.bold = True

        # Mạch nội dung
        row_cells[2].text = item["strand"]
        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p2.runs[0].font.name = "Times New Roman"
        p2.runs[0].font.size = Pt(9.5)

        # Chủ đề con
        row_cells[3].text = item["subtopic"]
        p3 = row_cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p3.runs[0].font.name = "Times New Roman"
        p3.runs[0].font.size = Pt(9.5)

        # YCCĐ chuẩn hóa (Gồm Mã đậm + YCCĐ + Diễn giải sư phạm nghiêng)
        cell4 = row_cells[4]
        p4 = cell4.paragraphs[0]
        p4.paragraph_format.line_spacing = 1.15
        p4.paragraph_format.space_after = Pt(4)
        
        # Run 1: Code
        r_code = p4.add_run(f"{item['code']}: ")
        r_code.font.name = "Times New Roman"
        r_code.font.size = Pt(9.5)
        r_code.font.bold = True
        r_code.font.color.rgb = RGBColor(0, 32, 96)
        
        # Run 2: YCCĐ text
        r_yccd = p4.add_run(item["yccd"])
        r_yccd.font.name = "Times New Roman"
        r_yccd.font.size = Pt(9.5)

        # Paragraph 2: Pedagogy
        p4_sub = cell4.add_paragraph()
        p4_sub.paragraph_format.line_spacing = 1.15
        p4_sub.paragraph_format.space_after = Pt(2)
        r_ped = p4_sub.add_run(item["pedagogy"])
        r_ped.font.name = "Times New Roman"
        r_ped.font.size = Pt(9)
        r_ped.font.italic = True
        r_ped.font.color.rgb = RGBColor(80, 80, 80)

    doc.save(output_path)

if __name__ == "__main__":
    out_dir = r"c:\Users\USER\Desktop\AI-vibe-code\KhungNLĐTAI\03_San_Pham_Chinh_Thuc"
    md_path = os.path.join(out_dir, "SP1_Khung_Chuong_Trinh_AI_Lop1_Lop2_Chuan_Hoa.md")
    docx_path = os.path.join(out_dir, "SP1_Khung_Chuong_Trinh_AI_Lop1_Lop2_Chuan_Hoa.docx")
    
    generate_markdown(md_path)
    generate_docx(docx_path)
    print("Exported successfully!")
