# -*- coding: utf-8 -*-
"""
Script tạo các file YCCĐ Lớp 2 chuẩn hóa (Sản phẩm 1):
1. SP1_Khung_YCCD_AI_Lop2_Day_Du.docx (Khổ ngang, 5 cột đối sánh chuẩn Bộ và Quốc tế)
2. SP1_Khung_YCCD_AI_Lop2_Day_Du.md (Markdown 5 cột đối sánh đầy đủ)
3. SP1_Khung_YCCD_AI_Lop2_Rut_Gon.docx (Khổ dọc, 4 cột tinh gọn STT | Mạch năng lực | Chủ đề con | YCCĐ)
4. SP1_Khung_YCCD_AI_Lop2_Rut_Gon.md (Markdown 4 cột rút gọn)
"""
import os
import sys
if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

data_lop2 = [
    # ==================== MẠCH A (11 YCCĐ) ====================
    {
        "stt": 1, "code": "2.A1.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Khi nào nên và không nên dùng AI",
        "yccd_vn": "Trình bày được một số tình huống trong cuộc sống mà AI có thể hỗ trợ con người hiệu quả (như tìm đường đi, tìm bài hát yêu thích, dịch một câu chào, phát hiện lỗi chính tả).",
        "intl_full": "• UNESCO (AI for Humanity): Nhận diện các công việc hàng ngày mà AI có thể hỗ trợ con người làm nhanh và thuận tiện hơn.\n• Singapore MOE (EdTech Masterplan): Nhận biết thời điểm phù hợp để sử dụng công nghệ thông minh trợ giúp công việc học tập và sinh hoạt.",
        "intl_compact": "Nhận biết các công việc quen thuộc mà AI hỗ trợ con người làm nhanh và thuận tiện hơn (như tìm đường, dịch câu, tìm bài hát). Biết lựa chọn thời điểm phù hợp để nhờ máy thông minh trợ giúp."
    },
    {
        "stt": 2, "code": "2.A1.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Khi nào nên và không nên dùng AI",
        "yccd_vn": "Trình bày được một số tình huống không nên hoặc cần thận trọng khi sử dụng AI (như khi AI có thể làm lộ thông tin cá nhân: địa chỉ nhà, khuôn mặt, giọng nói; không để máy làm hết bài tập thay cho mình).",
        "intl_full": "• Singapore MOE (Cyber Wellness - Self-Agency): Rèn luyện tư duy độc lập, kiên quyết chống ỷ lại vào máy; bảo vệ bí mật thông tin cá nhân.\n• OECD (Digital Safety for Children): Nhận thức ranh giới an toàn thông tin cá nhân và tính liêm chính học tập.",
        "intl_compact": "Hình thành hành vi an toàn số và tính tự giác: không chia sẻ thông tin riêng tư (địa chỉ nhà, khuôn mặt, thông tin gia đình) cho máy và không để máy làm hộ toàn bộ bài tập."
    },
    {
        "stt": 3, "code": "2.A1.3",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – AI làm việc, con người kiểm soát",
        "yccd_vn": "Nêu được ví dụ cụ thể về tình huống cần con người giám sát AI (như người lái xe vẫn cần quan sát khi xe tự lái hoạt động; bác sĩ cần kiểm tra lại kết quả do AI gợi ý).",
        "intl_full": "• UNESCO (Human Oversight): Khẳng định nguyên tắc con người luôn giám sát và đưa ra quyết định trong các tình huống ứng dụng AI quan trọng.\n• OECD (Trustworthy AI): Nhận thức tầm quan trọng của việc con người kiểm soát và can thiệp kịp thời vào hoạt động của AI.",
        "intl_compact": "Nhận thức rõ nguyên tắc con người cần giám sát và đưa ra quyết định khi dùng các hệ thống AI trong các tình huống quan trọng (như lái xe, khám chữa bệnh)."
    },
    {
        "stt": 4, "code": "2.A1.4",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – AI làm việc, con người kiểm soát",
        "yccd_vn": "Nêu được ví dụ cụ thể về thái độ đúng đắn khi sử dụng AI: biết tin tưởng vào công nghệ ở mức hợp lý, nhưng không tin tưởng tuyệt đối và luôn có sự kiểm tra, kiểm soát của con người.",
        "intl_full": "• Google for Education (Critically Evaluate): Biết giữ thái độ thận trọng, không phụ thuộc thụ động vào máy móc.\n• Singapore MOE (Critical Literacy): Chủ động tự suy nghĩ, biết hỏi ý kiến người lớn khi cần kiểm tra lại thông tin.",
        "intl_compact": "Hình thành thái độ sử dụng công nghệ đúng mực: biết tin cậy AI ở mức hợp lý nhưng không tin tưởng tuyệt đối; luôn tự suy nghĩ và hỏi ý kiến người lớn khi cần kiểm tra lại."
    },
    {
        "stt": 5, "code": "2.A2.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI trong gia đình",
        "yccd_vn": "Nêu được một số thiết bị hoặc ứng dụng trong gia đình có sử dụng AI (như loa thông minh, robot hút bụi, máy điều hoà tự động điều chỉnh nhiệt độ, camera nhận diện người trong nhà, tivi thông minh gợi ý chương trình yêu thích).",
        "intl_full": "• AI4K12 (Big Idea 5): Nhận diện các ứng dụng AI phổ biến trong không gian gia đình (Smart Home).\n• Singapore MOE (Smart Nation): Làm quen với hệ sinh thái công nghệ hỗ trợ sinh hoạt đời sống.",
        "intl_compact": "Nhận diện và liệt kê các ứng dụng công nghệ thông minh phổ biến trong không gian gia đình; mô tả cách chúng hỗ trợ sinh hoạt hàng ngày."
    },
    {
        "stt": 6, "code": "2.A2.MR1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI trong gia đình",
        "yccd_vn": "Thực hành sử dụng một số thiết bị hoặc ứng dụng có sử dụng AI với sự giám sát, hỗ trợ của giáo viên (như phát âm rõ câu lệnh để nhờ loa thông minh phát bài hát thiếu nhi).",
        "intl_full": "• UNESCO (Human-AI Interaction): Thực hành tương tác an toàn với hệ thống AI có sự giám sát sư phạm.\n• AI4K12 (Big Idea 4 - Natural Interaction): Rèn kỹ năng phát âm câu lệnh rõ ràng trong môi trường học tập an toàn.",
        "intl_compact": "Thực hành tương tác an toàn và chuẩn mực với hệ thống trợ lý thông minh; phát âm câu lệnh rõ ràng, mạch lạc dưới sự hướng dẫn của giáo viên."
    },
    {
        "stt": 7, "code": "2.A2.MR2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI trong gia đình",
        "yccd_vn": "Nêu được rằng không được sử dụng thiết bị hoặc ứng dụng có sử dụng AI mà không có sự giám sát của người lớn.",
        "intl_full": "• OECD (Digital Protection for Minors): Tuân thủ quy định an toàn số dành cho lứa tuổi nhỏ; luôn có sự đồng hành của người bảo hộ.\n• Common Sense Education: Kỷ luật số trong gia đình và trường học.",
        "intl_compact": "Tuân thủ nghiêm ngặt quy tắc an toàn số: luôn xin phép và có sự đồng hành, giám sát của người lớn khi sử dụng các thiết bị công nghệ thông minh."
    },
    {
        "stt": 8, "code": "2.A2.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI hỗ trợ mọi người",
        "yccd_vn": "Nhận biết và kể tên được các thành viên trong gia đình mà AI có thể hỗ trợ (như nhắc bà uống thuốc đúng giờ, giúp mẹ lau sàn nhà, giúp bố hẹn giờ làm việc).",
        "intl_full": "• UNESCO (AI for Social Well-being): Nhận biết giá trị nhân văn của AI trong việc chăm sóc và hỗ trợ các thế hệ trong gia đình.\n• OECD (Inclusive Technology): Công nghệ hướng tới hỗ trợ đa dạng nhu cầu con người.",
        "intl_compact": "Nhận biết và trình bày giá trị nhân văn của công nghệ AI trong việc chăm sóc, hỗ trợ công việc và sinh hoạt cho các thế hệ khác nhau trong gia đình."
    },
    {
        "stt": 9, "code": "2.A2.MR3",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI hỗ trợ mọi người",
        "yccd_vn": "Nêu được rằng mục tiêu của AI trong bối cảnh gia đình là giúp ích cho cuộc sống của các thành viên tiện nghi hơn, an toàn hơn và có nhiều thời gian cho nhau hơn.",
        "intl_full": "• OECD (Digital Well-being - Quality of Life): Hiểu mục đích của công nghệ trong gia đình là hỗ trợ đời sống tiện nghi, an toàn và gắn kết.\n• Singapore MOE (Family Connectedness): Công nghệ giúp tiết kiệm thời gian làm việc nhà để vun đắp tình cảm gia đình.",
        "intl_compact": "Nhận biết mục đích nhân văn của AI trong gia đình: hỗ trợ các công việc nhà để cuộc sống tiện nghi, an toàn hơn và giúp các thành viên có thêm thời gian vui chơi, trò chuyện cùng nhau."
    },
    {
        "stt": 10, "code": "2.A3.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Công dân trong kỉ nguyên AI – Con người dạy AI qua tương tác",
        "yccd_vn": "Nhận biết được rằng mỗi lần con người nói chuyện, đặt câu hỏi, tìm kiếm hoặc đưa ra lựa chọn với các hệ thống AI (như phần mềm hội thoại tự động, ứng dụng gợi ý video hoặc nhạc), thì dữ liệu tương tác đó có thể được ghi nhận/sử dụng tùy hệ thống để học hỏi cách con người phản hồi.",
        "intl_full": "• AI4K12 (Big Idea 3 - Learning from User Interaction): Nhận biết rằng câu hỏi và lựa chọn của người dùng có thể được ghi nhận/sử dụng tùy hệ thống làm dữ liệu cải tiến mô hình.\n• Common Sense Education (Digital Footprint): Ý thức bước đầu về dấu chân kỹ thuật số khi tương tác trực tuyến.",
        "intl_compact": "Nhận biết rằng các câu hỏi và lựa chọn của người dùng có thể được hệ thống ghi nhận/sử dụng tùy hệ thống làm dữ liệu học để cải tiến phản hồi; bước đầu làm quen với khái niệm dấu chân kỹ thuật số."
    },
    {
        "stt": 11, "code": "2.A3.MR1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Công dân trong kỉ nguyên AI – Con người dạy AI qua tương tác",
        "yccd_vn": "Nêu được rằng con người có vai trò quan trọng trong việc cung cấp thông tin cho AI thông qua cách tương tác: khi con người phản hồi đúng, lịch sự và có trách nhiệm thì giúp hệ thống có thông tin tích cực; ngược lại, nếu cung cấp thông tin sai lệch thì AI có thể đưa ra kết quả sai.",
        "intl_full": "• Singapore MOE (Digital Civility): Trách nhiệm văn hóa ứng xử số; tương tác văn minh và trung thực khi dùng công nghệ.\n• UNESCO (Ethical Data Contribution): Ý thức trách nhiệm của người dùng trong việc cung cấp thông tin đúng đắn, lành mạnh.",
        "intl_compact": "Xác định trách nhiệm văn hóa ứng xử số và tính trung thực: giao tiếp lịch thiệp, cung cấp thông tin đúng đắn và có trách nhiệm khi tương tác với các ứng dụng công nghệ."
    },

    # ==================== MẠCH B (7 YCCĐ) ====================
    {
        "stt": 12, "code": "2.B1.1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Các khía cạnh đạo đức của AI – Sự đối xử không công bằng",
        "yccd_vn": "Nhận biết được rằng AI đôi khi có thể thiên kiến (thiên lệch), tức là đối xử không công bằng với một số nhóm người.",
        "intl_full": "• Google DeepMind (Experience AI - Bias in, bias out): Nhận diện sơ khai hiện tượng thiên lệch (bias) trong các ứng dụng AI.\n• UNESCO (Fairness & Non-discrimination): Nhận biết máy móc không hoàn hảo và có thể đối xử thiếu công bằng.",
        "intl_compact": "Làm quen với khái niệm thiên lệch (bias); nhận biết máy thông minh có thể đưa ra kết quả thiếu công bằng hoặc phân biệt giữa các nhóm người."
    },
    {
        "stt": 13, "code": "2.B1.MR1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Các khía cạnh đạo đức của AI – Sự đối xử không công bằng",
        "yccd_vn": "Nêu được ví dụ về các tình huống thể hiện sự thiên kiến (thiên lệch) của AI (như máy nhận diện giọng nói chỉ hiểu giọng bạn miền Bắc mà chưa hiểu được giọng bạn miền Trung hay miền Nam).",
        "intl_full": "• Common Sense Education (AI Bias Case Studies): Phân tích tình huống thiên lệch giọng nói do thiếu dữ liệu đa dạng vùng miền.\n• AI4K12 (Big Idea 5 - Diversity): Nhận diện sai lệch của mô hình xử lý ngôn ngữ khi đối mặt với sự đa dạng ngữ âm.",
        "intl_compact": "Nêu được ví dụ cụ thể về sự thiên lệch của công nghệ (như nhận diện giọng nói vùng miền); chỉ ra nguyên nhân do hệ thống chưa được làm quen với sự đa dạng ngôn ngữ."
    },
    {
        "stt": 14, "code": "2.B1.MR2",
        "strand": "B. Đạo đức AI",
        "subtopic": "Các khía cạnh đạo đức của AI – Sự đối xử không công bằng",
        "yccd_vn": "Giải thích được rằng nếu dữ liệu chưa đa dạng hoặc chưa công bằng thì AI có thể học theo cách thiên vị (như chỉ dạy máy bằng hình ảnh một nhóm người thì máy sẽ nhận diện kém những người khác).",
        "intl_full": "• Google DeepMind (Foundations of AI - Dataset Diversity): Giải thích nguyên lý: tập dữ liệu huấn luyện bị lệch mẫu sẽ dẫn đến mô hình bị thiên vị.\n• OECD (Data Quality & Equity): Tầm quan trọng của tính đại diện dữ liệu đối với sự công bằng xã hội.",
        "intl_compact": "Giải thích được nguyên nhân gốc rễ của thiên lệch công nghệ: nếu dữ liệu huấn luyện chỉ đại diện cho một nhóm hẹp, mô hình sẽ nhận diện kém và thiếu công bằng với những nhóm khác."
    },
    {
        "stt": 15, "code": "2.B3.1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "yccd_vn": "Kể được tên một số thứ là của riêng em và một số thứ là của người khác (như hộp bút, quyển vở, bức tranh của em vẽ và của bạn khác trong lớp học).",
        "intl_full": "• UNESCO (Respect for Intellectual Property): Ý thức nền tảng về ranh giới quyền sở hữu sản phẩm cá nhân trong môi trường học đường.\n• Singapore MOE (CCE - Respecting Others): Tôn trọng sản phẩm và công sức sáng tạo của bạn bè.",
        "intl_compact": "Xác định rõ ràng ranh giới quyền sở hữu sản phẩm cá nhân; tôn trọng thành quả và công sức lao động học tập của bản thân và bạn bè."
    },
    {
        "stt": 16, "code": "2.B3.2",
        "strand": "B. Đạo đức AI",
        "subtopic": "Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "yccd_vn": "Nêu được ví dụ về quyền sở hữu đối với sản phẩm do con người hoặc AI tạo ra (như bức tranh do em tự vẽ là của em; còn bức tranh do máy tạo ra thì phải ghi rõ là máy tạo ra, không nhận là do mình tự vẽ).",
        "intl_full": "• UNESCO (Academic Integrity in GenAI): Khởi đầu của liêm chính học thuật: minh bạch nguồn gốc sản phẩm do AI tạo ra (Attribution).\n• Common Sense Education (AI & Cheating): Trung thực số, không nhận tác phẩm của AI là do mình tự sáng tác.",
        "intl_compact": "Hình thành ý thức liêm chính học thuật và trung thực số ban đầu: công khai minh bạch nguồn gốc sản phẩm do AI tạo ra, không nhận sản phẩm của máy là do mình tự sáng tác."
    },
    {
        "stt": 17, "code": "2.B3.MR1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "yccd_vn": "Phân biệt được thứ của mình với thứ của người khác trong một vài tình huống đơn giản (như đồ vật nào được phép tự do dùng và đồ vật nào phải hỏi trước khi chạm vào).",
        "intl_full": "• OECD (Digital Ethics & Boundaries): Nhận diện ranh giới cho phép (Permissions & Boundaries) khi sử dụng tài sản chung và riêng.\n• Singapore MOE (Cyber Wellness): Tôn trọng quyền riêng tư và quyền sở hữu trong các hoạt động hợp tác.",
        "intl_compact": "Phân định rõ ràng ranh giới quyền hạn sử dụng tài sản trong tình huống thực tế; nhận biết đồ vật nào được tự do dùng và đồ vật nào phải xin phép trước khi sử dụng."
    },
    {
        "stt": 18, "code": "2.B3.MR2",
        "strand": "B. Đạo đức AI",
        "subtopic": "Nguyên tắc đạo đức và trách nhiệm xã hội – Của bạn và của tớ",
        "yccd_vn": "Nêu được rằng muốn dùng thứ của người khác thì phải hỏi và được người đó đồng ý; thực hành nói lời xin phép lịch sự trước khi mượn đồ dùng hoặc tranh vẽ của bạn.",
        "intl_full": "• UNESCO (Principle of Consent): Thực hành nguyên tắc đồng thuận (Consent) văn minh trong giao tiếp và hợp tác học tập.\n• CASEL (Relationship Skills): Rèn luyện kỹ năng quan hệ xã hội tôn trọng và lịch thiệp.",
        "intl_compact": "Rèn luyện nguyên tắc đồng thuận văn minh: luôn hỏi ý kiến và chỉ sử dụng tài sản, sản phẩm của người khác khi nhận được sự đồng ý vui vẻ."
    },

    # ==================== MẠCH C (5 YCCĐ) ====================
    {
        "stt": 19, "code": "2.C1.1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Đặc điểm chính của AI – Cách AI học và học liệu của AI",
        "yccd_vn": "Giải thích được “dữ liệu” là những ví dụ (hình ảnh, âm thanh) mà con người dùng để dạy cho AI.",
        "intl_full": "• AI4K12 (Big Idea 3 - What is Data): Định nghĩa dữ liệu trực quan: tập hợp hình ảnh, âm thanh mẫu phục vụ việc dạy máy tính.\n• Google DeepMind (Experience AI): Hiểu dữ liệu là nguồn cung cấp kiến thức nền tảng để xây dựng mô hình AI.",
        "intl_compact": "Nêu được khái niệm cơ bản về dữ liệu: là tập hợp các hình ảnh, âm thanh mẫu do con người thu thập và chọn lọc để làm bài học dạy cho AI."
    },
    {
        "stt": 20, "code": "2.C1.MR1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Đặc điểm chính của AI – Cách AI học và học liệu của AI",
        "yccd_vn": "So sánh được ở mức độ cơ bản giữa cách học của con người và AI (con người học qua nhìn, chơi đùa và trải nghiệm thật trong cuộc sống; còn AI học bằng cách xem qua nhiều tranh ảnh, âm thanh mẫu do con người nạp vào).",
        "intl_full": "• UNESCO (Human Intelligence vs Machine Learning): So sánh phương thức học tập trực tiếp đa giác quan của con người với cách máy học tính toán từ dữ liệu số.\n• AI4K12 (Big Idea 3): Hiểu bản chất khác biệt giữa nhận thức sinh học và học máy.",
        "intl_compact": "So sánh sự khác biệt cơ bản: con người học tập qua trải nghiệm thực tế và cảm xúc đa giác quan; máy tính học bằng cách tính toán và phân tích các tập dữ liệu mẫu số hóa."
    },
    {
        "stt": 21, "code": "2.C3.1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Công nghệ AI – Sơ lược cách AI phân loại đồ vật",
        "yccd_vn": "Nhận biết được rằng AI có thể phân loại (như nhận biết và chia đồ vật thành nhóm quả cam và nhóm quả chuối).",
        "intl_full": "• Google DeepMind (Experience AI - Classification): Nhận diện bài toán phân loại (Classification) như một dạng bài toán quen thuộc của học máy.\n• AI4K12 (Big Idea 3 - Pattern Categorization): Khả năng nhận dạng đặc trưng để phân nhóm đối tượng của AI.",
        "intl_compact": "Nhận biết bài toán phân loại (Classification) – một dạng bài toán quen thuộc của học máy: hệ thống AI có khả năng nhận diện đặc trưng đơn giản để chia các đối tượng vào từng nhóm cụ thể."
    },
    {
        "stt": 22, "code": "2.C3.2",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Công nghệ AI – Sơ lược cách AI phân loại đồ vật",
        "yccd_vn": "Nêu được rằng AI có thể phân loại sai (như khi đồ vật nhìn quá giống nhau: nhầm quả bóng màu cam với quả cam).",
        "intl_full": "• AI4K12 (Big Idea 1 & 3 - Edge Cases): Nhận thức về giới hạn của thị giác máy tính khi các đặc trưng bên ngoài quá tương đồng.\n• OECD (Critical Digital Skills): Hiểu nguyên nhân gây sai sót của mô hình AI khi thiếu đặc trưng chiều sâu.",
        "intl_compact": "Nhận thức rõ giới hạn kỹ thuật của thị giác máy tính: hệ thống AI có thể nhận định sai lầm nếu các đối tượng có đặc trưng hình ảnh bên ngoài quá tương đồng nhau."
    },
    {
        "stt": 23, "code": "2.C3.MR1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Công nghệ AI – Sơ lược cách AI phân loại đồ vật",
        "yccd_vn": "Quan sát được thao tác phân loại đồ vật do giáo viên thực hiện trên thiết bị điện tử hoặc mô phỏng; từ kết quả quan sát đó, so sánh được cách AI phân loại với cách con người phân loại (nêu được rằng khi đồ vật bị che khuất một phần, con người có thể tận dụng kinh nghiệm và ngữ cảnh thực tế để nhận ra, còn AI có thể gặp khó khăn hoặc phân loại nhầm).",
        "intl_full": "• UNESCO (Contextual Inference): Nhận biết khả năng liên tưởng ngữ cảnh của con người và giới hạn của AI khi đồ vật bị che khuất một phần (Occlusion).\n• AI4K12 (Big Idea 1 - Perception Limitations): Hạn chế của AI khi đối mặt với vật thể bị che khuất.",
        "intl_compact": "So sánh khả năng nhận biết giữa người và máy: khi hình ảnh bị che khuất một phần, con người có thể liên tưởng nhờ kinh nghiệm thực tế, trong khi hệ thống AI có thể gặp khó khăn hoặc đưa ra kết quả chưa chính xác."
    },

    # ==================== MẠCH D (6 YCCĐ) ====================
    {
        "stt": 24, "code": "2.D1.1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Nhận diện và hình thành giải pháp – Máy thông minh giúp giải quyết vấn đề quanh em",
        "yccd_vn": "Nêu được một số vấn đề đơn giản, gần gũi trong đời sống có thể áp dụng AI để giải quyết (như một công việc vất vả trong lớp hoặc ở nhà mà em mong có máy thông minh giúp đỡ: lau bảng trên cao, nhặt rác dưới gầm bàn).",
        "intl_full": "• Stanford AI Quests / Experience AI Module 2: Xác định vấn đề thực tế xuất phát từ nhu cầu con người (Problem Formulation).\n• UNESCO (Human-Centred Problem Identification): Đặt khó khăn của con người làm trọng tâm để tìm kiếm giải pháp công nghệ.",
        "intl_compact": "Bước đầu nhận diện các tình huống khó khăn, vất vả gần gũi quanh em (như lau bảng trên cao, dọn vệ sinh) mà máy thông minh có thể hỗ trợ con người giải quyết."
    },
    {
        "stt": 25, "code": "2.D1.2",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Nhận diện và hình thành giải pháp – Ý tưởng máy thông minh",
        "yccd_vn": "Nêu được một số ý tưởng máy thông minh (như vẽ tranh hoặc kể tên một chú robot mơ ước do em tưởng tượng ra để giúp đỡ mọi người).",
        "intl_full": "• MIT Day of AI (Creative AI Design): Phát triển năng lực sáng tạo và hình dung giải pháp hệ thống phục vụ cộng đồng.\n• Singapore MOE (Design Thinking for Primary): Thể hiện ý tưởng thiết kế công nghệ qua tranh vẽ hoặc kể chuyện.",
        "intl_compact": "Phát triển tư duy sáng tạo và thiết kế giải pháp: thể hiện ý tưởng về một cỗ máy thông minh phục vụ cộng đồng thông qua ngôn ngữ kể chuyện hoặc tranh vẽ trực quan."
    },
    {
        "stt": 26, "code": "2.D1.MR1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Nhận diện và hình thành giải pháp – Ý tưởng máy thông minh",
        "yccd_vn": "Nêu một số ví dụ phù hợp để “dạy” AI trong một tình huống cụ thể (như chọn được một số hình ảnh mẫu đúng để dạy cho chú robot của em nhận biết quả táo hay quả cam).",
        "intl_full": "• Google DeepMind (Experience AI - Dataset Curation): Thực hành vai trò lựa chọn ví dụ mẫu để dạy cho mô hình.\n• AI4K12 (Big Idea 3): Nhận biết tầm quan trọng của việc chọn đúng hình ảnh mẫu phù hợp.",
        "intl_compact": "Thực hành chọn một số hình ảnh mẫu phù hợp trong tình huống cụ thể để chuẩn bị làm ví dụ giúp máy thông minh nhận biết đồ vật."
    },
    {
        "stt": 27, "code": "2.D2.1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Vai trò của dữ liệu",
        "yccd_vn": "Giải thích được ở mức độ cơ bản về vai trò của dữ liệu trong việc “dạy” AI trong một tình huống cụ thể: đối với các hệ thống AI học máy, máy tính cần có tập ví dụ mẫu (như tranh ảnh, âm thanh) do con người chuẩn bị thì mới có thể học và nhận biết được.",
        "intl_full": "• AI4K12 (Big Idea 3 - Data Dependency in Machine Learning): Hiểu nguyên lý phụ thuộc dữ liệu của học máy: không có ví dụ mẫu thì mô hình học máy không thể học được.\n• UNESCO (System Architecture): Vai trò quan trọng của dữ liệu mẫu đầu vào trong hệ thống AI học máy.",
        "intl_compact": "Giải thích được vai trò cơ bản của dữ liệu trong học máy: hệ thống AI học máy cần có các tập ví dụ mẫu do con người chuẩn bị để học cách nhận biết và thực hiện nhiệm vụ."
    },
    {
        "stt": 28, "code": "2.D2.2",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Vai trò của dữ liệu",
        "yccd_vn": "Giải thích được ở mức độ cơ bản rằng cần cung cấp cho AI dữ liệu chính xác, rõ ràng để AI đưa ra kết quả đúng; nếu đưa cho AI hình ảnh bị mờ, chụp tối hoặc sai thông tin thì AI có thể nhận diện kém hoặc đưa ra kết quả sai.",
        "intl_full": "• Google DeepMind (Experience AI - Garbage in, Garbage out): Hiểu nguyên lý chất lượng dữ liệu: dữ liệu đầu vào mờ/nhiễu/sai thông tin có thể làm mô hình dự đoán kém hoặc sai sót.\n• OECD (Data Quality Principles): Tầm quan trọng của hình ảnh, âm thanh rõ ràng đối với kết quả của AI.",
        "intl_compact": "Giải thích được tầm quan trọng của chất lượng dữ liệu: dữ liệu đầu vào bị mờ, nhiễu hoặc sai thông tin có thể khiến mô hình học máy nhận diện kém hoặc đưa ra kết quả sai lệch."
    },
    {
        "stt": 29, "code": "2.D2.MR1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Vai trò của dữ liệu",
        "yccd_vn": "Thực hành thu thập một số dữ liệu để “dạy” AI trong một tình huống cụ thể (như tìm và sắp xếp 5-10 bức tranh rõ nét từ sách báo cũ để tạo thành bộ ví dụ mẫu dạy máy nhận diện đồ vật).",
        "intl_full": "• MIT Day of AI & Google DeepMind (Hands-on Dataset Building): Trải nghiệm quy trình thu thập và sắp xếp tập hình ảnh mẫu rõ ràng trong tình huống thực tế.\n• Singapore MOE (Computational Practice): Rèn luyện thao tác sắp xếp dữ liệu ngăn nắp, khoa học.",
        "intl_compact": "Thực hành thu thập và sắp xếp tập dữ liệu mẫu: tự tay tìm kiếm, sàng lọc và sắp xếp tập hình ảnh rõ nét để làm bài học mẫu dạy máy nhận diện đồ vật."
    }
]

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def generate_markdown_day_du(output_path):
    lines = [
        "# KHUNG CHƯƠNG TRÌNH GIÁO DỤC TRÍ TUỆ NHÂN TẠO – LỚP 2 (BẢN ĐẦY ĐỦ)",
        "",
        "> **Căn cứ pháp lý:** Quyết định số 2422/QĐ-BGDĐT ngày 18/8/2026 của Bộ Giáo dục và Đào tạo.  ",
        "> **Căn cứ tham chiếu quốc tế:** UNESCO AI Competency Framework, OECD Learning Compass 2030, Singapore MOE, AI4K12, Google DeepMind (Experience AI).",
        "",
        "---",
        "",
        "| STT | Mã YCCĐ | Mạch nội dung & Chủ đề con (QĐ 2422) | Yêu cầu cần đạt Chuẩn Quốc gia (Bộ GD&ĐT - QĐ 2422) | Chuẩn đầu ra tham chiếu Đa nguồn Quốc tế (UNESCO, OECD, Singapore MOE, AI4K12, Google) |",
        "| :---: | :---: | :--- | :--- | :--- |"
    ]
    for it in data_lop2:
        intl_fmt = it["intl_full"].replace("\n", "<br>")
        row = f"| **{it['stt']}** | **{it['code']}** | **{it['strand']}**<br>• *{it['subtopic']}* | {it['yccd_vn']} | {intl_fmt} |"
        lines.append(row)
    
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"Exported MD full: {output_path}")

def generate_markdown_rut_gon(output_path):
    lines = [
        "# KHUNG NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO – LỚP 2",
        "*(Bản rút gọn chuẩn hóa: Mạch năng lực, Chủ đề con và Yêu cầu cần đạt đầu ra)*",
        "",
        "> **Căn cứ pháp lý:** Quyết định số 2422/QĐ-BGDĐT của Bộ Giáo dục và Đào tạo.  ",
        "> **Đặc điểm:** Tinh gọn các chú giải nguồn, diễn đạt thành các Chuẩn đầu ra năng lực (Exit Competencies) liền mạch, súc tích, phục vụ trực tiếp việc xây dựng kế hoạch dạy học và giảng dạy thực tế.",
        "",
        "---",
        "",
        "| STT | Mạch năng lực | Chủ đề con | YCCĐ (Chuẩn đầu ra năng lực) |",
        "| :---: | :--- | :--- | :--- |"
    ]
    for it in data_lop2:
        row = f"| **{it['stt']}** | **{it['strand']}** | {it['subtopic']} | {it['intl_compact']} |"
        lines.append(row)
    
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"Exported MD compact: {output_path}")

def generate_docx_day_du(output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BẢNG ĐỐI SÁNH CHUẨN ĐẦU RA AI LỚP 2 (QĐ 2422/QĐ-BGDĐT – ĐA NGUỒN QUỐC TẾ)")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 32, 96)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Bản đầy đủ: Đối sánh nguyên văn chuẩn Bộ GD&ĐT và Chuẩn đầu ra tham chiếu Quốc tế (UNESCO, OECD, Singapore MOE, AI4K12, Google DeepMind)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)
    p_sub.paragraph_format.space_after = Pt(10)

    headers = [
        ("STT", Inches(0.6)),
        ("Mã YCCĐ", Inches(0.9)),
        ("Mạch nội dung & Chủ đề con", Inches(2.3)),
        ("Yêu cầu cần đạt Chuẩn Quốc gia\n(Bộ GD&ĐT - QĐ 2422)", Inches(3.3)),
        ("Chuẩn đầu ra tham chiếu Đa nguồn Quốc tế\n(UNESCO, OECD, Singapore MOE, AI4K12, Google)", Inches(3.4))
    ]
    
    table = doc.add_table(rows=len(data_lop2) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    hdr_cells = table.rows[0].cells
    for idx, (h_text, w) in enumerate(headers):
        hdr_cells[idx].width = w
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=100, right=100)
        set_cell_shading(hdr_cells[idx], "1B365D")
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))

    for row_idx, it in enumerate(data_lop2, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F7F9FC" if row_idx % 2 == 0 else "FFFFFF"

        for idx, (_, w) in enumerate(headers):
            row_cells[idx].width = w
            set_cell_margins(row_cells[idx], top=80, bottom=80, left=90, right=90)
            set_cell_shading(row_cells[idx], bg_color)

        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(str(it["stt"]))
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9.5)
        r0.font.bold = True

        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(it["code"])
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(0, 32, 96)

        p2 = row_cells[2].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        p2.paragraph_format.space_after = Pt(2)
        r2_a = p2.add_run(it["strand"] + "\n")
        r2_a.font.name = "Times New Roman"
        r2_a.font.size = Pt(9)
        r2_a.font.bold = True
        r2_b = p2.add_run("• " + it["subtopic"])
        r2_b.font.name = "Times New Roman"
        r2_b.font.size = Pt(8.5)
        r2_b.font.italic = True
        r2_b.font.color.rgb = RGBColor(60, 60, 60)

        p3 = row_cells[3].paragraphs[0]
        p3.paragraph_format.line_spacing = 1.15
        r3 = p3.add_run(it["yccd_vn"])
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(9)

        p4 = row_cells[4].paragraphs[0]
        p4.paragraph_format.line_spacing = 1.15
        lines_intl = it["intl_full"].split("\n")
        for l_idx, line_txt in enumerate(lines_intl):
            if l_idx > 0:
                p4 = row_cells[4].add_paragraph()
                p4.paragraph_format.line_spacing = 1.15
                p4.paragraph_format.space_after = Pt(2)
            parts = line_txt.split(":", 1)
            if len(parts) == 2:
                r_bold = p4.add_run(parts[0] + ":")
                r_bold.font.name = "Times New Roman"
                r_bold.font.size = Pt(8.5)
                r_bold.font.bold = True
                r_bold.font.color.rgb = RGBColor(0, 51, 102)
                r_norm = p4.add_run(parts[1])
                r_norm.font.name = "Times New Roman"
                r_norm.font.size = Pt(8.5)
            else:
                r_all = p4.add_run(line_txt)
                r_all.font.name = "Times New Roman"
                r_all.font.size = Pt(8.5)

    doc.save(output_path)
    print(f"Exported DOCX full: {output_path}")

def generate_docx_rut_gon(output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("KHUNG NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO – LỚP 2")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 32, 96)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("(Bản rút gọn chuẩn hóa: Mạch năng lực, Chủ đề con và Yêu cầu cần đạt đầu ra)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)
    p_sub.paragraph_format.space_after = Pt(12)

    headers = [
        ("STT", Inches(0.6)),
        ("Mạch năng lực", Inches(1.8)),
        ("Chủ đề con", Inches(2.0)),
        ("YCCĐ", Inches(2.5))
    ]
    
    table = doc.add_table(rows=len(data_lop2) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    hdr_cells = table.rows[0].cells
    for idx, (h_text, w) in enumerate(headers):
        hdr_cells[idx].width = w
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=100, right=100)
        set_cell_shading(hdr_cells[idx], "1B365D")
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))

    for row_idx, it in enumerate(data_lop2, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F7F9FC" if row_idx % 2 == 0 else "FFFFFF"

        for idx, (_, w) in enumerate(headers):
            row_cells[idx].width = w
            set_cell_margins(row_cells[idx], top=80, bottom=80, left=90, right=90)
            set_cell_shading(row_cells[idx], bg_color)

        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(str(it["stt"]))
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9.5)
        r0.font.bold = True

        p1 = row_cells[1].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(it["strand"])
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(9.5)
        r1.font.bold = True

        p2 = row_cells[2].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(it["subtopic"])
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(9.5)

        p3 = row_cells[3].paragraphs[0]
        p3.paragraph_format.line_spacing = 1.15
        r3 = p3.add_run(it["intl_compact"])
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(9.5)

    doc.save(output_path)
    print(f"Exported DOCX compact: {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\Users\USER\Desktop\AI-vibe-code\KhungNLĐTAI\03_San_Pham_Chinh_Thuc"
    
    md_day_du = os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop2_Day_Du.md")
    md_rut_gon = os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop2_Rut_Gon.md")
    docx_day_du = os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop2_Day_Du.docx")
    docx_rut_gon = os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop2_Rut_Gon.docx")
    
    generate_markdown_day_du(md_day_du)
    generate_markdown_rut_gon(md_rut_gon)
    generate_docx_day_du(docx_day_du)
    generate_docx_rut_gon(docx_rut_gon)
    print("All Grade 2 files generated successfully!")
