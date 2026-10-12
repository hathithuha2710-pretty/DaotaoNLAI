"""Project Word tables -> educational fields only. Raw documents remain local."""
from pathlib import Path
from docx import Document
import argparse,json,re,unicodedata
SITE=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser(description='Trích dữ liệu giáo dục từ ba tài liệu Word.')
parser.add_argument('--input-dir',type=Path,required=True)
parser.add_argument('--output',type=Path,default=SITE/'data/curriculum.json')
args=parser.parse_args()
INPUT=args.input_dir.resolve()
NAMES=['SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.docx','SP2_Bang_Ma_Tran_Ky_Nang_AI_Tieu_Hoc_Lop1_5_Moi.docx','SP3_ Phan_phoi_Chuong_Trinh_AI_Lop_1_5_36_Bai.docx']
paths=[INPUT/n for n in NAMES]
assert all(p.is_file() for p in paths)
def plain(value):
    value=unicodedata.normalize('NFC',value).replace('\u00a0',' ')
    value=re.sub(r'\*\*([^*]+)\*\*',r'\1',value)
    value=re.sub(r'(?<!\w)\*([^*\n]+)\*',r'\1',value)
    return value.strip()
def codes(value):return list(dict.fromkeys(re.findall(r'\b[1-5]\.[A-D][1-5]\.(?:MR\d+|\d+)\b',value)))
yccd=[]
for grade,table in enumerate(Document(paths[0]).tables,1):
    for row in table.rows[1:]:
        cells=[plain(c.text) for c in row.cells]
        assert len(cells)==5
        yccd.append({'grade':grade,'code':cells[1],'strand':cells[2].split('.')[0].strip(),'strandName':cells[2],'topic':cells[3],'text':cells[4]})
skilldoc=Document(paths[1])
groupNames=[plain(p.text) for p in skilldoc.paragraphs if p.text.strip().startswith('NHÓM ')]
assert len(groupNames)==6
skillGroups=[];skills=[]
for number,(table,title) in enumerate(zip(skilldoc.tables,groupNames),1):
    name=re.sub(r'^NHÓM\s*\d+\s*:\s*','',title)
    skillGroups.append({'id':number,'name':name})
    for row in table.rows[1:]:
        cells=[plain(c.text) for c in row.cells]
        assert len(cells)==8
        sid=int(cells[0]);name,*definition=cells[1].splitlines()
        progression=[]
        for grade,value in enumerate(cells[3:],1):
            lines=[s.strip() for s in value.splitlines() if s.strip()]
            level=next((s for s in lines if s.startswith('Mức ')),'')
            outcomes=list(dict.fromkeys(re.findall(r'\[(VN-CORE|VN-MR|INT-ENRICH)\]',value)))
            anchor=next((s for s in lines if 'Mã YCCĐ:' in s),'')
            activities=[]
            for line in lines:
                if line.startswith(('•','–','- ')) and line!='-':activities.append(line.lstrip('•–- ').strip())
                elif activities and not line.startswith(('Mức ','[','(','Nguồn:')):activities[-1]+=' '+line
            progression.append({'grade':grade,'level':level,'labels':outcomes,'codes':codes(anchor),'activities':activities})
        skills.append({'id':sid,'group':number,'name':name,'definition':'\n'.join(definition),'strand':cells[2],'progression':progression})
lessons=[]
for grade,table in enumerate(Document(paths[2]).tables,1):
    assert len(table.rows)==37
    for position,row in enumerate(table.rows[1:],1):
        raw=[c.text for c in row.cells]
        assert len(raw)==5
        code,name,focus,materials,anchor=raw
        refs=list(dict.fromkeys(int(s) for s in re.findall(r'\bKN\s*(\d+)',focus)))
        assert all(1<=s<=46 for s in refs)
        focusClean=re.sub(r'\s*(?:-\s*)?Tích hợp\s*(?:\[KN[^\]]*\][,;\s]*)+\.?','',focus,flags=re.I)
        focusClean=re.sub(r'\s*\[KN[^\]]*\]','',focusClean)
        focusClean=re.sub(r'[ \t]+',' ',focusClean).strip(' -')
        lessons.append({'id':f'{grade}-{position:02}','grade':grade,'position':position,'code':plain(code),'name':name,'focus':plain(focusClean),'materials':plain(materials),'codes':codes(anchor),'assessment':plain(anchor) if not codes(anchor) else '', 'skillIds':refs})
assert len(yccd)==123 and len(skills)==46 and len(lessons)==180
assert {s['id'] for s in skills}==set(range(1,47))
data={'grades':[{'id':g,'name':f'Lớp {g}','count':36} for g in range(1,6)],'strands':[{'id':'A','name':'Tư duy lấy con người làm trung tâm'},{'id':'B','name':'Đạo đức AI'},{'id':'C','name':'Các kĩ thuật và ứng dụng AI'},{'id':'D','name':'Thiết kế hệ thống AI'}],'skillGroups':skillGroups,'skills':skills,'yccd':yccd,'lessons':lessons}
args.output.parent.mkdir(parents=True,exist_ok=True)
args.output.write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
print(json.dumps({'yccd':len(yccd),'skills':len(skills),'lessons':len(lessons)},ensure_ascii=False))
