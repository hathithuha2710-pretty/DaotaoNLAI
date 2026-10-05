import os
import sys
import re
import docx

sys.stdout.reconfigure(encoding='utf-8')

# Read QĐ 2422 primary section
qd2422_path = "06_He_Thong_Ky_Thuat/extracted_data/qd2422_full.txt"
with open(qd2422_path, encoding='utf-8') as f:
    qd_text = f.read()

# Let's inspect where primary school starts and ends in qd2422_full.txt
print("Length of qd2422_full.txt:", len(qd_text))

# Let's check master files
master_full_md = "03_San_Pham_Chinh_Thuc/YCCĐ Lớp 1 - 5/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Day_Du.md"
master_comp_md = "03_San_Pham_Chinh_Thuc/YCCĐ Lớp 1 - 5/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.md"
master_full_docx = "03_San_Pham_Chinh_Thuc/YCCĐ Lớp 1 - 5/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Day_Du.docx"
master_comp_docx = "03_San_Pham_Chinh_Thuc/YCCĐ Lớp 1 - 5/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.docx"

def parse_md_table(md_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.readlines()
    rows = []
    current_grade = None
    for line in lines:
        line_s = line.strip()
        if line_s.startswith("# ") or line_s.startswith("## ") or line_s.startswith("### "):
            m = re.search(r"Lớp\s*(\d)", line_s, re.IGNORECASE)
            if m:
                current_grade = int(m.group(1))
        if line_s.startswith("| **"):
            parts = [p.strip() for p in line_s.split("|")[1:-1]]
            rows.append({
                "grade": current_grade,
                "raw": line_s,
                "parts": parts
            })
    return rows

full_rows = parse_md_table(master_full_md)
comp_rows = parse_md_table(master_comp_md)

print(f"Master Full MD total rows: {len(full_rows)}")
print(f"Master Comp MD total rows: {len(comp_rows)}")

# Count by grade
for g in range(1, 6):
    g_full = [r for r in full_rows if r['grade'] == g]
    g_comp = [r for r in comp_rows if r['grade'] == g]
    print(f"Lớp {g}: Full={len(g_full)}, Comp={len(g_comp)}")

# Check docx
doc_f = docx.Document(master_full_docx)
doc_c = docx.Document(master_comp_docx)

f_docx_rows = sum(len(t.rows)-1 for t in doc_f.tables)
c_docx_rows = sum(len(t.rows)-1 for t in doc_c.tables)
print(f"Master Full DOCX total table data rows: {f_docx_rows}")
print(f"Master Comp DOCX total table data rows: {c_docx_rows}")

# Let's inspect the columns of Full and Comp
if full_rows:
    print("\nSample Full Row columns count:", len(full_rows[0]['parts']))
    for i, p in enumerate(full_rows[0]['parts']):
        print(f" Col {i}: {p[:60]}...")

if comp_rows:
    print("\nSample Comp Row columns count:", len(comp_rows[0]['parts']))
    for i, p in enumerate(comp_rows[0]['parts']):
        print(f" Col {i}: {p[:60]}...")
