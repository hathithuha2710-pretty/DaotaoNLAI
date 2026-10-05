import os
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

docx_files = [
    "03_San_Pham_Chinh_Thuc/YCCĐ Lớp 1 - 5/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Day_Du.docx",
    "03_San_Pham_Chinh_Thuc/YCCĐ Lớp 1 - 5/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.docx",
]
for g in range(1, 6):
    docx_files.append(f"03_San_Pham_Chinh_Thuc/file lẻ/SP1_Khung_YCCD_AI_Lop{g}_Day_Du.docx")
    docx_files.append(f"03_San_Pham_Chinh_Thuc/file lẻ/SP1_Khung_YCCD_AI_Lop{g}_Rut_Gon.docx")

print("=== CHECKING DOCX FILES INTEGRITY ===")
for p in docx_files:
    if not os.path.exists(p):
        print(f"❌ Missing: {p}")
        continue
    doc = docx.Document(p)
    tables_count = len(doc.tables)
    rows_count = sum(len(t.rows)-1 for t in doc.tables)
    paragraphs_count = len(doc.paragraphs)
    size_kb = os.path.getsize(p) / 1024
    print(f"✅ {os.path.basename(p):<45} | Size: {size_kb:6.1f} KB | Tables: {tables_count} | Data Rows: {rows_count:3d} | Paras: {paragraphs_count:3d}")
