# -*- coding: utf-8 -*-
"""
Script cập nhật file Word chứa 2 bảng Phân phối chương trình Giáo dục AI Lớp 1 và Lớp 6
Được hiệu chỉnh theo yêu cầu:
1. Cột Tên bài học chỉ có tên bài, bỏ tiền tố 'Bài 1:', 'Bài 2:',...
2. Xóa bỏ cột 'Phân loại', bài nào là cốt lõi thì toàn bộ khung/dòng tô màu vàng nổi bật.
3. Giữ nguyên 100% nội dung, đối chiếu chính xác từng bài cốt lõi và mở rộng theo QĐ 2422.
"""

import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="B0BEC5", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = parse_xml(f'<w:cantSplit {nsdecls("w")}/>')
    trPr.append(cantSplit)

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = parse_xml(f'<w:tblHeader {nsdecls("w")}/>')
    trPr.append(tblHeader)

def clean_lesson_name(name):
    # Loại bỏ tiền tố 'Bài 1:', 'Bài 2:', 'Bài 12: '...
    return re.sub(r'^Bài\s*\d+\s*:\s*', '', name.strip())

# Dữ liệu 36 tiết Lớp 1 (đã đối chiếu chính xác phân loại)
DATA_LOP_1_RAW = [
    ("1", "Bài 1: Cảm xúc của em và cảm xúc của bạn bè",
     "Nêu ví dụ về cảm xúc con người (vui, buồn, sợ hãi...); nhận ra con người có cảm xúc thật còn AI thì không (chỉ mô phỏng).",
     "1.A1.1, 1.A1.2", True), # True = Cốt lõi (Tô vàng)
    ("2", "Bài 2: Nụ cười của Robot do ai lập trình?",
     "Nhận biết biểu hiện cảm xúc của AI (cười, chào) là do con người lập trình/thiết kế trước; AI không tự vui buồn.",
     "1.A1.3, 1.A1.4", True),
    ("3", "Bài 3: Trò chơi đóng vai: Robot biết cười",
     "Trải nghiệm đóng vai Robot phản hồi nụ cười và lời chào theo kịch bản quy ước trước.",
     "1.A1.MR", False),
    ("4", "Bài 4: Bàn tay con người thiết kế Robot",
     "Khám phá công việc của các kỹ sư sáng tạo ra các cử chỉ thân thiện cho robot.",
     "1.A1.MR", False),
    ("5", "Bài 5: Cách AI thể hiện cảm xúc thân thiện",
     "Nêu được các cách mà AI thể hiện cảm xúc qua hình ảnh, giọng nói hoặc hành vi (biểu tượng khuôn mặt vui, giọng nói thân thiện, lời chào).",
     "1.A2.1", True),
    ("6", "Bài 6: Những thiết bị AI giúp ích quanh em",
     "Kể tên thiết bị có sử dụng AI (loa thông minh, điện thoại); mô tả công dụng chính giúp cuộc sống tiện ích hơn.",
     "1.A2.2, 1.A2.3", True),
    ("7", "Bài 7: Lời chào dễ thương giúp hợp tác vui vẻ",
     "Nêu được rằng những biểu hiện thân thiện giúp AI giao tiếp tự nhiên hơn, khiến người dùng cảm thấy thoải mái, dễ hợp tác.",
     "1.A2.MR1", False),
    ("8", "Bài 8: Đôi mắt và đôi tai của thiết bị thông minh",
     "Nhận biết AI trong một số ví dụ cụ thể; nhận ra thiết bị thông minh có các bộ phận như con người: Camera là 'mắt', Micro là 'tai'.",
     "1.C1.1, 1.C1.3", True),
    ("9", "Bài 9: Khám phá 'Mắt thần' Camera nhận diện đồ vật",
     "Nhận biết được AI có khả năng xử lí hình ảnh nhận được để nhận diện các đồ vật xung quanh em.",
     "1.C1.MR1", False),
    ("10", "Bài 10: Trò chuyện cùng trợ lý thông minh",
     "Nhận diện trợ lý ảo quen thuộc; nhận biết AI hiểu mệnh lệnh đơn giản và trò chuyện theo kịch bản định sẵn dưới sự giám sát.",
     "1.C1.2, 1.C1.4", True),
    ("11", "Bài 11: Đi tìm thiết bị AI trong trường học của em",
     "Quan sát và kể tên các thiết bị thông minh có trong trường học (camera cổng trường, màn hình tương tác).",
     "1.A2.MR", False),
    ("12", "Bài 12: Máy thông minh làm việc tốt",
     "Nêu ví dụ con người sử dụng AI đúng cách vì mục đích tốt đẹp (giúp học tập, chỉ đường, hỗ trợ bác sĩ).",
     "1.B3.2", True),
    ("13", "Bài 13: Vẽ tranh: Người bạn AI làm việc tốt",
     "Sáng tạo tranh vẽ ý tưởng thiết bị thông minh giúp đỡ người khuyết tật, người già.",
     "1.B3.MR", False),
    ("14", "Bài 14: Không dùng thiết bị thông minh làm việc xấu",
     "Nhận biết hành vi dùng AI gây hại cho người khác; hiểu rằng tuyệt đối không được phép sử dụng AI với mục đích làm hại người khác.",
     "1.B1.1, 1.B3.1", True),
    ("15", "Bài 15: Thói quen tốt khi sử dụng thiết bị thông minh",
     "Hình thành thói quen ngồi đúng tư thế, giữ khoảng cách mắt an toàn và nhờ người lớn hướng dẫn khi dùng máy.",
     "1.B2.MR", False),
    ("16", "Bài 16: Kể chuyện: Bạn nhỏ và chú Robot lễ phép",
     "Đọc truyện và thảo luận về cách chào hỏi văn minh, đúng mực khi tương tác với thiết bị số.",
     "1.B3.MR", False),
    ("17", "Bài 17: Ôn tập chuyên đề Học kỳ I: Cảm xúc và Đạo đức",
     "Hệ thống hóa các nhận thức về cảm xúc người vs máy, giác quan của máy và hành vi tốt khi dùng AI.",
     "Ôn tập MR", False),
    ("18", "Bài 18: Báo cáo hồ sơ học tập Học kỳ I",
     "Tổng hợp minh chứng quá trình: Trưng bày bộ tranh vẽ cảm xúc và bảng nhận diện thiết bị AI; giáo viên nhận xét sự tiến bộ.",
     "Mục VI QĐ 2422", False),
    ("19", "Bài 19: Khám phá 'Tai thính' Micro phân biệt âm thanh",
     "Nhận biết được AI có khả năng xử lí âm thanh để phân biệt các loại âm thanh khác nhau trong đời sống.",
     "1.C1.MR2", False),
    ("20", "Bài 20: Trò chơi: Máy lắng nghe tiếng còi và tiếng nhạc",
     "Trải nghiệm phân biệt âm thanh cảnh báo và âm thanh giai điệu qua giác quan Micro của máy.",
     "1.C1.MR", False),
    ("21", "Bài 21: Trải nghiệm nét vẽ kỳ diệu cùng AI",
     "Làm quen với việc máy gợi ý nét vẽ hoàn chỉnh từ hình vẽ phác thảo đơn giản của học sinh.",
     "1.C2.MR", False),
    ("22", "Bài 22: Máy thông minh học từ ví dụ",
     "Nêu ví dụ AI 'học' từ hình ảnh hoặc thông tin con người cung cấp (AI học nhận biết quả cam qua nhiều bức ảnh quả cam).",
     "1.D1.1", True),
    ("23", "Bài 23: Cần nhiều ví dụ khác nhau để máy học đúng",
     "Biết được rằng để AI trả lời đúng cần có nhiều ví dụ đúng và khác nhau để AI học.",
     "1.D1.MR1", False),
    ("24", "Bài 24: Trò chơi phân loại: Dạy máy nhận biết quả",
     "Học sinh cung cấp các thẻ hình ảnh mẫu để máy phân loại quả tròn và quả dài.",
     "1.D1.MR", False),
    ("25", "Bài 25: Nhận biết máy thông minh đơn năng",
     "Nhận biết được rằng có loại máy thông minh chỉ làm được một việc duy nhất (như máy nhận diện khuôn mặt).",
     "1.D2.1", True),
    ("26", "Bài 26: Xếp nhóm máy làm một việc và nhiều việc",
     "Xếp được một số máy thông minh vào hai nhóm: nhóm chỉ làm được một việc và nhóm làm được nhiều việc khác nhau.",
     "1.D2.2", True),
    ("27", "Bài 27: Thực hành phân loại máy đa năng trong đời sống",
     "Khám phá điện thoại thông minh, máy tính bảng làm được nhiều việc khác nhau.",
     "1.D2.MR", False),
    ("28", "Bài 28: Trò chơi: Đi tìm người bạn máy đơn năng",
     "Phân biệt các máy quét mã vạch ở siêu thị, máy đo nhiệt độ chỉ làm một nhiệm vụ.",
     "1.D2.MR", False),
    ("29", "Bài 29: Khởi động hoạt động: Chú Robot giấy của em",
     "Lên ý tưởng thiết kế một mô hình chú Robot giấy thông minh đồng hành trong học tập.",
     "1.D2.MR", False),
    ("30", "Bài 30: Chế tạo mô hình: Mắt Camera và Tai Micro",
     "Cắt dán các bộ phận cảm nhận cho robot giấy: Đôi mắt Camera và đôi tai Micro.",
     "1.D2.MR", False),
    ("31", "Bài 31: Vẽ tranh các đồ vật để nạp vào túi nhớ của Robot",
     "Vẽ và cắt các thẻ tranh đồ chơi, hoa quả để nạp vào 'bộ nhớ ví dụ' của chú Robot giấy.",
     "1.D1.MR", False),
    ("32", "Bài 32: Đóng vai ra lệnh và đối thoại cùng Robot giấy",
     "Thực hành nói câu lệnh rõ ràng và cho bạn robot giấy phản hồi lời chào thân thiện.",
     "1.A1.MR", False),
    ("33", "Bài 33: Tổng hợp minh chứng & Báo cáo hồ sơ học tập Học kỳ II",
     "Đánh giá quá trình qua hồ sơ bài tập và mô hình học tập; giáo viên nhận xét sự tiến bộ của học sinh.",
     "Mục VI QĐ 2422", False),
    ("34", "Bài 34: Triển lãm ngày hội: Chú Robot giấy thông minh",
     "Trưng bày gian hàng mô hình robot giấy và bộ sưu tập ví dụ học tập của cả lớp.",
     "Triển lãm MR", False),
    ("35", "Bài 35: Tổng kết chương trình AI Lớp 1: Em yêu công nghệ",
     "Hệ thống hóa toàn bộ 4 mạch nội dung Lớp 1; tuyên dương tinh thần tự tin và làm chủ thiết bị.",
     "Tổng kết MR", False),
    ("36", "Bài 36: Báo cáo chuyên đề cốt lõi: Em làm chủ thiết bị thông minh",
     "Báo cáo tranh vẽ người bạn AI tương lai; tổng kết hồ sơ minh chứng cốt lõi; khẳng định con người làm chủ AI vì việc tốt.",
     "1.A1.2, 1.B3.2, 1.D2.2", True)
]

# Dữ liệu 36 tiết Lớp 6 (đã đối chiếu chính xác phân loại)
DATA_LOP_6_RAW = [
    ("1", "Bài 1: AI là công cụ, con người ra quyết định",
     "Giải thích AI do con người tạo ra, lập trình và điều khiển để thực hiện nhiệm vụ cụ thể; AI chỉ là công cụ hỗ trợ, con người ra quyết định cuối cùng và chịu trách nhiệm.",
     "6.A1.1, 6.A1.2", True), # True = Cốt lõi (Tô vàng)
    ("2", "Bài 2: Kỹ năng kiểm chứng kết quả do AI cung cấp",
     "Thực hiện kiểm tra lại kết quả do AI đưa ra (đối chiếu SGK, nguồn tin cậy); nêu ví dụ tình huống con người ra quyết định với sự hỗ trợ của AI.",
     "6.A1.3, 6.A3.1", True),
    ("3", "Bài 3: Thực hành kỹ năng đối chiếu đa nguồn tin cậy",
     "Rèn luyện thói quen tra cứu chéo ít nhất 2 nguồn tin độc lập để xác minh câu trả lời của AI.",
     "6.A1.MR", False),
    ("4", "Bài 4: Nhận biết hiện tượng ảo giác AI (hallucination)",
     "Nhận diện các trường hợp AI đưa ra câu trả lời có vẻ hợp lý nhưng hoàn toàn sai sự thật.",
     "6.A1.MR", False),
    ("5", "Bài 5: Bảo vệ dữ liệu cá nhân và quyền riêng tư",
     "Giải thích dữ liệu cá nhân (ảnh, giọng nói, họ tên, vị trí) cần bảo vệ; trình bày khái niệm quyền riêng tư ở mức đơn giản và tác hại khi dữ liệu cá nhân bị sử dụng sai mục đích.",
     "6.A3.3, 6.A3.4", True),
    ("6", "Bài 6: Kỹ năng khử nhận dạng dữ liệu cá nhân",
     "Thực hành kỹ thuật che mặt, ẩn thông tin danh tính trước khi chia sẻ dữ liệu phục vụ học tập.",
     "6.A3.MR", False),
    ("7", "Bài 7: Tính hai mặt của các tính năng AI",
     "Chỉ ra mặt tích cực và hạn chế của một số tính năng AI cụ thể (thu thập dữ liệu giúp gợi ý chính xác nhưng tiềm ẩn nguy cơ xâm phạm đời tư).",
     "6.B1.1", True),
    ("8", "Bài 8: Tiêu chuẩn an toàn và minh bạch của ứng dụng AI",
     "Đặt được một số câu hỏi đơn giản để kiểm tra tính an toàn và minh bạch của ứng dụng AI ('Công cụ này có an toàn không?', 'Dữ liệu lưu ở đâu?').",
     "6.B2.1", True),
    ("9", "Bài 9: Tổng hợp minh chứng & Báo cáo học tập giữa Học kỳ I",
     "Tổng hợp minh chứng quá trình: Đánh giá phiếu kiểm chứng thông tin và bảng kiểm an toàn dữ liệu; giáo viên nhận xét sự tiến bộ.",
     "Mục VI QĐ 2422", False),
    ("10", "Bài 10: Khám phá cách AI học: Dữ liệu và Thuật toán",
     "Giải thích hai thành phần chính để huấn luyện AI là dữ liệu (văn bản, hình ảnh) và thuật toán (phương pháp học từ dữ liệu).",
     "6.C1.1", True),
    ("11", "Bài 11: Mối liên hệ giữa huấn luyện và hoạt động thực tế",
     "Chỉ ra được mối liên hệ giữa quá trình huấn luyện và hoạt động của một công cụ AI qua ví dụ đơn giản (trợ lý ảo hiểu tiếng Việt nhờ dữ liệu giọng nói tiếng Việt).",
     "6.C1.MR1", False),
    ("12", "Bài 12: Thử nghiệm đối kháng với công cụ AI",
     "Thực hiện được thử nghiệm đơn giản với một công cụ AI và nhận xét được các bước hoạt động của công cụ đó (đặt cùng câu lệnh cho các công cụ khác nhau).",
     "6.C1.MR2", False),
    ("13", "Bài 13: Kỹ thuật câu lệnh (prompt) có cấu trúc",
     "Thực hành đặt câu lệnh rõ ràng, cung cấp bối cảnh và yêu cầu cụ thể để nhận được kết quả tốt nhất.",
     "6.C1.MR", False),
    ("14", "Bài 14: Kỹ năng ghi nhật ký câu lệnh (prompt log)",
     "Thực hành ghi lại nguồn dữ liệu, câu lệnh đã cung cấp và phản hồi nhận được để thực hiện trách nhiệm giải trình.",
     "6.C1.MR", False),
    ("15", "Bài 15: Tác động của AI đối với bản thân và học tập",
     "Nêu ví dụ tác động tích cực và tiêu cực của AI đối với bản thân, gia đình; trình bày lợi ích của AI trong việc hỗ trợ con người học hỏi, rèn kĩ năng khi dùng đúng mực.",
     "6.C1.2, 6.A3.2", True),
    ("16", "Bài 16: Phân tích tác động của công cụ AI trong gia đình",
     "Phân tích được ở mức đơn giản tác động của một công cụ AI cụ thể đối với bản thân và gia đình (chỉ ra cả mặt tích cực và tiêu cực).",
     "6.C1.MR3", False),
    ("17", "Bài 17: Ôn tập chuyên đề Học kỳ I: Dữ liệu và Câu lệnh",
     "Hệ thống hóa các kiến thức về kiến trúc Dữ liệu - Thuật toán, kỹ thuật câu lệnh và liêm chính học thuật.",
     "Ôn tập MR", False),
    ("18", "Bài 18: Báo cáo hồ sơ học tập Học kỳ I",
     "Trưng bày sổ nhật ký câu lệnh (prompt log) và bảng so sánh công cụ AI; giáo viên nhận xét đánh giá quá trình.",
     "Mục VI QĐ 2422", False),
    ("19", "Bài 19: Phân biệt ứng dụng có AI và không có AI",
     "Phân biệt công cụ có AI và không có AI qua ví dụ gần gũi; kể tên và mô tả chức năng chính của một số công cụ AI thông dụng (trợ lý ảo, bản đồ, dịch).",
     "6.C2.1, 6.C2.2", True),
    ("20", "Bài 20: Ứng dụng AI gắn với thực tiễn tại Việt Nam",
     "Nêu được ví dụ về ứng dụng AI trong cuộc sống, ưu tiên các vấn đề gắn với thực tiễn ở Việt Nam (nông nghiệp thông minh, định danh số, dự báo thời tiết).",
     "6.C2.MR1", False),
    ("21", "Bài 21: Các công nghệ AI quen thuộc trong đời sống",
     "Kể được tên một số công nghệ AI quen thuộc trong đời sống (nhận dạng hình ảnh, chuyển đổi giữa văn bản và giọng nói).",
     "6.C3.1", True),
    ("22", "Bài 22: Tính năng AI tích hợp trong ứng dụng thường gặp",
     "Trình bày được một số tính năng AI tích hợp trong các ứng dụng thường gặp (tính năng gợi ý đề xuất nội dung, dịch tự động).",
     "6.C3.MR1", False),
    ("23", "Bài 23: Nhận biết hiện tượng 'bong bóng thông tin'",
     "Tìm hiểu nguyên nhân khi hệ thống chỉ gợi ý những gì người dùng thích, dẫn đến góc nhìn phiến diện.",
     "6.C3.MR", False),
    ("24", "Bài 24: Ranh giới sử dụng: Khi nào nên và không nên dùng AI?",
     "Nêu được một số tình huống quen thuộc nên hoặc không nên sử dụng AI (nên dùng luyện phát âm, tìm ý tưởng; không nên dùng khi làm bài kiểm tra).",
     "6.D1.1", True),
    ("25", "Bài 25: Tranh biện ý kiến cá nhân: Nên hay không nên dùng AI?",
     "Trình bày được ý kiến cá nhân về việc nên hay không nên sử dụng AI trong một số tình huống thực tế; giải thích lí do bảo vệ quan điểm.",
     "6.D1.MR1", False),
    ("26", "Bài 26: Xây dựng quy ước sử dụng AI trong học tập của lớp",
     "Thảo luận nhóm xây dựng bản cam kết: Sử dụng AI để gợi ý dàn ý, không sao chép nguyên văn.",
     "6.D1.MR", False),
    ("27", "Bài 27: Tổng hợp minh chứng & Báo cáo học tập giữa Học kỳ II",
     "Đánh giá quá trình: Đánh giá bài tranh biện và bản quy ước sử dụng AI; giáo viên nhận xét sự tiến bộ.",
     "Mục VI QĐ 2422", False),
    ("28", "Bài 28: Giới hạn của hệ thống AI so với con người",
     "Trình bày một số giới hạn của hệ thống AI so với con người qua các tình huống, công việc mà con người thực hiện phù hợp hơn AI (thấu cảm, đạo đức).",
     "6.D2.1", True),
    ("29", "Bài 29: Vì sao AI gặp giới hạn trong một số công việc?",
     "Giải thích được vì sao AI gặp giới hạn trong một số công việc so với con người (AI học từ dữ liệu quá khứ nên khó thích ứng tình huống mới lạ).",
     "6.D2.MR1", False),
    ("30", "Bài 30: Đề xuất mô hình kết hợp tối ưu giữa Người và AI",
     "Đề xuất được cách kết hợp giữa con người và AI trong một công việc cụ thể để phát huy thế mạnh của mỗi bên (AI gợi ý thông tin, người phán đoán).",
     "6.D2.MR2", False),
    ("31", "Bài 31: Xây dựng kịch bản cộng tác: Em và trợ lý AI học tập",
     "Thực hành thiết lập kịch bản làm việc nhóm: Phân công nhiệm vụ cho công cụ AI hỗ trợ tra cứu và nhóm học sinh tổng hợp.",
     "6.D2.MR", False),
    ("32", "Bài 32: Thử nghiệm kịch bản và đánh giá hiệu quả",
     "Vận hành kịch bản cộng tác đã thiết kế và nhận xét ưu điểm, hạn chế của việc phối hợp cùng AI.",
     "6.D2.MR", False),
    ("33", "Bài 33: Hoàn thiện hồ sơ lưu vết câu lệnh và minh chứng",
     "Rà soát toàn bộ nhật ký câu lệnh và các sản phẩm đã thực hiện trong cả năm học.",
     "6.D2.MR", False),
    ("34", "Bài 34: Triển lãm báo cáo: Ý tưởng kết hợp Người và AI",
     "Trưng bày poster kịch bản cộng tác Người - AI và hồ sơ nhật ký câu lệnh của các nhóm học sinh.",
     "Triển lãm MR", False),
    ("35", "Bài 35: Tổng kết chương trình AI Lớp 6: Tự hào công dân số",
     "Hệ thống hóa toàn bộ 4 mạch năng lực; tuyên dương tinh thần sáng tạo, làm chủ và liêm chính học thuật.",
     "Tổng kết MR", False),
    ("36", "Bài 36: Báo cáo chuyên đề cốt lõi: Sử dụng AI thông minh và trách nhiệm",
     "Tổng kết chuyên đề cốt lõi 12 tiết; hoàn thiện hồ sơ học tập và lưu vết minh chứng câu lệnh (prompt log); cam kết liêm chính số.",
     "6.A1.2, 6.A1.3, 6.D1.1, 6.D2.1", True)
]

def create_document():
    doc = docx.Document()
    
    # Page setup - A4, margins 1.8cm (~0.71 inches)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.71)
        section.bottom_margin = Inches(0.71)
        section.left_margin = Inches(0.71)
        section.right_margin = Inches(0.71)
    
    # Normal style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)
    
    # Header Banner
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — CHƯƠNG TRÌNH GIÁO DỤC TRÍ TUỆ NHÂN TẠO (AI)")
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(100, 116, 139)
    p_inst.paragraph_format.space_after = Pt(2)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("KẾ HOẠCH BÀI DẠY VÀ PHÂN PHỐI CHƯƠNG TRÌNH AI")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93) # Navy
    p_title.paragraph_format.space_after = Pt(4)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Căn cứ Quyết định số 2422/QĐ-BGDĐT của Bộ trưởng Bộ Giáo dục và Đào tạo\n(Bao phủ 12 tiết Cốt lõi bắt buộc và 24 tiết Mở rộng tăng cường)")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    p_sub.paragraph_format.space_after = Pt(12)
    
    # Legend - Ghi chú trực quan
    p_legend = doc.add_paragraph()
    p_legend.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_leg_title = p_legend.add_run("Quy ước đánh dấu màu:")
    r_leg_title.bold = True
    r_leg_title.font.size = Pt(10.5)
    
    p_legend2 = doc.add_paragraph()
    r_dot_yellow = p_legend2.add_run("■ ")
    r_dot_yellow.font.color.rgb = RGBColor(234, 179, 8) # Yellow icon
    r_dot_yellow.font.size = Pt(12)
    r_txt_yellow = p_legend2.add_run("Các dòng tô màu vàng: ")
    r_txt_yellow.bold = True
    p_legend2.add_run("Là các bài thuộc ")
    r_bold_core = p_legend2.add_run("Nội dung cốt lõi bắt buộc (12 tiết/lớp/năm học)")
    r_bold_core.bold = True
    p_legend2.add_run(", thực hiện theo chuyên đề giáo dục AI bắt buộc cho 100% học sinh.")
    
    p_legend3 = doc.add_paragraph()
    r_dot_white = p_legend3.add_run("□ ")
    r_dot_white.font.color.rgb = RGBColor(100, 116, 139)
    r_dot_white.font.size = Pt(12)
    r_txt_white = p_legend3.add_run("Các dòng nền trắng: ")
    r_txt_white.bold = True
    p_legend3.add_run("Là các bài thuộc ")
    r_bold_ext = p_legend3.add_run("Nội dung mở rộng tăng cường (24 tiết/lớp/năm học)")
    r_bold_ext.bold = True
    p_legend3.add_run(", phục vụ thực hành trải nghiệm, dự án và bồi dưỡng chuyên sâu.")
    p_legend3.paragraph_format.space_after = Pt(12)
    
    def build_table(class_title, grade_desc, data_rows):
        p_sec = doc.add_paragraph()
        p_sec.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_sec = p_sec.add_run(class_title)
        r_sec.font.name = 'Times New Roman'
        r_sec.font.size = Pt(13)
        r_sec.font.bold = True
        r_sec.font.color.rgb = RGBColor(27, 54, 93)
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(2)
        
        p_desc = doc.add_paragraph()
        r_d = p_desc.add_run(grade_desc)
        r_d.font.name = 'Times New Roman'
        r_d.font.size = Pt(10)
        r_d.font.italic = True
        r_d.font.color.rgb = RGBColor(100, 116, 139)
        p_desc.paragraph_format.space_after = Pt(8)
        
        # Bảng 4 cột: Tiết | Tên bài học | Nội dung | Mã YCCĐ (QĐ 2422)
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        set_table_borders(table, color="B0BEC5", sz="4")
        
        # Độ rộng cột (Tổng cộng = 6.85 inches ~ 17.4 cm)
        # Tiết: 0.6 in (1.5 cm)
        # Tên bài học: 2.1 in (5.3 cm)
        # Nội dung: 3.1 in (7.9 cm)
        # Mã YCCĐ: 1.05 in (2.7 cm)
        col_widths = [Inches(0.6), Inches(2.1), Inches(3.1), Inches(1.05)]
        
        # Hàng tiêu đề (Header Row)
        hdr_row = table.rows[0]
        make_row_header(hdr_row)
        make_row_cant_split(hdr_row)
        hdr_titles = ["Tiết", "Tên bài học", "Nội dung", "Mã YCCĐ\n(QĐ 2422)"]
        
        for c_idx, title in enumerate(hdr_titles):
            cell = hdr_row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "1B365D") # Navy blue
            set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            
            run = p.add_run(title)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        
        # Màu vàng nhạt dịu mắt chuẩn giáo dục cho bài Cốt lõi
        COLOR_CORE_YELLOW = "FEF08A" # Soft pastel yellow
        COLOR_WHITE = "FFFFFF"
        
        # Các hàng dữ liệu
        for r_idx, (tiet, raw_ten_bai, noi_dung, ma_yccd, is_core) in enumerate(data_rows, 1):
            row = table.add_row()
            make_row_cant_split(row)
            
            clean_ten = clean_lesson_name(raw_ten_bai)
            bg_color = COLOR_CORE_YELLOW if is_core else COLOR_WHITE
            
            row_data = [tiet, clean_ten, noi_dung, ma_yccd]
            
            for c_idx, text in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_background(cell, bg_color)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                
                if c_idx == 0: # Tiết
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(text)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(9.5)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(15, 23, 42)
                elif c_idx == 1: # Tên bài học (đã bỏ 'Bài X:')
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(text)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(9.5)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(15, 23, 42)
                elif c_idx == 2: # Nội dung
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(text)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(51, 65, 85)
                elif c_idx == 3: # Mã YCCĐ
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(text)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(9)
                    run.font.bold = True
                    if is_core:
                        run.font.color.rgb = RGBColor(133, 77, 14) # Nâu hổ phách đậm trên nền vàng
                    else:
                        run.font.color.rgb = RGBColor(71, 85, 105) # Xám đậm trên nền trắng
        
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(12)
    
    # Bảng 1: Lớp 1
    build_table(
        "BẢNG 1: KẾ HOẠCH BÀI DẠY VÀ PHÂN PHỐI CHƯƠNG TRÌNH AI - LỚP 1",
        "Khung thời lượng: 36 tiết/năm học (12 tiết Cốt lõi tô màu vàng + 24 tiết Mở rộng nền trắng) | Cấp Tiểu học",
        DATA_LOP_1_RAW
    )
    
    # Ngắt trang giữa 2 lớp
    doc.add_page_break()
    
    # Bảng 2: Lớp 6
    build_table(
        "BẢNG 2: KẾ HOẠCH BÀI DẠY VÀ PHÂN PHỐI CHƯƠNG TRÌNH AI - LỚP 6",
        "Khung thời lượng: 36 tiết/năm học (12 tiết Cốt lõi tô màu vàng + 24 tiết Mở rộng nền trắng) | Cấp THCS",
        DATA_LOP_6_RAW
    )
    
    output_filename = "Bang_Phan_Phoi_AI_Lop1_Lop6_QD2422.docx"
    doc.save(output_filename)
    print(f"Document successfully created and updated: {output_filename}")

if __name__ == "__main__":
    create_document()
