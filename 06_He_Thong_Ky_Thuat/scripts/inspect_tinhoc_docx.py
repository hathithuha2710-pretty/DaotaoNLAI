import sys, io
sys.stdout.reconfigure(encoding='utf-8')
import docx

docx_path = r"06_He_Thong_Ky_Thuat\extracted_data\Chuong-trinh-giao-duc-pho-thong-2018-mon-tin-hoc.docx"
doc = docx.Document(docx_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

headings = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text.startswith(('I.', 'II.', 'III.', 'IV.', 'V.', 'VI.', 'VII.', 'VIII.', '1.', '2.', '3.')) and len(text) < 100:
        if any(p.text.strip().isupper() for _ in [1]):
            headings.append((i, text))

print("Found headings (sample):")
for idx, h in headings[:25]:
    print(f"  P{idx}: {h}")
