# -*- coding: utf-8 -*-
"""
Script tạo trọn bộ file Khung YCCĐ Rút gọn (Markdown và Word) cho Lớp 3, Lớp 4, Lớp 5:
- SP1_Khung_YCCD_AI_Lop3_Rut_Gon.md & .docx
- SP1_Khung_YCCD_AI_Lop4_Rut_Gon.md & .docx
- SP1_Khung_YCCD_AI_Lop5_Rut_Gon.md & .docx
Đồng thời tạo cả bản Full Word (.docx) cho Lớp 3, 4, 5 nếu cần đồng bộ hoàn chỉnh.
"""
import os, sys
if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

# ==============================================================================
# DỮ LIỆU LỚP 3
# ==============================================================================
data_lop3 = [
    {
        "stt": 1, "code": "3.A1.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Cách sử dụng AI trong học tập",
        "compact": "Nhận biết được một số tình huống sử dụng AI hỗ trợ học sinh trong quá trình học tập (như tìm kiếm tư liệu, giải thích từ mới, gợi ý ý tưởng)."
    },
    {
        "stt": 2, "code": "3.A1.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Cách sử dụng AI trong học tập",
        "compact": "Nêu được cách sử dụng trợ lí học tập thông minh, ứng dụng học ngôn ngữ để hỗ trợ quá trình học tập và rèn luyện phát âm."
    },
    {
        "stt": 3, "code": "3.A1.MR1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Cách sử dụng AI trong học tập",
        "compact": "Thực hành sử dụng trợ lí học tập thông minh, ứng dụng học ngôn ngữ an toàn dưới sự hướng dẫn sư phạm của giáo viên."
    },
    {
        "stt": 4, "code": "3.A1.3",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Không phụ thuộc hoàn toàn vào AI",
        "compact": "Phân biệt được giữa việc sử dụng AI như công cụ hỗ trợ (chủ động, có kiểm soát) và việc để AI làm thay toàn bộ (thụ động, ỷ lại)."
    },
    {
        "stt": 5, "code": "3.A1.MR2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Không phụ thuộc hoàn toàn vào AI",
        "compact": "Nêu được hậu quả của việc phụ thuộc quá mức vào AI: làm giảm khả năng tư duy độc lập, giảm sáng tạo và giảm khả năng giải quyết vấn đề."
    },
    {
        "stt": 6, "code": "3.A1.4",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Suy nghĩ kĩ trước khi dùng AI",
        "compact": "Nêu được một số tình huống có thể gây rủi ro hoặc hậu quả không mong muốn khi sử dụng AI trong học tập và đời sống."
    },
    {
        "stt": 7, "code": "3.A1.5",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Suy nghĩ kĩ trước khi dùng AI",
        "compact": "Tập thói quen tự đặt câu hỏi và suy nghĩ trước khi dùng AI: tự đặt và trả lời một số câu hỏi trước khi quyết định sử dụng AI trong tình huống cụ thể."
    },
    {
        "stt": 8, "code": "3.A1.MR3",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Suy nghĩ kĩ trước khi dùng AI",
        "compact": "Giải thích được rằng việc suy nghĩ và kiểm tra trước khi dùng AI giúp con người phòng tránh rủi ro, bảo vệ an toàn và dữ liệu cá nhân."
    },
    {
        "stt": 9, "code": "3.A2.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI trong trường học",
        "compact": "Nhận biết được một số ứng dụng hoặc thiết bị có sử dụng AI trong trường học (robot trợ giảng, bảng thông minh, ứng dụng nhận dạng chữ viết)."
    },
    {
        "stt": 10, "code": "3.A2.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI trong trường học",
        "compact": "Nêu được mục đích của AI trong trường học: hỗ trợ giáo viên và học sinh, cá nhân hoá việc học theo năng lực và kích hoạt hứng thú học tập."
    },
    {
        "stt": 11, "code": "3.A2.MR1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI trong trường học",
        "compact": "Mô tả được mô hình trường học mong muốn trong tương lai – nơi có sự hỗ trợ của các thiết bị và ứng dụng AI thông minh, nhân văn."
    },
    {
        "stt": 12, "code": "3.A2.MR2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI trong trường học",
        "compact": "Đề xuất được một số nguyên tắc và quy ước văn hóa, đạo đức khi sử dụng công cụ AI trong học tập tại nhà trường."
    },
    {
        "stt": 13, "code": "3.A2.3",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI hỗ trợ mọi người",
        "compact": "Nêu được các đối tượng trong trường học mà AI có thể hỗ trợ (học sinh, giáo viên, nhân viên nhà trường)."
    },
    {
        "stt": 14, "code": "3.A2.MR3",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI hỗ trợ mọi người",
        "compact": "Nêu được mục tiêu của AI trong bối cảnh trường học là hỗ trợ tất cả thành viên, giúp việc dạy và học dễ dàng, hiệu quả và công bằng hơn."
    },
    {
        "stt": 15, "code": "3.A3.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Công dân trong kỉ nguyên AI – Kiểm tra và phản biện kết quả của AI",
        "compact": "Nêu được ví dụ cụ thể về việc AI có thể sai (dịch sai nghĩa, nhận dạng nhầm ảnh, phần mềm hội thoại tự động đưa thông tin sai sự thật)."
    },
    {
        "stt": 16, "code": "3.A3.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Công dân trong kỉ nguyên AI – Kiểm tra và phản biện kết quả của AI",
        "compact": "Trình bày được rằng khi sử dụng AI cần kiểm tra lại kết quả, so sánh với nhiều nguồn thông tin và hỏi ý kiến giáo viên để xác minh."
    },
    {
        "stt": 17, "code": "3.B2.1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Sử dụng AI an toàn và có trách nhiệm – Phân biệt thật và giả",
        "compact": "Nhận biết và nêu được ví dụ về việc thông tin hoặc sản phẩm do AI tạo ra có thể không đúng sự thật."
    },
    {
        "stt": 18, "code": "3.B3.1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Nguyên tắc đạo đức và trách nhiệm xã hội – Cùng máy thông minh làm việc tốt",
        "compact": "Nhận biết được rằng không nên tạo ra hoặc sử dụng AI cho các mục đích xấu như lừa đảo, bắt nạt hay gây hại cho người khác."
    },
    {
        "stt": 19, "code": "3.B3.MR1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Nguyên tắc đạo đức và trách nhiệm xã hội – Cùng máy thông minh làm việc tốt",
        "compact": "Nêu được rằng con người cần cung cấp dữ liệu đúng và đa dạng để AI hoạt động chính xác, công bằng và tránh thiên lệch."
    },
    {
        "stt": 20, "code": "3.C4.1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Dữ liệu trong AI – Dữ liệu học máy",
        "compact": "Trình bày được khái niệm dữ liệu thông qua các ví dụ cụ thể (hình ảnh, con số, chữ viết và âm thanh)."
    },
    {
        "stt": 21, "code": "3.C4.MR1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Dữ liệu trong AI – Dữ liệu học máy",
        "compact": "Xác định được các thuộc tính, đặc trưng nhận diện cơ bản (màu sắc, hình dạng, kích thước) của một bộ dữ liệu đơn giản."
    },
    {
        "stt": 22, "code": "3.C5.1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Kĩ thuật AI dựa trên luật",
        "compact": "Mô tả được cấu trúc logic điều kiện 'Nếu … thì …' trong việc giải quyết tình huống hoặc phân loại đồ vật."
    },
    {
        "stt": 23, "code": "3.C5.MR1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Kĩ thuật AI dựa trên luật",
        "compact": "Nêu được một số công cụ AI đơn giản hoạt động dựa trên cấu trúc quy tắc logic 'Nếu … thì …' do con người cài đặt."
    },
    {
        "stt": 24, "code": "3.C5.2",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Kĩ thuật học máy",
        "compact": "Nêu được đặc điểm chính của học máy là học từ dữ liệu mẫu (máy tính học quy luật dựa trên các ví dụ do con người cung cấp)."
    },
    {
        "stt": 25, "code": "3.C5.3",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Kĩ thuật học máy",
        "compact": "Nêu được một số tình huống ứng dụng học máy trong thực tiễn đời sống (phân loại nông sản, dự báo mưa lũ)."
    },
    {
        "stt": 26, "code": "3.C5.4",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Kĩ thuật học máy",
        "compact": "Nhận biết được bài toán học máy phân loại (Classification) và bài toán học máy dự đoán (Prediction) thông qua tình huống trực quan, cụ thể."
    },
    {
        "stt": 27, "code": "3.C5.MR2",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Kĩ thuật học máy",
        "compact": "Minh họa được bài toán học máy phân loại và bài toán học máy dự đoán bằng các công cụ tương tác trực quan (Teachable Machine, ML for Kids)."
    },
    {
        "stt": 28, "code": "3.D1.1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Nhận diện và hình thành giải pháp – Quá trình huấn luyện máy thông minh",
        "compact": "Trình bày được quy trình học máy tối giản gồm 2 bước: thu thập ví dụ mẫu và cho máy tính học hỏi từ tập mẫu đó."
    },
    {
        "stt": 29, "code": "3.D2.1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Dữ liệu tốt cho máy thông minh",
        "compact": "Nêu được rằng dữ liệu dùng để dạy AI có thể bị thiếu hụt, nhầm lẫn hoặc chứa thông tin không đúng sự thật."
    },
    {
        "stt": 30, "code": "3.D2.2",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Dữ liệu tốt cho máy thông minh",
        "compact": "Nêu được một số yêu cầu cơ bản đối với dữ liệu huấn luyện: số lượng đầy đủ, nhãn dán chính xác và hình ảnh rõ ràng."
    },
    {
        "stt": 31, "code": "3.D2.MR1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Dữ liệu tốt cho máy thông minh",
        "compact": "Thực hành thao tác làm sạch dữ liệu (Data Cleaning): kiểm tra bộ thẻ mẫu và loại bỏ các thẻ bị gán nhãn nhầm hoặc sai sự thật."
    },
    {
        "stt": 32, "code": "3.D2.3",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Máy thông minh có thể học sai",
        "compact": "Trình bày được rằng nếu dữ liệu đầu vào sai hoặc không tốt thì AI có thể hoạt động kém hiệu quả hoặc đưa ra kết quả sai."
    }
]

# ==============================================================================
# DỮ LIỆU LỚP 4
# ==============================================================================
data_lop4 = [
    {
        "stt": 1, "code": "4.A1.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – AI trong công việc hằng ngày",
        "compact": "Trình bày được một số lĩnh vực hoặc công việc cụ thể mà AI có thể hỗ trợ con người (như nông nghiệp chính xác, y tế, giao thông vận tải)."
    },
    {
        "stt": 2, "code": "4.A1.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – AI hỗ trợ, con người suy nghĩ",
        "compact": "Nhận biết được rằng AI được tạo ra để hỗ trợ con người học tập và làm việc hiệu quả hơn nhưng không thể thay thế tư duy, cảm xúc và sự sáng tạo của con người."
    },
    {
        "stt": 3, "code": "4.A1.MR1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – AI hỗ trợ, con người suy nghĩ",
        "compact": "Nêu được ví dụ cụ thể về AI hỗ trợ học tập và giải thích được vì sao người học vẫn cần tự suy nghĩ, kiểm tra và hiểu bài thay vì chỉ chép kết quả của AI."
    },
    {
        "stt": 4, "code": "4.A2.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI vì cuộc sống tốt đẹp hơn",
        "compact": "Nêu được rằng AI được con người tạo ra để hỗ trợ các hoạt động trong đời sống xã hội, giúp giải quyết vấn đề, tiết kiệm thời gian và nâng cao chất lượng cuộc sống."
    },
    {
        "stt": 5, "code": "4.A2.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI trong xã hội",
        "compact": "Trình bày được các đối tượng trong xã hội mà AI có thể hỗ trợ, bao gồm: người lao động, người yếu thế vượt qua rào cản cuộc sống."
    },
    {
        "stt": 6, "code": "4.A3.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Công dân trong kỉ nguyên AI – Con người quyết định khi dùng AI",
        "compact": "Trình bày được rằng việc quyết định có sử dụng AI hay không phụ thuộc vào mục đích, nhu cầu thực tế và sự an toàn của con người."
    },
    {
        "stt": 7, "code": "4.A3.MR1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Công dân trong kỉ nguyên AI – Con người quyết định khi dùng AI",
        "compact": "Thực hành đưa ra quyết định có sử dụng AI hay không trong tình huống cụ thể: chọn dùng khi an toàn, có ích và từ chối khi có nguy cơ xâm phạm riêng tư hay đạo đức."
    },
    {
        "stt": 8, "code": "4.B2.1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Sử dụng AI an toàn và có trách nhiệm – Bảo vệ thông tin cá nhân",
        "compact": "Nêu được một số loại thông tin cá nhân cần giữ bí mật và không nên chia sẻ cho AI (như họ tên đầy đủ, ngày sinh, địa chỉ nhà, trường lớp, số điện thoại, mật khẩu)."
    },
    {
        "stt": 9, "code": "4.B2.2",
        "strand": "B. Đạo đức AI",
        "subtopic": "Sử dụng AI an toàn và có trách nhiệm – Bảo vệ thông tin cá nhân",
        "compact": "Nêu được một số hậu quả khi thông tin cá nhân bị lộ hoặc bị lợi dụng trên không gian mạng (như bị làm phiền, bị mạo danh, bị lừa đảo)."
    },
    {
        "stt": 10, "code": "4.B2.MR1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Sử dụng AI an toàn và có trách nhiệm – Bảo vệ thông tin cá nhân",
        "compact": "Thực hành xử lí theo hướng dẫn một tình huống giả định đơn giản về việc bị lộ thông tin cá nhân khi dùng AI (như lập tức báo ngay cho thầy cô, bố mẹ và dừng ứng dụng)."
    },
    {
        "stt": 11, "code": "4.C2.1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Ứng dụng AI trong học tập và cuộc sống – Một số ứng dụng AI quen thuộc",
        "compact": "Nêu được các ứng dụng của AI trong học tập và đời sống, đặc biệt là các ứng dụng gần gũi với bối cảnh Việt Nam (nhận diện tiếng Việt, nông nghiệp thông minh)."
    },
    {
        "stt": 12, "code": "4.C2.MR1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Ứng dụng AI trong học tập và cuộc sống – Một số ứng dụng AI quen thuộc",
        "compact": "Thực hành trải nghiệm một số ứng dụng AI trong học tập và đời sống dưới sự hướng dẫn an toàn của giáo viên."
    },
    {
        "stt": 13, "code": "4.C5.MR1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Làm quen với một số công cụ trải nghiệm kĩ thuật học máy",
        "compact": "Thực hiện lại được từng bước theo hướng dẫn mẫu của giáo viên các thao tác cơ bản với công cụ học máy trực quan (như Teachable Machine, ML for Kids): (1) Đặt tên nhóm; (2) Đưa ảnh mẫu vào; (3) Bấm lệnh cho máy học; (4) Đưa ảnh mới vào quan sát kết quả."
    },
    {
        "stt": 14, "code": "4.C5.MR2",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Làm quen với một số công cụ trải nghiệm kĩ thuật học máy",
        "compact": "Lặp lại được đầy đủ 4 bước trên với một nhóm dữ liệu tự chọn khác và mô tả được thao tác ở mỗi bước đã làm dưới sự hỗ trợ của giáo viên."
    },
    {
        "stt": 15, "code": "4.D1.1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Nhận diện và hình thành giải pháp – Từ vấn đề đến ý tưởng AI",
        "compact": "Nêu được một số ý tưởng ban đầu về cách AI giúp giải quyết vấn đề gần gũi với bối cảnh Việt Nam (ví dụ: AI học từ ảnh các loại rác để phân loại rác tái chế)."
    },
    {
        "stt": 16, "code": "4.D1.MR1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Nhận diện và hình thành giải pháp – Từ vấn đề đến ý tưởng AI",
        "compact": "Trình bày được cách sử dụng AI để giải quyết vấn đề trong tình huống đơn giản gần gũi (ví dụ: AI học từ giọng nói, chữ viết dân tộc thiểu số để dịch sang tiếng phổ thông)."
    },
    {
        "stt": 17, "code": "4.D2.1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Liên tục cải tiến AI",
        "compact": "Trình bày được rằng con người cần liên tục đánh giá và nâng cấp sản phẩm để hệ thống AI cho ra kết quả tốt hơn."
    }
]

# ==============================================================================
# DỮ LIỆU LỚP 5
# ==============================================================================
data_lop5 = [
    {
        "stt": 1, "code": "5.A1.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Con người chịu trách nhiệm",
        "compact": "Trình bày được vai trò của tự động hóa thông minh trong các công việc lặp lại, nguy hiểm hoặc đòi hỏi độ chính xác cao (robot nhà máy, kiểm tra lỗi sản xuất, xe tự lái)."
    },
    {
        "stt": 2, "code": "5.A1.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Con người chịu trách nhiệm",
        "compact": "Nêu được ví dụ cụ thể khẳng định trách nhiệm của con người: khi AI chẩn đoán nhầm bệnh, bác sĩ phải kiểm tra lại và chịu trách nhiệm điều chỉnh."
    },
    {
        "stt": 3, "code": "5.A1.MR1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Tính chủ động của con người – Con người chịu trách nhiệm",
        "compact": "Trình bày được con người chịu trách nhiệm cuối cùng về mọi quyết định hoặc kết quả do AI tạo ra; không phó mặc hay đổ lỗi cho máy tính."
    },
    {
        "stt": 4, "code": "5.A2.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI không thay thế con người",
        "compact": "Giải thích được rằng mục đích chính của AI là hỗ trợ con người, không thể thay thế vai trò, tư duy, cảm xúc và trách nhiệm của con người trong xã hội."
    },
    {
        "stt": 5, "code": "5.A2.MR1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI không thay thế con người",
        "compact": "Khẳng định những năng lực con người giữ vai trò trung tâm: sáng tạo, thấu cảm, phán đoán đạo đức – con người luôn cần kiểm soát sự phát triển của AI."
    },
    {
        "stt": 6, "code": "5.A2.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI phục vụ lợi ích chung",
        "compact": "Trình bày được rằng AI được tạo ra nhằm phục vụ lợi ích chung của xã hội, cải thiện chất lượng cuộc sống và giải quyết các vấn đề phức tạp."
    },
    {
        "stt": 7, "code": "5.A2.3",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "AI vì sự tiến bộ của con người – AI phục vụ lợi ích chung",
        "compact": "Nêu được ví dụ cụ thể về việc AI mang lại lợi ích cho cộng đồng trong các lĩnh vực: y tế, giáo dục và bảo vệ môi trường."
    },
    {
        "stt": 8, "code": "5.A3.1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Công dân trong kỉ nguyên AI – Con người trong kỉ nguyên AI",
        "compact": "Trình bày được rằng mọi công dân thuộc mọi thế hệ (trẻ em, người lớn, người cao tuổi) đều cần hiểu và biết cách sử dụng AI an toàn, hiệu quả."
    },
    {
        "stt": 9, "code": "5.A3.2",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Công dân trong kỉ nguyên AI – Con người trong kỉ nguyên AI",
        "compact": "Trình bày được việc sử dụng AI phải phù hợp nhu cầu, tránh lạm dụng, không chia sẻ thông tin cá nhân và không tin tưởng tuyệt đối vào AI."
    },
    {
        "stt": 10, "code": "5.A3.MR1",
        "strand": "A. Tư duy lấy con người làm trung tâm",
        "subtopic": "Công dân trong kỉ nguyên AI – Con người trong kỉ nguyên AI",
        "compact": "Trình bày được ý tưởng về một số tình huống trong tương lai mà AI hỗ trợ trẻ em, người lớn hoặc người cao tuổi trong đời sống."
    },
    {
        "stt": 11, "code": "5.B1.1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Các khía cạnh đạo đức của AI – Hệ thống AI công bằng",
        "compact": "Nêu được ví dụ về sự công bằng hoặc không công bằng trong một số tình huống đơn giản trong đời sống và khi AI phục vụ con người."
    },
    {
        "stt": 12, "code": "5.B1.2",
        "strand": "B. Đạo đức AI",
        "subtopic": "Các khía cạnh đạo đức của AI – Hệ thống AI công bằng",
        "compact": "Giải thích được rằng AI cần phục vụ con người công bằng, không phân biệt giới tính, vùng miền hay điều kiện kinh tế."
    },
    {
        "stt": 13, "code": "5.B2.1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Sử dụng AI an toàn và có trách nhiệm – Giúp AI công bằng",
        "compact": "Nêu được một số cách giúp AI hoạt động công bằng hơn: sử dụng dữ liệu đa dạng, tránh định kiến và kiểm tra kết quả trên nhiều trường hợp/nhóm khác nhau."
    },
    {
        "stt": 14, "code": "5.B2.MR1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Sử dụng AI an toàn và có trách nhiệm – Giúp AI công bằng",
        "compact": "Đề xuất được một số quy tắc khi thu thập dữ liệu nhằm đảm bảo tính đại diện đa dạng và giảm thiểu thiên lệch dữ liệu."
    },
    {
        "stt": 15, "code": "5.B3.1",
        "strand": "B. Đạo đức AI",
        "subtopic": "Nguyên tắc đạo đức và trách nhiệm xã hội – Cần hiểu cách AI suy nghĩ",
        "compact": "Giải thích được vì sao con người cần hiểu cách AI đưa ra quyết định để đảm bảo tính minh bạch và độ tin cậy khi sử dụng AI."
    },
    {
        "stt": 16, "code": "5.C5.1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Thuật toán AI dựa trên luật",
        "compact": "Nêu được cách sử dụng cấu trúc 'Nếu … thì …' trong lập trình thuật toán AI đơn giản để phản hồi thông minh theo quy tắc."
    },
    {
        "stt": 17, "code": "5.C5.MR1",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Thuật toán AI dựa trên luật",
        "compact": "Thực hành sử dụng cấu trúc 'Nếu … thì …' trong lập trình khối lệnh Scratch/Blockly để tạo nhân vật phản hồi thông minh theo luật."
    },
    {
        "stt": 18, "code": "5.C5.2",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Làm quen với một số công cụ ứng dụng học máy trực quan",
        "compact": "Thực hiện được các thao tác cơ bản trên công cụ trải nghiệm AI trực quan: tạo nhóm nhãn, nạp mẫu và quan sát kết quả dự đoán (có thể quan sát điểm số tin cậy nếu có)."
    },
    {
        "stt": 19, "code": "5.C5.MR2",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Làm quen với một số công cụ ứng dụng học máy trực quan",
        "compact": "Huấn luyện được một mô hình phân loại đơn giản trên công cụ học máy trực quan (ví dụ: phân loại ảnh lá cây khỏe mạnh và lá cây sâu bệnh)."
    },
    {
        "stt": 20, "code": "5.C5.MR3",
        "strand": "C. Các kĩ thuật và ứng dụng AI",
        "subtopic": "Kĩ thuật và thuật toán AI – Làm quen với một số công cụ ứng dụng học máy trực quan",
        "compact": "Kiểm tra được kết quả do mô hình đưa ra với các mẫu ảnh mới và chỉ ra được những trường hợp mô hình phân loại sai."
    },
    {
        "stt": 21, "code": "5.D1.1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Nhận diện và hình thành giải pháp – Quy trình huấn luyện AI",
        "compact": "Mô tả được ở mức độ đơn giản 4 bước trong quy trình huấn luyện AI: Xác định vấn đề → Thu thập dữ liệu → Dạy máy học → Kiểm tra và đánh giá kết quả."
    },
    {
        "stt": 22, "code": "5.D1.MR1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Nhận diện và hình thành giải pháp – Quy trình huấn luyện AI",
        "compact": "Thực hành thu thập và chuẩn bị dữ liệu mẫu nhằm huấn luyện mô hình AI giải quyết một vấn đề đơn giản."
    },
    {
        "stt": 23, "code": "5.D2.1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Cải tiến hệ thống AI bằng dữ liệu",
        "compact": "Giải thích được bằng ví dụ rằng hệ thống AI có thể được cải tiến khi dữ liệu được bổ sung, cập nhật và hệ thống được huấn luyện, điều chỉnh phù hợp."
    },
    {
        "stt": 24, "code": "5.D2.MR1",
        "strand": "D. Thiết kế hệ thống AI",
        "subtopic": "Cấu trúc và tương tác, cải tiến hệ thống – Cải tiến hệ thống AI bằng dữ liệu",
        "compact": "Thực hành cải tiến hệ thống AI bằng cách bổ sung dữ liệu mới, thực hiện tái huấn luyện (Retrain) và kiểm tra xem kết quả có cải thiện hay không."
    }
]

def generate_markdown_rut_gon(grade, data, output_path):
    lines = [
        f"# KHUNG NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO – LỚP {grade}",
        "*(Bản rút gọn chuẩn hóa: Mạch năng lực, Chủ đề con và Yêu cầu cần đạt đầu ra)*",
        "",
        "> **Căn cứ pháp lý:** Quyết định số 2422/QĐ-BGDĐT của Bộ Giáo dục và Đào tạo.  ",
        "> **Đặc điểm:** Tinh gọn các chú giải nguồn, diễn đạt thành các Chuẩn đầu ra năng lực (Exit Competencies) liền mạch, súc tích, phục vụ trực tiếp việc xây dựng kế hoạch dạy học và giảng dạy thực tế.",
        "",
        "---",
        "",
        "| STT | Mạch năng lực | Chủ đề con | YCCĐ (Chuẩn đầu ra năng lực) |",
        "| :---: | :--- | :--- | :--- |"
    ]
    for it in data:
        row = f"| **{it['stt']}** | **{it['strand']}** | {it['subtopic']} | {it['compact']} |"
        lines.append(row)
    
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"Exported MD compact: {output_path}")

def generate_docx_rut_gon(grade, data, output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"KHUNG NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO – LỚP {grade}")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 32, 96)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("(Bản rút gọn chuẩn hóa: Mạch năng lực, Chủ đề con và Yêu cầu cần đạt đầu ra)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)
    p_sub.paragraph_format.space_after = Pt(12)

    headers = [
        ("STT", Inches(0.6)),
        ("Mạch năng lực", Inches(1.8)),
        ("Chủ đề con", Inches(2.0)),
        ("YCCĐ", Inches(2.5))
    ]
    
    table = doc.add_table(rows=len(data) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    hdr_cells = table.rows[0].cells
    for idx, (h_text, w) in enumerate(headers):
        hdr_cells[idx].width = w
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=100, right=100)
        set_cell_shading(hdr_cells[idx], "1B365D")
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))

    for row_idx, it in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F7F9FC" if row_idx % 2 == 0 else "FFFFFF"

        for idx, (_, w) in enumerate(headers):
            row_cells[idx].width = w
            set_cell_margins(row_cells[idx], top=80, bottom=80, left=90, right=90)
            set_cell_shading(row_cells[idx], bg_color)

        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(str(it["stt"]))
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9.5)
        r0.font.bold = True

        p1 = row_cells[1].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(it["strand"])
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(9.5)
        r1.font.bold = True

        p2 = row_cells[2].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(it["subtopic"])
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(9.5)

        p3 = row_cells[3].paragraphs[0]
        p3.paragraph_format.line_spacing = 1.15
        r3 = p3.add_run(it["compact"])
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(9.5)

    doc.save(output_path)
    print(f"Exported DOCX compact: {output_path}")

def generate_docx_from_full_md(grade, md_path, output_path):
    """Tạo file docx bản đầy đủ (5 cột khổ ngang) từ file Markdown đầy đủ đã có"""
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = [line.strip() for line in f if line.strip().startswith('| **')]
    
    doc = docx.Document()
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"BẢNG ĐỐI SÁNH CHUẨN ĐẦU RA AI LỚP {grade} (QĐ 2422/QĐ-BGDĐT – ĐA NGUỒN QUỐC TẾ)")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 32, 96)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Bản đầy đủ: Đối sánh nguyên văn chuẩn Bộ GD&ĐT và Chuẩn đầu ra tham chiếu Quốc tế (UNESCO, OECD, Singapore MOE, AI4K12, Google DeepMind)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)
    p_sub.paragraph_format.space_after = Pt(10)

    headers = [
        ("STT", Inches(0.6)),
        ("Mã YCCĐ", Inches(0.9)),
        ("Mạch nội dung & Chủ đề con", Inches(2.3)),
        ("Yêu cầu cần đạt Chuẩn Quốc gia\n(Bộ GD&ĐT - QĐ 2422)", Inches(3.3)),
        ("Chuẩn đầu ra tham chiếu Đa nguồn Quốc tế\n(UNESCO, OECD, Singapore MOE, AI4K12, Google)", Inches(3.4))
    ]
    
    table = doc.add_table(rows=len(lines) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    hdr_cells = table.rows[0].cells
    for idx, (h_text, w) in enumerate(headers):
        hdr_cells[idx].width = w
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=100, right=100)
        set_cell_shading(hdr_cells[idx], "1B365D")
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))

    for row_idx, line in enumerate(lines, start=1):
        parts = [p.strip() for p in line.split('|')[1:-1]]
        stt_txt = parts[0].replace('**', '')
        code_txt = parts[1].replace('**', '')
        strand_topic_txt = parts[2].replace('**', '')
        vn_txt = parts[3]
        intl_txt = parts[4]

        row_cells = table.rows[row_idx].cells
        bg_color = "F7F9FC" if row_idx % 2 == 0 else "FFFFFF"

        for idx, (_, w) in enumerate(headers):
            row_cells[idx].width = w
            set_cell_margins(row_cells[idx], top=80, bottom=80, left=90, right=90)
            set_cell_shading(row_cells[idx], bg_color)

        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(stt_txt)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9.5)
        r0.font.bold = True

        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(code_txt)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(0, 32, 96)

        p2 = row_cells[2].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        p2.paragraph_format.space_after = Pt(2)
        st_parts = strand_topic_txt.split('<br>')
        r2_a = p2.add_run(st_parts[0] + "\n")
        r2_a.font.name = "Times New Roman"
        r2_a.font.size = Pt(9)
        r2_a.font.bold = True
        if len(st_parts) > 1:
            r2_b = p2.add_run(st_parts[1].replace('*', ''))
            r2_b.font.name = "Times New Roman"
            r2_b.font.size = Pt(8.5)
            r2_b.font.italic = True
            r2_b.font.color.rgb = RGBColor(60, 60, 60)

        p3 = row_cells[3].paragraphs[0]
        p3.paragraph_format.line_spacing = 1.15
        r3 = p3.add_run(vn_txt)
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(9)

        p4 = row_cells[4].paragraphs[0]
        p4.paragraph_format.line_spacing = 1.15
        lines_intl = intl_txt.split("<br>")
        for l_idx, line_txt in enumerate(lines_intl):
            if l_idx > 0:
                p4 = row_cells[4].add_paragraph()
                p4.paragraph_format.line_spacing = 1.15
                p4.paragraph_format.space_after = Pt(2)
            c_parts = line_txt.split(":", 1)
            if len(c_parts) == 2:
                r_bold = p4.add_run(c_parts[0] + ":")
                r_bold.font.name = "Times New Roman"
                r_bold.font.size = Pt(8.5)
                r_bold.font.bold = True
                r_bold.font.color.rgb = RGBColor(0, 51, 102)
                r_norm = p4.add_run(c_parts[1])
                r_norm.font.name = "Times New Roman"
                r_norm.font.size = Pt(8.5)
            else:
                r_all = p4.add_run(line_txt)
                r_all.font.name = "Times New Roman"
                r_all.font.size = Pt(8.5)

    doc.save(output_path)
    print(f"Exported DOCX full: {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\Users\USER\Desktop\AI-vibe-code\KhungNLĐTAI\03_San_Pham_Chinh_Thuc"
    
    # 1. Tạo bản rút gọn Lớp 3
    generate_markdown_rut_gon(3, data_lop3, os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop3_Rut_Gon.md"))
    generate_docx_rut_gon(3, data_lop3, os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop3_Rut_Gon.docx"))
    generate_docx_from_full_md(3, os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop3_Day_Du.md"), os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop3_Day_Du.docx"))
    
    # 2. Tạo bản rút gọn Lớp 4
    generate_markdown_rut_gon(4, data_lop4, os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop4_Rut_Gon.md"))
    generate_docx_rut_gon(4, data_lop4, os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop4_Rut_Gon.docx"))
    generate_docx_from_full_md(4, os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop4_Day_Du.md"), os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop4_Day_Du.docx"))
    
    # 3. Tạo bản rút gọn Lớp 5
    generate_markdown_rut_gon(5, data_lop5, os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop5_Rut_Gon.md"))
    generate_docx_rut_gon(5, data_lop5, os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop5_Rut_Gon.docx"))
    generate_docx_from_full_md(5, os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop5_Day_Du.md"), os.path.join(out_dir, "SP1_Khung_YCCD_AI_Lop5_Day_Du.docx"))

    print("SUCCESS: Toàn bộ file Word và Markdown Lớp 3, 4, 5 (Đầy đủ & Rút gọn) đã được tạo thành công!")
