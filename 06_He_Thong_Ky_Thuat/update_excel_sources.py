# -*- coding: utf-8 -*-
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')

excel_path = r"c:\Users\USER\Desktop\AI-vibe-code\KhungNLĐTAI\04_Hoc_Lieu_Va_Tham_Khao\tong_hop_chuong_trinh_AI (1).xlsx"
wb = openpyxl.load_workbook(excel_path)

# Create or replace sheet 'Học liệu AI Quốc tế Chi tiết'
sheet_name = 'Học liệu AI Quốc tế Chi tiết'
if sheet_name in wb.sheetnames:
    del wb[sheet_name]

ws = wb.create_sheet(title=sheet_name)
ws.views.sheetView[0].showGridLines = True

headers = [
    "STT", "Nhóm phân loại", "Tên chương trình / Học liệu", "Tổ chức bảo trợ", 
    "Đối tượng lứa tuổi", "Đường dẫn chính thức (URL)", "Nội dung & Tài nguyên cung cấp", "Ánh xạ Mạch năng lực QĐ 2422"
]

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

data_rows = [
    (
        1, "I. K-12 Curriculum", "Experience AI", "Google DeepMind + Raspberry Pi Foundation (Anh Quốc)",
        "8–12 tuổi (Tiểu học) & 11–14 tuổi (THCS)", "https://experience-ai.org/en/units",
        "Trọn gói Lesson plans, Teacher guides, Slide PPTX, Worksheets, Glossary, Rubric. Khóa Discovering AI (8-12) & 2 Module chuyên sâu (11-14) kèm Foundations of AI (6 bài học thực hành huấn luyện mô hình).",
        "Mạch C (Kỹ thuật AI, Học máy), Mạch B (Thiên lệch - Bias in bias out), Mạch D (Thiết kế giải pháp AI)"
    ),
    (
        2, "I. K-12 Curriculum", "AI Literacy Lessons for Grades 6–12", "Common Sense Education (Hoa Kỳ)",
        "Lớp 6 đến Lớp 12 (11–18 tuổi)", "https://www.commonsense.org/education/collections/ai-literacy-lessons-for-grades-6-12",
        "Giáo trình AI Literacy, an toàn dữ liệu, nhận diện tin giả / Deepfake, liêm chính học thuật, đạo đức và tác động xã hội của AI tạo sinh.",
        "Mạch B (Đạo đức AI, An toàn số), Mạch A (Công dân trong kỷ nguyên AI, Trách nhiệm số)"
    ),
    (
        3, "I. K-12 Curriculum", "Day of AI Curriculum Resources", "MIT RAISE + i2 Learning (Hoa Kỳ)",
        "Toàn bộ K-12 (Tiểu học, THCS, THPT)", "https://dayofai.org/curriculum-resources",
        "Giáo trình ngày hội AI của MIT, bài tập không máy tính (Unplugged), kịch bản trải nghiệm, dự án sáng tạo AI, bộ slide bài giảng và phiếu thoát vé (Exit ticket).",
        "Mạch A (Tư duy người làm trung tâm), Mạch C (Nhận diện mẫu, Dự đoán), Mạch D (Ý tưởng robot thông minh)"
    ),
    (
        4, "I. K-12 Curriculum", "AI Quests", "Google Research + Stanford University (Hoa Kỳ)",
        "11–14 tuổi (THCS)", "https://grow.google/intl/europe/ai-skills/",
        "Game-based learning / Problem-solving: Nhiệm vụ -> Vấn đề -> Dữ liệu -> AI -> Ra quyết định -> Phản tư (Reflection).",
        "Mạch D (Nhận diện & hình thành giải pháp hệ thống), Mạch A (Tính chủ động & ra quyết định)"
    ),
    (
        5, "I. K-12 Curriculum", "Teaching Responsible Use of AI", "Google for Education",
        "Học sinh phổ thông K-12", "https://edu.google.com/ai/education/",
        "Khung 5 nguyên tắc đạo đức: AI là công nghệ; AI hỗ trợ - không thay thế; Luôn phản biện; Nhận diện deepfake; Bảo vệ dữ liệu cá nhân.",
        "Mạch A (Tư duy người làm trung tâm), Mạch B (Sử dụng AI an toàn và có trách nhiệm)"
    ),
    (
        6, "II. Teacher Training", "Google AI Educator Series", "Google + ISTE / ASCD",
        "Giáo viên K-12", "https://edu.google.com/learning-center/google-ai-educator-series/",
        "Micro-learning (15 phút/module): Hiểu nền tảng AI, ứng dụng thiết kế sư phạm, tối ưu soạn bài giảng và đánh giá rubric.",
        "Bồi dưỡng năng lực giáo viên triển khai 4 mạch QĐ 2422"
    ),
    (
        7, "II. Teacher Training", "Generative AI for Educators with Gemini", "Google + MIT RAISE",
        "Giáo viên mọi cấp học", "https://grow.google/intl/europe/ai-skills/",
        "Khóa đào tạo 2 giờ: Kỹ thuật câu lệnh (Prompting), tạo tài liệu phân hóa, cá nhân hóa học tập, đạo đức và kiểm soát ảo giác AI.",
        "Nâng cao kỹ năng ứng dụng AI tạo sinh cho đội ngũ sư phạm"
    ),
    (
        8, "Khung chuẩn quốc tế", "UNESCO AI Competency Framework for Students", "UNESCO (Liên Hợp Quốc)",
        "Học sinh phổ thông K-12", "https://www.unesco.org",
        "Khung tham chiếu trung tâm 4 chiều năng lực (Human-centred, Ethics, Techniques, System Design) x 3 mức (Understand, Apply, Create).",
        "Là Khung tham chiếu quốc tế trung tâm định hình toàn bộ chương trình QĐ 2422"
    ),
    (
        9, "Khung chuẩn quốc tế", "AI4K12 - Five Big Ideas in AI", "AAAI + CSTA + NSF (Hoa Kỳ)",
        "K-12 (chia 4 dải tuổi K-2, 3-5, 6-8, 9-12)", "https://ai4k12.org",
        "Tiến trình 5 Big Ideas: Perception, Representation & Reasoning, Machine Learning, Natural Interaction, Societal Impact.",
        "Cơ sở khoa học máy tính sâu sắc cho Mạch C và Mạch D"
    ),
    (
        10, "Khung chính sách quốc gia", "OECD Learning Compass 2030 & Singapore MOE NDLP", "OECD & Bộ Giáo dục Singapore",
        "Học sinh phổ thông", "https://www.oecd.org & https://www.moe.gov.sg",
        "Năng lực tự chủ số (Digital Agency), chống ỷ lại nhận thức công nghệ (Cognitive Offloading), bảo vệ an toàn trẻ em và Cyber Wellness.",
        "Trọng tâm định hướng Mạch A (Tính chủ động) và Mạch B (Trách nhiệm xã hội)"
    )
]

for row_idx, rdata in enumerate(data_rows, 2):
    for col_idx, val in enumerate(rdata, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = Font(name="Calibri", size=10)
        cell.border = border_thin
        if col_idx in [1, 5]:
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif col_idx in [2, 4]:
            cell.alignment = Alignment(horizontal="left", vertical="top")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

# Adjust column widths
col_widths = {1: 8, 2: 22, 3: 30, 4: 28, 5: 22, 6: 35, 7: 45, 8: 35}
for col_idx, w in col_widths.items():
    ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

wb.save(excel_path)
print("Updated excel successfully with new sheet: 'Học liệu AI Quốc tế Chi tiết'!")
