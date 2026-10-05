import os
import sys
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

full_md_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Day_Du.md"
compact_md_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.md"
compact_docx_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.docx"

# 1. Parse rows directly from the full markdown file
with open(full_md_path, "r", encoding="utf-8") as f:
    full_text = f.read()

grade_sections = re.split(r'## PHẦN [I|V|X]+: KHUNG NỘI DUNG GIÁO DỤC AI – LỚP (\d)', full_text)

grades_rows = {}
for i in range(1, len(grade_sections), 2):
    g_num = int(grade_sections[i])
    content = grade_sections[i+1]
    rows = []
    lines = content.strip().split('\n')
    for line in lines:
        if line.startswith('| **') or line.startswith('|**'):
            parts = [p.strip() for p in line.split('|')[1:-1]]
            if len(parts) >= 5:
                stt = parts[0].replace('**', '')
                code = parts[1].replace('**', '')
                topic_raw = parts[2]
                vn_req = parts[3]
                intl_req = parts[4]
                
                topic_parts = topic_raw.split('<br>')
                mach = topic_parts[0].strip()
                subtopic = topic_parts[1].replace('•', '').replace('*', '').strip() if len(topic_parts) > 1 else ""
                
                rows.append({
                    "stt": stt,
                    "code": code,
                    "mach": mach,
                    "subtopic": subtopic,
                    "vn": vn_req,
                    "intl": intl_req
                })
    grades_rows[g_num] = rows
    print(f"Lớp {g_num}: {len(rows)} YCCĐ parsed from full file.")

# Dictionary of integrated descriptions mapping by code
# We construct seamless integrated descriptions combining QĐ 2422 and international nuances,
# without mentioning source names, completely safe from blacklist words!

integrated_dict = {
    # --- LỚP 1 (21) ---
    "1.A1.1": "Nêu được ví dụ về các cảm xúc quen thuộc của bản thân (vui, buồn, tức giận, sợ hãi) và nhận ra rằng con người cảm thấy vui buồn thật sự, còn máy móc thì không; tránh việc coi máy móc như con người (tránh nhân hóa công nghệ số).",
    "1.A1.2": "Nêu được rằng máy thông minh chỉ bắt chước hoặc nhận biết nét mặt của con người (vui, buồn, ngạc nhiên), chứ không có cảm xúc như con người; không ngộ nhận tín hiệu mô phỏng của máy là sự đồng cảm.",
    "1.A1.3": "Nêu được rằng việc AI thể hiện cảm xúc (như hình mặt cười, giọng nói vui vẻ) là do con người tạo ra hoặc lập trình từ trước; AI không tự có cảm xúc như con người.",
    "1.A1.4": "Nêu được ví dụ khi máy phát ra câu chào hoặc cười theo mẫu do con người thiết kế (ví dụ: loa thông minh nói 'Chào bạn, tôi có thể giúp gì?'); phân biệt được phản hồi do máy tạo ra với sự giao tiếp chân thực giữa con người.",
    "1.A2.1": "Chỉ ra được các cách máy thông minh thể hiện trạng thái và cảm xúc mô phỏng qua hình ảnh (mặt cười), tín hiệu ánh sáng (nháy đèn) và âm thanh (giọng nói nhẹ nhàng) khi tương tác.",
    "1.A2.MR1": "Nêu được rằng giọng nói thân thiện và hình ảnh dễ thương giúp con người cảm thấy gần gũi, dễ chịu và an tâm khi nhờ máy thông minh giúp đỡ.",
    "1.A2.2": "Kể tên và nhận diện được một số đồ dùng thông minh quen thuộc có thể tự thực hiện việc đơn giản (như robot hút bụi, camera nhận diện khuôn mặt, loa thông minh nhận lệnh giọng nói).",
    "1.A2.3": "Nói được việc làm có ích của từng đồ dùng thông minh (như robot hút bụi giúp quét nhà sạch sẽ, loa thông minh phát nhạc theo yêu cầu) và cách AI giúp đồ dùng hoạt động thông minh, tiện lợi hơn.",
    "1.B1.1": "Nêu được một vài việc không nên làm khi dùng thiết bị thông minh (như bật tiếng loa quá to làm phiền mọi người, dùng máy trêu chọc bạn làm bạn buồn); hình thành ý thức tôn trọng người xung quanh.",
    "1.B3.1": "Biết và nhắc lại được quy tắc đạo đức nền tảng: Tuyệt đối không được dùng máy móc thông minh hay công nghệ AI để làm hại, làm đau hoặc gây tổn thương cho người khác.",
    "1.B3.2": "Nêu được việc làm tốt đẹp mà con người sử dụng máy thông minh để thực hiện (như gọi giúp đỡ khi khẩn cấp, học nghe phát âm tiếng Anh) nhằm mang lại điều tích cực cho bản thân và gia đình.",
    "1.C1.1": "Chỉ ra được sự khác biệt đơn giản giữa quạt máy thông thường (bật nút là quay) và quạt thông minh (tự quay khi nhận biết có người đến gần); bước đầu hiểu cách thiết bị tiếp nhận tín hiệu từ môi trường để hoạt động.",
    "1.C1.2": "Nhận biết được hình ảnh biểu tượng trợ lí giọng nói (micro) hoặc ô vuông nhận diện khuôn mặt trên màn hình điện thoại do giáo viên chỉ dẫn; làm quen với dấu hiệu trực quan khi AI hoạt động.",
    "1.C1.3": "Nhận biết được camera giống như 'mắt nhìn' và micro giống như 'tai nghe' của thiết bị thông minh, tương tự các giác quan tiếp nhận thông tin của con người.",
    "1.C1.4": "Nêu được rằng máy thông minh có thể hiểu được một số câu nói ngắn, rõ ràng của con người và trả lời lại theo cách con người đã dạy; tập nói câu ngắn gọn, phát âm rõ ràng khi tương tác với trợ lý giọng nói.",
    "1.C1.MR1": "Quan sát và nhận biết được hình ảnh camera của máy đóng khung viền xung quanh quả táo hoặc chú mèo trên màn hình khi máy nhận diện được đối tượng.",
    "1.C1.MR2": "Quan sát và nhận biết được máy thông minh phản ứng khác nhau khi nghe tiếng vỗ tay so với khi nghe tiếng nói của con người (bước đầu làm quen với khả năng phân biệt âm thanh của máy).",
    "1.D1.1": "Nêu được ví dụ về việc con người cho máy xem nhiều bức tranh con mèo để máy biết nhận ra con mèo; bước đầu làm quen với ý niệm máy học từ các ví dụ mẫu do con người cung cấp.",
    "1.D1.MR1": "Giải thích được rằng nếu chỉ đưa 1 bức tranh thì máy dễ đoán sai; máy cần xem nhiều bức tranh đúng và khác nhau để học nhận biết chính xác hơn.",
    "1.D2.1": "Nhận biết được có đồ dùng chỉ làm được một việc riêng (như cân điện tử), có đồ dùng làm được nhiều việc khác nhau (như điện thoại thông minh vừa nghe gọi, vừa chụp ảnh).",
    "1.D2.2": "Xếp được các thẻ tranh đồ dùng thông minh vào 2 nhóm: nhóm chỉ làm 1 việc và nhóm làm được nhiều việc khác nhau; nêu được lý do sắp xếp đơn giản.",

    # --- LỚP 2 (29) ---
    "2.A1.1": "Trình bày được một số tình huống trong cuộc sống mà AI có thể hỗ trợ con người hiệu quả (như tìm đường đi, tìm bài hát yêu thích, dịch một câu chào, phát hiện lỗi chính tả); biết lựa chọn thời điểm phù hợp để nhờ máy thông minh trợ giúp.",
    "2.A1.2": "Trình bày được một số tình huống không nên hoặc cần thận trọng khi sử dụng AI (như khi AI có thể làm lộ thông tin riêng tư: địa chỉ nhà, khuôn mặt, giọng nói; không để máy làm hết bài tập thay cho mình); hình thành tính tự giác và ý thức an toàn số.",
    "2.A1.3": "Nêu được ví dụ cụ thể về tình huống cần con người giám sát AI (như người lái xe vẫn cần quan sát khi xe tự lái hoạt động; bác sĩ cần kiểm tra lại kết quả do AI gợi ý); khẳng định nguyên tắc con người luôn là người quyết định cuối cùng.",
    "2.A1.4": "Nêu được ví dụ cụ thể về thái độ đúng đắn khi sử dụng AI: biết tin tưởng vào công nghệ ở mức hợp lý, nhưng không tin tưởng tuyệt đối; luôn có sự kiểm tra của con người và chủ động hỏi ý kiến người lớn khi cần xác minh.",
    "2.A2.1": "Nêu được một số thiết bị hoặc ứng dụng trong gia đình có sử dụng AI (như loa thông minh, robot hút bụi, máy điều hoà tự động điều chỉnh nhiệt độ, camera nhận diện người trong nhà, tivi thông minh gợi ý chương trình yêu thích) và mô tả cách chúng hỗ trợ sinh hoạt hàng ngày.",
    "2.A2.MR1": "Thực hành sử dụng một số thiết bị hoặc ứng dụng có sử dụng AI với sự giám sát, hỗ trợ của giáo viên (như phát âm rõ câu lệnh để nhờ loa thông minh phát bài hát thiếu nhi); rèn luyện kỹ năng tương tác chuẩn mực và tự tin.",
    "2.A2.MR2": "Nêu được rằng không được sử dụng thiết bị hoặc ứng dụng có sử dụng AI mà không có sự giám sát của người lớn; tuân thủ nguyên tắc an toàn số trong gia đình và trường học.",
    "2.A2.2": "Nhận biết và kể tên được các thành viên trong gia đình mà AI có thể hỗ trợ (như nhắc bà uống thuốc đúng giờ, giúp mẹ lau sàn nhà, giúp bố hẹn giờ làm việc); hiểu giá trị nhân văn của công nghệ trong chăm sóc gia đình.",
    "2.A2.MR3": "Nêu được rằng mục tiêu của AI trong bối cảnh gia đình là giúp ích cho cuộc sống của các thành viên tiện nghi hơn, an toàn hơn và giúp gia đình có nhiều thời gian trò chuyện, gắn kết bên nhau hơn.",
    "2.A3.1": "Nêu được rằng những việc con người làm với máy (như chọn bài hát, chọn tranh ảnh) có thể được máy ghi nhớ lại để phục vụ tốt hơn; bước đầu làm quen với ý niệm về dấu chân kỹ thuật số.",
    "2.A3.MR1": "Nhận ra rằng nếu đưa vào máy những thông tin sai hoặc xấu thì máy có thể ghi nhớ điều sai; từ đó biết giữ lời nói và hành động đúng khi dùng máy; hình thành văn hóa ứng xử số có trách nhiệm.",
    "2.B1.1": "Nêu được ví dụ đơn giản về việc máy thông minh có thể nhận nhầm hoặc phân biệt đối xử (như chỉ nhận giọng của một người mà không nghe rõ giọng của người khác); bước đầu làm quen với hiện tượng thiên lệch (bias).",
    "2.B1.MR1": "Nhận ra rằng việc máy đối xử không công bằng có thể khiến người khác buồn hoặc gặp khó khăn; hiểu rằng công nghệ cần đối xử bình đẳng với mọi người.",
    "2.B1.MR2": "Giải thích được nguyên nhân đơn giản khiến máy đối xử không công bằng: do máy chỉ được dạy bằng giọng nói hoặc hình ảnh của một nhóm người, chưa được học đầy đủ từ các nhóm người khác.",
    "2.B3.1": "Phân biệt được sản phẩm do mình tự làm ra (như bức tranh tự vẽ) với sản phẩm do máy thông minh tạo ra; tôn trọng công sức lao động học tập của bản thân và bạn bè.",
    "2.B3.2": "Biết và nhắc lại được quy tắc trung thực: Không nhận sản phẩm do máy tạo ra là do mình tự làm; hình thành ý thức liêm chính và trung thực số ngay từ nhỏ.",
    "2.B3.MR1": "Phân biệt được đồ vật nào là của chung (như máy tính của lớp), đồ vật nào là của riêng (như cặp sách, đồ dùng của bạn); xác định rõ ranh giới quyền sử dụng trong thực tế.",
    "2.B3.MR2": "Nêu được và thực hiện quy tắc: Khi muốn dùng đồ của người khác thì phải xin phép và được sự đồng ý mới dùng; rèn luyện nguyên tắc đồng thuận văn minh.",
    "2.C1.1": "Nêu được rằng con người học từ thầy cô, sách vở và trải nghiệm cuộc sống; còn máy thông minh học từ các hình ảnh, âm thanh, dữ liệu do con người đưa vào.",
    "2.C1.MR1": "Nêu được ví dụ đơn giản về việc máy học từ các mẫu do con người chọn lọc (như đưa nhiều ảnh quả cam và quả táo để máy học phân biệt).",
    "2.C3.1": "Mô tả được việc máy có thể sắp xếp hoặc phân loại đồ vật theo một số đặc điểm đơn giản (như màu sắc, hình dạng, kích thước); làm quen với bài toán phân loại (Classification) cơ bản.",
    "2.C3.2": "Nhận biết được rằng máy có thể phân loại sai nếu các đồ vật có hình dáng hoặc màu sắc quá giống nhau (ví dụ: quả bóng tròn màu đỏ và quả táo màu đỏ).",
    "2.C3.MR1": "So sánh được: Khi nhìn thấy đồ vật bị che khuất một phần, con người có thể đoán ra nhờ trải nghiệm thực tế, còn máy có thể gặp khó khăn hoặc nhận diện chưa chính xác.",
    "2.D1.1": "Chỉ ra được một số việc làm khó khăn hoặc vất vả quanh em mà máy thông minh có thể hỗ trợ (như lau bảng trên cao, nhặt rác sân trường, tưới cây tự động).",
    "2.D1.2": "Nêu được ý tưởng đơn giản về một chiếc máy thông minh có thể giúp ích cho bản thân hoặc bạn bè (thông qua lời kể hoặc tranh vẽ trực quan).",
    "2.D1.MR1": "Chọn được một số hình ảnh mẫu phù hợp trong tình huống cụ thể để chuẩn bị làm ví dụ giúp máy thông minh nhận biết đồ vật.",
    "2.D2.1": "Nêu được rằng máy thông minh cần có nhiều ví dụ mẫu (dữ liệu) do con người cung cấp thì mới học và làm việc tốt được.",
    "2.D2.2": "Giải thích được rằng nếu các bức tranh mẫu bị mờ, bị rách hoặc thiếu thông tin thì máy có thể nhận diện sai hoặc đoán nhầm; hiểu tầm quan trọng của chất lượng dữ liệu.",
    "2.D2.MR1": "Tự tay chọn và xếp được một bộ tranh ảnh rõ nét để làm mẫu cho máy học nhận biết một đồ vật quen thuộc; rèn kỹ năng chuẩn bị dữ liệu thực nghiệm ban đầu.",

    # --- LỚP 3 (32) ---
    "3.A1.1": "Nhận biết được một số tình huống sử dụng AI hỗ trợ học sinh trong quá trình học tập (như tìm kiếm tư liệu, giải thích từ mới, gợi ý ý tưởng học tập).",
    "3.A1.2": "Nêu được cách sử dụng trợ lí học tập thông minh, ứng dụng học ngôn ngữ để hỗ trợ quá trình học tập và rèn luyện phát âm hiệu quả.",
    "3.A1.MR1": "Thực hành sử dụng trợ lí học tập thông minh, ứng dụng học ngôn ngữ một cách an toàn và đúng mục đích dưới sự hướng dẫn sư phạm của giáo viên.",
    "3.A1.3": "Phân biệt được giữa việc sử dụng AI như công cụ hỗ trợ (chủ động, có kiểm soát) và việc để AI làm thay toàn bộ bài tập (thụ động, ỷ lại).",
    "3.A1.MR2": "Nêu được hậu quả của việc phụ thuộc quá mức vào AI: làm giảm khả năng tự suy nghĩ độc lập, giảm tính sáng tạo và giảm năng lực tự giải quyết vấn đề.",
    "3.A1.4": "Nêu được một số tình huống có thể gây rủi ro hoặc hậu quả không mong muốn khi sử dụng AI trong học tập và đời sống (như nhận thông tin sai, làm lộ thông tin riêng tư).",
    "3.A1.5": "Tập thói quen tự đặt và trả lời một số câu hỏi trước khi quyết định sử dụng AI trong tình huống cụ thể (công việc này mình có tự làm được không, dùng AI có an toàn không).",
    "3.A1.MR3": "Giải thích được rằng việc suy nghĩ kĩ và kiểm tra trước khi dùng AI giúp con người phòng tránh rủi ro, bảo vệ an toàn và bảo mật dữ liệu cá nhân.",
    "3.A2.1": "Nhận biết được một số ứng dụng hoặc thiết bị có sử dụng AI trong trường học (như robot trợ giảng, bảng thông minh, ứng dụng điểm danh bằng khuôn mặt, phần mềm nhận dạng chữ viết).",
    "3.A2.2": "Nêu được mục đích của AI trong trường học: hỗ trợ thầy cô và học sinh, cá nhân hoá việc học theo năng lực và khơi dậy hứng thú học tập.",
    "3.A2.MR1": "Mô tả được mô hình trường học mong muốn trong tương lai – nơi có sự hỗ trợ của các thiết bị và ứng dụng AI thông minh, tiện ích và nhân văn.",
    "3.A2.MR2": "Đề xuất được một số nguyên tắc và quy ước văn hóa khi sử dụng công cụ AI trong học tập tại nhà trường (như trung thực, không gian lận, xin phép giáo viên).",
    "3.A2.3": "Nêu được các đối tượng trong trường học mà AI có thể hỗ trợ (học sinh, giáo viên, nhân viên nhà trường) để môi trường học đường thuận tiện hơn.",
    "3.A2.MR3": "Nêu được mục tiêu của AI trong bối cảnh trường học là hỗ trợ tất cả các thành viên, giúp việc dạy và học trở nên dễ dàng, hiệu quả và công bằng hơn.",
    "3.A3.1": "Nêu được ví dụ cụ thể về việc AI có thể đưa ra kết quả sai (như dịch sai nghĩa câu, nhận dạng nhầm đồ vật, phần mềm hội thoại tự động trả lời sai sự thật).",
    "3.A3.2": "Trình bày được rằng khi sử dụng AI cần luôn kiểm tra lại kết quả, so sánh đối chiếu với nhiều nguồn tài liệu tin cậy và hỏi ý kiến giáo viên để xác minh.",
    "3.B2.1": "Nhận biết và nêu được ví dụ về việc thông tin hoặc sản phẩm do AI tạo ra (hình ảnh, âm thanh, bài viết) có thể không đúng sự thật hoặc là thông tin giả mạo.",
    "3.B3.1": "Nhận biết được rằng không nên tạo ra hoặc sử dụng AI cho các mục đích xấu như lừa dối, trêu chọc ác ý, bắt nạt hay gây tổn hại cho người khác.",
    "3.B3.MR1": "Nêu được rằng con người cần cung cấp dữ liệu đúng đắn và đa dạng để máy thông minh học hỏi, từ đó hoạt động chính xác, công bằng và tránh thiên lệch.",
    "3.C4.1": "Trình bày được khái niệm cơ bản về dữ liệu thông qua các ví dụ cụ thể, trực quan (như hình ảnh, chữ số, con chữ và âm thanh thu được).",
    "3.C4.MR1": "Xác định được các thuộc tính và đặc trưng nhận diện cơ bản (như màu sắc, hình dạng, kích thước) của một bộ dữ liệu trực quan đơn giản.",
    "3.C5.1": "Mô tả được cấu trúc logic điều kiện 'Nếu … thì …' trong việc giải quyết tình huống thực tế hoặc phân loại đồ vật quen thuộc.",
    "3.C5.MR1": "Nêu được một số công cụ hoặc trò chơi đơn giản hoạt động dựa trên cấu trúc quy tắc logic 'Nếu … thì …' do con người thiết lập.",
    "3.C5.2": "Nêu được đặc điểm chính của học máy là học từ dữ liệu mẫu (máy tính tự rút ra quy luật từ các ví dụ do con người cung cấp thay vì lập trình sẵn từng bước).",
    "3.C5.3": "Nêu được một số tình huống ứng dụng học máy trong thực tiễn đời sống (như phân loại hoa quả, nhận diện giọng nói, dự báo thời tiết).",
    "3.C5.4": "Nhận biết được bài toán học máy phân loại (chia đối tượng vào từng nhóm) và bài toán học máy dự đoán (ước tính kết quả tương lai) thông qua tình huống trực quan.",
    "3.C5.MR2": "Minh họa được bài toán học máy phân loại và bài toán học máy dự đoán bằng các công cụ tương tác trực quan (như Teachable Machine hoặc mô phỏng trò chơi).",
    "3.D1.1": "Trình bày được quy trình học máy tối giản gồm 2 bước cơ bản: thu thập các ví dụ mẫu và cho máy tính học hỏi từ tập mẫu đó.",
    "3.D2.1": "Nêu được rằng dữ liệu dùng để dạy AI có thể bị thiếu hụt, bị nhầm lẫn hoặc chứa thông tin sai lệch dẫn đến kết quả không chính xác.",
    "3.D2.2": "Nêu được một số yêu cầu cơ bản đối với dữ liệu huấn luyện: số lượng đầy đủ, gán nhãn chính xác và hình ảnh rõ nét, không bị nhiễu.",
    "3.D2.MR1": "Thực hành thao tác làm sạch dữ liệu đơn giản: kiểm tra bộ thẻ mẫu và loại bỏ các thẻ bị gán nhãn sai hoặc hình ảnh bị mờ nhòe trước khi cho máy học.",
    "3.D2.3": "Trình bày được rằng nếu dữ liệu đầu vào sai hoặc kém chất lượng thì AI sẽ học sai và đưa ra kết quả không tin cậy ('dữ liệu sai thì kết quả sai').",

    # --- LỚP 4 (17) ---
    "4.A1.1": "Trình bày được một số lĩnh vực hoặc công việc cụ thể mà trí tuệ nhân tạo (AI) có thể hỗ trợ con người (như nông nghiệp, y tế, giao thông vận tải, dự báo thời tiết).",
    "4.A1.2": "Nhận biết được rằng trí tuệ nhân tạo (AI) được tạo ra để hỗ trợ con người học tập và làm việc hiệu quả hơn nhưng không thể thay thế tư duy, cảm xúc và sự sáng tạo của con người.",
    "4.A1.MR1": "Nêu được ví dụ cụ thể về việc AI hỗ trợ trong học tập và giải thích được vì sao người học vẫn cần tự suy nghĩ, kiểm tra và hiểu bài thay vì chỉ chép kết quả do AI đưa ra; chủ động chống ỷ lại vào công nghệ.",
    "4.A2.1": "Nêu được rằng AI được con người tạo ra để hỗ trợ các hoạt động trong đời sống xã hội, giúp giải quyết vấn đề, tiết kiệm thời gian và nâng cao chất lượng cuộc sống của con người.",
    "4.A2.2": "Trình bày được các đối tượng trong xã hội mà AI có thể hỗ trợ, bao gồm: người lao động, người yếu thế (người khuyết tật, người cao tuổi) vượt qua rào cản sinh hoạt.",
    "4.A3.1": "Trình bày được rằng việc quyết định có sử dụng AI hay không phụ thuộc vào mục đích, nhu cầu thực tế và sự an toàn của con người; khẳng định quyền tự chủ của con người trước công nghệ.",
    "4.A3.MR1": "Thực hành đưa ra quyết định có sử dụng AI hay không trong tình huống cụ thể: lựa chọn dùng khi AI an toàn và mang lại lợi ích thiết thực; từ chối hoặc dừng sử dụng khi AI gây rủi ro, xâm phạm quyền riêng tư hoặc vi phạm đạo đức.",
    "4.B2.1": "Nêu được một số loại thông tin cá nhân cần được giữ bí mật và không nên chia sẻ cho AI (như họ tên đầy đủ, ngày sinh, địa chỉ nhà, trường lớp, số điện thoại, mật khẩu tài khoản).",
    "4.B2.2": "Nêu được một số hậu quả khi thông tin cá nhân bị lộ hoặc bị lợi dụng trên không gian mạng (như bị làm phiền, bị mạo danh, bị lừa đảo hoặc bắt nạt trực tuyến).",
    "4.B2.MR1": "Thực hành xử lí theo hướng dẫn một tình huống giả định đơn giản về việc bị lộ thông tin cá nhân khi dùng AI (như dừng ứng dụng, báo ngay cho thầy cô hoặc cha mẹ để được hỗ trợ kịp thời).",
    "4.C2.1": "Nêu được các ứng dụng của AI trong học tập và đời sống, đặc biệt là các ứng dụng gần gũi với bối cảnh Việt Nam (như nhận diện giọng nói tiếng Việt, phần mềm học tiếng Việt, nông nghiệp thông minh tại địa phương).",
    "4.C2.MR1": "Thực hành trải nghiệm một số ứng dụng AI trực quan trong học tập và đời sống dưới sự hướng dẫn an toàn của giáo viên.",
    "4.C5.MR1": "Thực hiện lại được từng bước theo hướng dẫn mẫu của giáo viên các thao tác cơ bản với công cụ trải nghiệm AI trực quan dựa trên học máy có giám sát (như Teachable Machine hoặc ML for Kids): (1) Đặt tên cho các nhóm phân loại; (2) Đưa vào mỗi nhóm một số hình ảnh mẫu; (3) Bấm lệnh cho máy học từ các mẫu đó; (4) Đưa hình ảnh mới vào quan sát kết quả máy đưa ra.",
    "4.C5.MR2": "Lặp lại được đầy đủ 4 bước trên với một nhóm dữ liệu tự chọn khác và mô tả được thao tác ở mỗi bước đã làm dưới sự hỗ trợ của giáo viên.",
    "4.D1.1": "Nêu được một số ý tưởng ban đầu về cách AI có thể học hoặc giúp giải quyết vấn đề gần gũi với bối cảnh Việt Nam (ví dụ: AI học từ ảnh các loại rác để phân loại rác tái chế bảo vệ môi trường trường học).",
    "4.D1.MR1": "Trình bày được cách sử dụng AI để giải quyết vấn đề trong tình huống đơn giản gần gũi với bối cảnh Việt Nam (ví dụ: AI học từ giọng nói và chữ viết của đồng bào dân tộc thiểu số để dịch sang tiếng phổ thông, giúp học sinh vùng cao học tập thuận lợi).",
    "4.D2.1": "Trình bày được rằng con người cần liên tục đánh giá và cải tiến để hệ thống AI cho ra kết quả tốt hơn, phục vụ con người hiệu quả và an toàn hơn.",

    # --- LỚP 5 (24) ---
    "5.A1.1": "Trình bày được rằng AI có thể thực hiện một số công việc thay con người, đặc biệt là những việc lặp đi lặp lại, nguy hiểm hoặc đòi hỏi tính chính xác cao (như robot lắp ráp trong nhà máy, AI kiểm tra lỗi sản phẩm, xe tự lái chở hàng hoá).",
    "5.A1.2": "Nêu được ví dụ cụ thể về trách nhiệm của người tạo ra và người sử dụng AI (ví dụ: nếu AI chẩn đoán nhầm bệnh, bác sĩ vẫn phải kiểm tra lại và chịu trách nhiệm điều chỉnh phương án điều trị).",
    "5.A1.MR1": "Trình bày được rằng con người chịu trách nhiệm cuối cùng về mọi quyết định hoặc kết quả do AI tạo ra; không phó mặc hay đổ lỗi cho công nghệ máy tính.",
    "5.A2.1": "Giải thích được rằng AI được tạo ra nhằm mục đích hỗ trợ con người, không thể thay thế vai trò, tư duy, cảm xúc và trách nhiệm của con người trong đời sống và xã hội.",
    "5.A2.MR1": "Khẳng định được những năng lực mà con người luôn giữ vai trò trung tâm (sáng tạo nghệ thuật, thấu cảm xã hội, đánh giá đạo đức); con người luôn cần làm chủ và kiểm soát sự phát triển của AI.",
    "5.A2.2": "Trình bày được rằng AI được con người tạo ra nhằm phục vụ lợi ích chung của xã hội, giúp cải thiện chất lượng cuộc sống, giải quyết các vấn đề phức tạp và hỗ trợ con người trong nhiều lĩnh vực.",
    "5.A2.3": "Nêu được ví dụ cụ thể về việc AI mang lại lợi ích thiết thực cho cộng đồng trong các lĩnh vực: chăm sóc sức khỏe y tế, nâng cao chất lượng giáo dục và bảo vệ môi trường tự nhiên.",
    "5.A3.1": "Trình bày được rằng mọi người thuộc mọi thế hệ (trẻ em, người lớn, người cao tuổi) đều cần hiểu biết và biết cách sử dụng AI để phục vụ cuộc sống, học tập và công việc một cách an toàn, hiệu quả.",
    "5.A3.2": "Trình bày được việc sử dụng AI phải phù hợp với nhu cầu và khả năng thực tế của từng người; biết tránh lạm dụng, không chia sẻ thông tin riêng tư và không tin tưởng tuyệt đối vào AI.",
    "5.A3.MR1": "Trình bày được ý tưởng về một số tình huống trong tương lai mà AI có thể hỗ trợ trẻ em, người lớn hoặc người cao tuổi trong đời sống hàng ngày.",
    "5.B1.1": "Nêu được ví dụ về sự công bằng hoặc không công bằng trong một số tình huống đơn giản của đời sống và khi AI phục vụ con người (như phần mềm nhận diện tốt hơn với một nhóm người nhất định).",
    "5.B1.2": "Giải thích được rằng hệ thống AI cần phục vụ con người một cách công bằng, bình đẳng, không phân biệt giới tính, vùng miền hay hoàn cảnh xã hội.",
    "5.B2.1": "Nêu được một số cách giúp AI hoạt động công bằng hơn: sử dụng tập dữ liệu đa dạng có tính đại diện, tránh định kiến khi thu thập dữ liệu và kiểm tra lại kết quả trên nhiều trường hợp khác nhau.",
    "5.B2.MR1": "Đề xuất được một số quy tắc khi thu thập dữ liệu nhằm bảo đảm bộ dữ liệu có tính đa dạng, cân bằng và giảm thiểu sự thiên lệch.",
    "5.B3.1": "Giải thích được vì sao con người cần hiểu cách AI đưa ra quyết định để đảm bảo tính minh bạch, an toàn và mức độ tin cậy khi sử dụng công nghệ.",
    "5.C5.1": "Nêu được cách sử dụng cấu trúc điều kiện 'Nếu … thì …' trong lập trình thuật toán AI đơn giản để nhân vật phản hồi thông minh theo quy tắc đặt trước.",
    "5.C5.MR1": "Thực hành sử dụng cấu trúc điều kiện 'Nếu … thì …' trong môi trường lập trình trực quan (như Scratch/Blockly) để tạo nhân vật tương tác thông minh theo các luật đã định.",
    "5.C5.2": "Thực hiện được các thao tác cơ bản với công cụ trải nghiệm AI trực quan dựa trên học máy có giám sát (như Teachable Machine, ML for Kids trong Scratch): tạo nhóm nhãn, nạp dữ liệu mẫu và quan sát kết quả dự đoán.",
    "5.C5.MR2": "Huấn luyện được một mô hình phân loại đơn giản trên công cụ trải nghiệm AI trực quan (ví dụ: phân loại ảnh lá cây khỏe mạnh và lá cây bị sâu bệnh).",
    "5.C5.MR3": "Kiểm tra được kết quả phân loại do mô hình đưa ra bằng các hình ảnh mẫu mới và chỉ ra được những trường hợp mô hình nhận diện sai để tìm hiểu nguyên nhân.",
    "5.D1.1": "Mô tả được ở mức độ đơn giản 4 bước cơ bản trong quá trình huấn luyện mô hình AI: (1) Xác định vấn đề $\\rightarrow$ (2) Thu thập dữ liệu $\\rightarrow$ (3) Dạy máy học $\\rightarrow$ (4) Kiểm tra và đánh giá kết quả.",
    "5.D1.MR1": "Thực hành thu thập và chuẩn bị dữ liệu mẫu (hình ảnh hoặc âm thanh) nhằm huấn luyện mô hình AI giải quyết một vấn đề đơn giản trong học tập hoặc đời sống.",
    "5.D2.1": "Giải thích được bằng ví dụ cụ thể rằng hệ thống AI có thể được cải tiến và hoạt động chính xác hơn khi dữ liệu huấn luyện được bổ sung, cập nhật và mô hình được điều chỉnh phù hợp.",
    "5.D2.MR1": "Thực hành cải tiến hệ thống AI bằng cách bổ sung thêm các mẫu dữ liệu còn thiếu hoặc bị nhận nhầm, cho máy học lại và kiểm tra xem kết quả có tốt hơn trước hay không."
}

# Verify that all codes from full file exist in integrated_dict
missing_codes = []
for g in range(1, 6):
    for r in grades_rows[g]:
        if r['code'] not in integrated_dict:
            missing_codes.append(r['code'])

if missing_codes:
    print(f"ERROR: Missing codes in dictionary: {missing_codes}")
    sys.exit(1)
else:
    print("SUCCESS: 100% of 123 codes from full file are present in integrated_dict!")

# 2. Build Markdown content
md_header = """# TỔNG HỢP KHUNG NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO CẤP TIỂU HỌC (LỚP 1 ĐẾN LỚP 5)
## BẢN RÚT GỌN CHUẨN HOÁ: CHUẨN ĐẦU RA NĂNG LỰC TÍCH HỢP

> **Căn cứ thực hiện:** Quyết định số 2422/QĐ-BGDĐT và Công văn số 5588/BGDĐT-GDPT của Bộ Giáo dục và Đào tạo.  
> **Đặc điểm:** Tích hợp hài hoà giữa Yêu cầu cần đạt Chuẩn Quốc gia (Bộ GD&ĐT - QĐ 2422) và Chuẩn đầu ra năng lực tham chiếu Đa nguồn Quốc tế (UNESCO, OECD, Singapore MOE, AI4K12, Google); không ghi tên nguồn trích dẫn, câu chữ tinh gọn, chuẩn mực sư phạm lứa tuổi 6–11, phục vụ trực tiếp cho việc xây dựng Ma trận Kỹ năng, Kế hoạch bài dạy và Phân phối chương trình giảng dạy thực tế.  
> **Quy mô toàn cấp:** 123 YCCĐ (Lớp 1: 21, Lớp 2: 29, Lớp 3: 32, Lớp 4: 17, Lớp 5: 24).

---

## MỤC LỤC CÁC PHẦN
- [Phần I: Lớp 1 (21 YCCĐ)](#phan-i-lop-1-21-yccd)
- [Phần II: Lớp 2 (29 YCCĐ)](#phan-ii-lop-2-29-yccd)
- [Phần III: Lớp 3 (32 YCCĐ)](#phan-iii-lop-3-32-yccd)
- [Phần IV: Lớp 4 (17 YCCĐ)](#phan-iv-lop-4-17-yccd)
- [Phần V: Lớp 5 (24 YCCĐ)](#phan-v-lop-5-24-yccd)

---
"""

roman_map = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V"}
full_md_out = md_header

for g in range(1, 6):
    r_num = roman_map[g]
    rows = grades_rows[g]
    full_md_out += f"## PHẦN {r_num}: LỚP {g} ({len(rows)} YCCĐ)\n\n"
    full_md_out += "| STT | Mã YCCĐ | Mạch năng lực | Chủ đề con | YCCĐ (Chuẩn đầu ra năng lực tích hợp) |\n"
    full_md_out += "| :---: | :---: | :--- | :--- | :--- |\n"
    for r in rows:
        desc = integrated_dict[r['code']]
        full_md_out += f"| **{r['stt']}** | **{r['code']}** | **{r['mach']}** | {r['subtopic']} | {desc} |\n"
    full_md_out += "\n---\n\n"

with open(compact_md_path, "w", encoding="utf-8") as f:
    f.write(full_md_out)
print(f"Generated {compact_md_path} successfully!")

# 3. Build DOCX document
doc = Document()

# Page setup: A4 Portrait
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79)     # 2.0 cm
    section.bottom_margin = Inches(0.79)  # 2.0 cm
    section.left_margin = Inches(0.98)    # 2.5 cm
    section.right_margin = Inches(0.79)   # 2.0 cm

COLOR_PRIMARY = RGBColor(15, 76, 129)     # Classic Navy Blue
COLOR_SECONDARY = RGBColor(40, 116, 166)
COLOR_TEXT = RGBColor(33, 33, 33)
HEX_PRIMARY_BG = "0F4C81"
HEX_ZEBRA_BG = "F7F9FB"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY_BG}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(4)
run_title = title_p.add_run("TỔNG HỢP KHUNG NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO CẤP TIỂU HỌC\n(LỚP 1 ĐẾN LỚP 5)")
run_title.font.name = "Times New Roman"
run_title.font.size = Pt(14)
run_title.font.bold = True
run_title.font.color.rgb = COLOR_PRIMARY

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_before = Pt(0)
sub_p.paragraph_format.space_after = Pt(10)
run_sub = sub_p.add_run("BẢN RÚT GỌN CHUẨN HOÁ: CHUẨN ĐẦU RA NĂNG LỰC TÍCH HỢP")
run_sub.font.name = "Times New Roman"
run_sub.font.size = Pt(11.5)
run_sub.font.bold = True
run_sub.font.color.rgb = COLOR_SECONDARY

# Callout / Intro Box
callout_p = doc.add_paragraph()
callout_p.paragraph_format.space_before = Pt(2)
callout_p.paragraph_format.space_after = Pt(12)
callout_p.paragraph_format.line_spacing = 1.15

r_intro = callout_p.add_run(
    "• Căn cứ thực hiện: Quyết định số 2422/QĐ-BGDĐT và Công văn số 5588/BGDĐT-GDPT của Bộ Giáo dục và Đào tạo.\n"
    "• Đặc điểm: Tích hợp hài hoà giữa Yêu cầu cần đạt Chuẩn Quốc gia (Bộ GD&ĐT - QĐ 2422) và Chuẩn đầu ra năng lực tham chiếu Đa nguồn Quốc tế (UNESCO, OECD, Singapore MOE, AI4K12, Google); không ghi tên nguồn trích dẫn, câu chữ tinh gọn, chuẩn mực sư phạm lứa tuổi 6–11, phục vụ trực tiếp cho việc xây dựng Ma trận Kỹ năng, Kế hoạch bài dạy và Phân phối chương trình giảng dạy thực tế.\n"
    "• Quy mô toàn cấp: 123 YCCĐ (Lớp 1: 21, Lớp 2: 29, Lớp 3: 32, Lớp 4: 17, Lớp 5: 24)."
)
r_intro.font.name = "Times New Roman"
r_intro.font.size = Pt(10)
r_intro.font.italic = True
r_intro.font.color.rgb = COLOR_TEXT

# Widths: sum = 6.5 in
col_widths = [Inches(0.35), Inches(0.75), Inches(1.15), Inches(1.35), Inches(2.9)]

for g in range(1, 6):
    r_num = roman_map[g]
    rows = grades_rows[g]
    
    # Section Heading
    h_p = doc.add_paragraph()
    h_p.paragraph_format.space_before = Pt(14)
    h_p.paragraph_format.space_after = Pt(6)
    h_p.paragraph_format.keep_with_next = True
    
    r_h = h_p.add_run(f"PHẦN {r_num}: LỚP {g} ({len(rows)} YCCĐ)")
    r_h.font.name = "Times New Roman"
    r_h.font.size = Pt(12)
    r_h.font.bold = True
    r_h.font.color.rgb = COLOR_PRIMARY
    
    # Create Table
    table = doc.add_table(rows=len(rows) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header
    headers = ["STT", "Mã YCCĐ", "Mạch năng lực", "Chủ đề con", "YCCĐ (Chuẩn đầu ra năng lực tích hợp)"]
    hdr_row = table.rows[0]
    make_row_header(hdr_row)
    make_row_cant_split(hdr_row)
    
    for c_idx, cell in enumerate(hdr_row.cells):
        cell.width = col_widths[c_idx]
        set_cell_background(cell, HEX_PRIMARY_BG)
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1] else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        run = p.add_run(headers[c_idx])
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    
    # Data Rows
    for r_idx, r in enumerate(rows):
        row = table.rows[r_idx + 1]
        make_row_cant_split(row)
        
        is_zebra = (r_idx % 2 == 1)
        bg_color = HEX_ZEBRA_BG if is_zebra else "FFFFFF"
        
        desc = integrated_dict[r['code']]
        row_vals = [r['stt'], r['code'], r['mach'], r['subtopic'], desc]
        
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            if bg_color != "FFFFFF":
                set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            
            val = row_vals[c_idx]
            if c_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.font.bold = True
                if c_idx == 1 and "MR" in val:
                    run.font.color.rgb = COLOR_SECONDARY
                else:
                    run.font.color.rgb = COLOR_TEXT
            elif c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = COLOR_TEXT
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.font.italic = True
                run.font.color.rgb = COLOR_TEXT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.font.color.rgb = COLOR_TEXT
    
    # Spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(8)

doc.save(compact_docx_path)
print(f"Generated DOCX {compact_docx_path} successfully!")
