import os
import sys
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

md_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP2_Bang_Ma_Tran_Ky_Nang_AI_Tieu_Hoc_Lop1_5_Moi.md"
docx_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP2_Bang_Ma_Tran_Ky_Nang_AI_Tieu_Hoc_Lop1_5_Moi.docx"
xlsx_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP2_Bang_Ma_Tran_Ky_Nang_AI_Tieu_Hoc_Lop1_5_Moi.xlsx"

with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

# -------------------------------------------------------------
# 1. PARSE GROUPS AND SKILLS
# -------------------------------------------------------------
groups = [] # list of (group_name, list_of_skills)
current_group = "KHUNG MA TRẬN KỸ NĂNG"
current_skills = []

for line in lines:
    line_s = line.strip()
    if line_s.startswith("### NHÓM"):
        if current_skills:
            groups.append((current_group, current_skills))
            current_skills = []
        current_group = line_s.replace("###", "").strip()
    elif line_s.startswith("| **") or line_s.startswith("|**"):
        parts = [p.strip() for p in line_s.split("|")[1:-1]]
        if len(parts) >= 8:
            try:
                stt_num = int(parts[0].replace("**", ""))
                current_skills.append({
                    "stt": parts[0].replace("**", ""),
                    "name": parts[1],
                    "mach": parts[2].replace("`", ""),
                    "g1": parts[3],
                    "g2": parts[4],
                    "g3": parts[5],
                    "g4": parts[6],
                    "g5": parts[7]
                })
            except ValueError:
                pass

if current_skills:
    groups.append((current_group, current_skills))

total_parsed = sum(len(skills) for _, skills in groups)
print(f"Parsed {len(groups)} groups with total {total_parsed} skills.")

# -------------------------------------------------------------
# 2. GENERATE DOCX (A4 Landscape)
# -------------------------------------------------------------
doc = Document()
for section in doc.sections:
    section.page_width = Inches(11.69)   # A4 Landscape width
    section.page_height = Inches(8.27)   # A4 Landscape height
    section.top_margin = Inches(0.79)    # 2.0 cm
    section.bottom_margin = Inches(0.79) # 2.0 cm
    section.left_margin = Inches(0.98)   # 2.5 cm
    section.right_margin = Inches(0.79)  # 2.0 cm

COLOR_PRIMARY = RGBColor(15, 76, 129)     # Classic Navy Blue
COLOR_SECONDARY = RGBColor(40, 116, 166)
COLOR_TEXT = RGBColor(33, 33, 33)
HEX_PRIMARY_BG = "0F4C81"
HEX_ZEBRA_BG = "F7F9FB"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY_BG}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

# Title
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.space_before = Pt(0)
tp.paragraph_format.space_after = Pt(2)
rt = tp.add_run("BẢNG MA TRẬN PHÂN BỔ 46 KỸ NĂNG GIÁO DỤC AI CẤP TIỂU HỌC (LỚP 1 – 5)")
rt.font.name = "Times New Roman"
rt.font.size = Pt(14)
rt.font.bold = True
rt.font.color.rgb = COLOR_PRIMARY

sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sp.paragraph_format.space_before = Pt(0)
sp.paragraph_format.space_after = Pt(8)
rs = sp.add_run("BẢN CHUẨN HOÁ TIỂU HỌC THEO QUYẾT ĐỊNH 2422/QĐ-BGDĐT VÀ THAM CHIẾU QUỐC TẾ (SP2)")
rs.font.name = "Times New Roman"
rs.font.size = Pt(11)
rs.font.bold = True
rs.font.color.rgb = COLOR_SECONDARY

# Preamble note
cp = doc.add_paragraph()
cp.paragraph_format.space_before = Pt(0)
cp.paragraph_format.space_after = Pt(10)
cp.paragraph_format.line_spacing = 1.15
rc = cp.add_run(
    "• Căn cứ: Quyết định số 2422/QĐ-BGDĐT và Công văn số 5588/BGDĐT-GDPT của Bộ GD&ĐT.\n"
    "• Quy mô: 46 Kỹ năng tích hợp phân bổ trên 5 khối lớp, phủ trọn vẹn 100% (123/123) YCCĐ của Bộ GD&ĐT.\n"
    "• Rào chắn sư phạm: 100% phù hợp nhận thức lứa tuổi 6–11; loại bỏ hoàn toàn các khái niệm và công cụ THCS."
)
rc.font.name = "Times New Roman"
rc.font.size = Pt(9.5)
rc.font.italic = True
rc.font.color.rgb = COLOR_TEXT

# Printable width = 11.69 - 0.98 - 0.79 = 9.92 inches
# 8 cols: STT=0.35, Tên=1.45, Mạch=0.45, L1=1.5, L2=1.5, L3=1.5, L4=1.5, L5=1.67 -> sum = 9.92 in
col_widths = [Inches(0.35), Inches(1.45), Inches(0.45), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.67)]
headers = ["STT", "Tên kỹ năng", "Mạch", "Lớp 1", "Lớp 2", "Lớp 3", "Lớp 4", "Lớp 5"]

for g_title, g_skills in groups:
    gp = doc.add_paragraph()
    gp.paragraph_format.space_before = Pt(12)
    gp.paragraph_format.space_after = Pt(4)
    gp.paragraph_format.keep_with_next = True
    rg = gp.add_run(g_title.upper())
    rg.font.name = "Times New Roman"
    rg.font.size = Pt(11)
    rg.font.bold = True
    rg.font.color.rgb = COLOR_PRIMARY

    tbl = doc.add_table(rows=len(g_skills) + 1, cols=8)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)

    # Header Row
    hr = tbl.rows[0]
    make_row_header(hr)
    make_row_cant_split(hr)
    for c_idx, cell in enumerate(hr.cells):
        cell.width = col_widths[c_idx]
        set_cell_background(cell, HEX_PRIMARY_BG)
        set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(headers[c_idx])
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, sk in enumerate(g_skills):
        row = tbl.rows[r_idx + 1]
        make_row_cant_split(row)
        is_zebra = (r_idx % 2 == 1)
        bg = HEX_ZEBRA_BG if is_zebra else "FFFFFF"

        clean_name = sk['name'].replace('<br>', '\n').replace('*', '')
        vals = [sk['stt'], clean_name, sk['mach'],
                sk['g1'].replace('<br>', '\n'),
                sk['g2'].replace('<br>', '\n'),
                sk['g3'].replace('<br>', '\n'),
                sk['g4'].replace('<br>', '\n'),
                sk['g5'].replace('<br>', '\n')]

        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            if bg != "FFFFFF":
                set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1

            txt = vals[c_idx]
            if c_idx in [0, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(txt)
                run.font.name = "Times New Roman"
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = COLOR_TEXT
            elif c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(txt)
                run.font.name = "Times New Roman"
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = COLOR_TEXT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(txt)
                run.font.name = "Times New Roman"
                run.font.size = Pt(8)
                run.font.color.rgb = COLOR_TEXT

    sp_end = doc.add_paragraph()
    sp_end.paragraph_format.space_before = Pt(0)
    sp_end.paragraph_format.space_after = Pt(6)

doc.save(docx_path)
print(f"Generated DOCX: {docx_path} successfully!")

# -------------------------------------------------------------
# 3. GENERATE XLSX (Excel)
# -------------------------------------------------------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Ma Tran Ky Nang AI (1-5)"

font_header = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="0F4C81", end_color="0F4C81", fill_type="solid")
font_group = Font(name="Times New Roman", size=11, bold=True, color="0F4C81")
fill_group = PatternFill(start_color="EAF2F8", end_color="EAF2F8", fill_type="solid")
font_cell = Font(name="Times New Roman", size=9)
font_cell_bold = Font(name="Times New Roman", size=9, bold=True)
fill_zebra = PatternFill(start_color="F7F9FB", end_color="F7F9FB", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

align_center = Alignment(horizontal="center", vertical="top", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

# Title Rows
ws.merge_cells("A1:H1")
ws["A1"] = "BẢNG MA TRẬN PHÂN BỔ 46 KỸ NĂNG GIÁO DỤC AI CẤP TIỂU HỌC (LỚP 1 – 5)"
ws["A1"].font = Font(name="Times New Roman", size=14, bold=True, color="0F4C81")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells("A2:H2")
ws["A2"] = "BẢN CHUẨN HOÁ TIỂU HỌC THEO QUYẾT ĐỊNH 2422/QĐ-BGDĐT VÀ THAM CHIẾU QUỐC TẾ (SP2)"
ws["A2"].font = Font(name="Times New Roman", size=11, bold=True, color="2874A6")
ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

row_cursor = 4

# Header row in Excel
for col_idx, h_text in enumerate(headers, 1):
    cell = ws.cell(row=row_cursor, column=col_idx, value=h_text)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = thin_border
ws.row_dimensions[row_cursor].height = 25
row_cursor += 1

# Groups and data
for g_title, g_skills in groups:
    # Group header row
    ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=8)
    g_cell = ws.cell(row=row_cursor, column=1, value=g_title.upper())
    g_cell.font = font_group
    g_cell.fill = fill_group
    g_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row_cursor].height = 22
    row_cursor += 1

    for r_idx, sk in enumerate(g_skills):
        is_zebra = (r_idx % 2 == 1)
        row_fill = fill_zebra if is_zebra else None

        clean_name = sk['name'].replace('<br>', '\n').replace('*', '')
        row_data = [
            int(sk['stt']),
            clean_name,
            sk['mach'],
            sk['g1'].replace('<br>', '\n'),
            sk['g2'].replace('<br>', '\n'),
            sk['g3'].replace('<br>', '\n'),
            sk['g4'].replace('<br>', '\n'),
            sk['g5'].replace('<br>', '\n')
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_cursor, column=col_idx, value=val)
            cell.font = font_cell_bold if col_idx in [1, 2, 3] else font_cell
            if row_fill:
                cell.fill = row_fill
            cell.alignment = align_center if col_idx in [1, 3] else align_left
            cell.border = thin_border

        ws.row_dimensions[row_cursor].height = 65
        row_cursor += 1

# Set column widths
ws.column_dimensions['A'].width = 6
ws.column_dimensions['B'].width = 25
ws.column_dimensions['C'].width = 8
ws.column_dimensions['D'].width = 32
ws.column_dimensions['E'].width = 32
ws.column_dimensions['F'].width = 32
ws.column_dimensions['G'].width = 32
ws.column_dimensions['H'].width = 35

wb.save(xlsx_path)
print(f"Generated XLSX: {xlsx_path} successfully!")
