import os
import sys
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

md_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.md"
docx_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP1_Khung_YCCD_AI_Tieu_Hoc_Lop1_5_Rut_Gon.docx"

doc = Document()

# Page setup: A4 Portrait
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79)     # 2.0 cm
    section.bottom_margin = Inches(0.79)  # 2.0 cm
    section.left_margin = Inches(0.98)    # 2.5 cm
    section.right_margin = Inches(0.79)   # 2.0 cm

# Color palette: Elegant Navy
COLOR_PRIMARY = RGBColor(15, 76, 129)     # Classic Blue / Navy
COLOR_SECONDARY = RGBColor(40, 116, 166)
COLOR_TEXT = RGBColor(33, 33, 33)
COLOR_MUTED = RGBColor(100, 100, 100)
HEX_PRIMARY_BG = "0F4C81"
HEX_ZEBRA_BG = "F7F9FB"
HEX_BORDER = "CCCCCC"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY_BG}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(4)
run_title = title_p.add_run("TỔNG HỢP KHUNG NỘI DUNG GIÁO DỤC TRÍ TUỆ NHÂN TẠO CẤP TIỂU HỌC\n(LỚP 1 ĐẾN LỚP 5)")
run_title.font.name = "Times New Roman"
run_title.font.size = Pt(15)
run_title.font.bold = True
run_title.font.color.rgb = COLOR_PRIMARY

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_before = Pt(0)
sub_p.paragraph_format.space_after = Pt(12)
run_sub = sub_p.add_run("BẢN RÚT GỌN CHUẨN HOÁ: CHUẨN ĐẦU RA NĂNG LỰC TÍCH HỢP")
run_sub.font.name = "Times New Roman"
run_sub.font.size = Pt(12)
run_sub.font.bold = True
run_sub.font.color.rgb = COLOR_SECONDARY

# Callout Box / Metadata
callout_p = doc.add_paragraph()
callout_p.paragraph_format.space_before = Pt(4)
callout_p.paragraph_format.space_after = Pt(14)
callout_p.paragraph_format.line_spacing = 1.15

r_intro = callout_p.add_run(
    "• Căn cứ thực hiện: Quyết định số 2422/QĐ-BGDĐT và Công văn số 5588/BGDĐT-GDPT của Bộ Giáo dục và Đào tạo.\n"
    "• Đặc điểm: Tích hợp hài hoà giữa Yêu cầu cần đạt Chuẩn Quốc gia (Bộ GD&ĐT - QĐ 2422) và Chuẩn đầu ra năng lực tham chiếu Đa nguồn Quốc tế (UNESCO, OECD, Singapore MOE, AI4K12, Google); không ghi tên nguồn trích dẫn, câu chữ tinh gọn, chuẩn mực sư phạm lứa tuổi 6–11, phục vụ trực tiếp cho việc xây dựng Ma trận Kỹ năng, Kế hoạch bài dạy và Phân phối chương trình giảng dạy thực tế.\n"
    "• Quy mô toàn cấp: 123 YCCĐ (Lớp 1: 21, Lớp 2: 29, Lớp 3: 32, Lớp 4: 17, Lớp 5: 24)."
)
r_intro.font.name = "Times New Roman"
r_intro.font.size = Pt(10.5)
r_intro.font.italic = True
r_intro.font.color.rgb = COLOR_TEXT

# Read MD content
with open(md_path, "r", encoding="utf-8") as f:
    text = f.read()

sections = re.split(r'## PHẦN ([I|V|X]+): (LỚP \d \(.+?\))', text)

grade_idx = 1
for s in range(1, len(sections), 3):
    roman = sections[s]
    grade_title = sections[s+1]
    content = sections[s+2]
    
    # Section Heading
    h_p = doc.add_paragraph()
    h_p.paragraph_format.space_before = Pt(16)
    h_p.paragraph_format.space_after = Pt(8)
    h_p.paragraph_format.keep_with_next = True
    
    r_h = h_p.add_run(f"PHẦN {roman}: {grade_title.upper()}")
    r_h.font.name = "Times New Roman"
    r_h.font.size = Pt(13)
    r_h.font.bold = True
    r_h.font.color.rgb = COLOR_PRIMARY
    
    # Parse rows
    lines = content.strip().split('\n')
    table_rows = []
    for line in lines:
        if line.startswith('| **') or line.startswith('|**'):
            parts = [p.strip() for p in line.split('|')[1:-1]]
            if len(parts) >= 5:
                table_rows.append((
                    parts[0].replace('**', ''),
                    parts[1].replace('**', ''),
                    parts[2].replace('**', ''),
                    parts[3],
                    parts[4]
                ))
    
    if not table_rows:
        continue
    
    # Create Table: 5 columns
    # Printable width = 8.27 - 0.98 - 0.79 = 6.5 inches
    # Col widths: STT=0.35 in, Mã YCCĐ=0.75 in, Mạch=1.1 in, Chủ đề con=1.4 in, YCCĐ=2.9 in -> sum=6.5 in
    col_widths = [Inches(0.35), Inches(0.75), Inches(1.1), Inches(1.4), Inches(2.9)]
    
    table = doc.add_table(rows=len(table_rows) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header Row
    headers = ["STT", "Mã YCCĐ", "Mạch năng lực", "Chủ đề con", "YCCĐ (Chuẩn đầu ra năng lực tích hợp)"]
    hdr_row = table.rows[0]
    make_row_header(hdr_row)
    make_row_cant_split(hdr_row)
    
    for c_idx, cell in enumerate(hdr_row.cells):
        cell.width = col_widths[c_idx]
        set_cell_background(cell, HEX_PRIMARY_BG)
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1] else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        run = p.add_run(headers[c_idx])
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    
    # Data Rows
    for r_idx, row_data in enumerate(table_rows):
        row = table.rows[r_idx + 1]
        make_row_cant_split(row)
        
        is_zebra = (r_idx % 2 == 1)
        bg_color = HEX_ZEBRA_BG if is_zebra else "FFFFFF"
        
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            if bg_color != "FFFFFF":
                set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            
            text_val = row_data[c_idx]
            
            if c_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text_val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.bold = True
                if c_idx == 1 and "MR" in text_val:
                    run.font.color.rgb = COLOR_SECONDARY
                else:
                    run.font.color.rgb = COLOR_TEXT
            elif c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text_val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = COLOR_TEXT
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text_val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.italic = True
                run.font.color.rgb = COLOR_TEXT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text_val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.color.rgb = COLOR_TEXT
    
    # Space after table
    sp_p = doc.add_paragraph()
    sp_p.paragraph_format.space_before = Pt(0)
    sp_p.paragraph_format.space_after = Pt(10)

doc.save(docx_path)
print(f"Generated DOCX: {docx_path} successfully!")
