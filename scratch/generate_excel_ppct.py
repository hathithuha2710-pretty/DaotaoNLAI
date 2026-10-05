# -*- coding: utf-8 -*-
"""
Script tạo file Excel Phân phối chương trình Giáo dục AI Cấp Tiểu học (Lớp 1 - 5)
Chuyên nghiệp, đa sheet, chuẩn phong cách QĐ 2422 & Bảng ma trận 46 kỹ năng.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import json
import re

from ppct_data import GRADES_CONFIG
from lessons_grade1_v3 import GRADE1_12_CORE, GRADE1_36_EXPANDED
from lessons_grade2_v3 import GRADE2_12_CORE, GRADE2_36_EXPANDED
from lessons_grade3_v3 import GRADE3_12_CORE, GRADE3_36_EXPANDED
from lessons_grade4_v3 import GRADE4_12_CORE, GRADE4_36_EXPANDED
from lessons_grade5_v3 import GRADE5_12_CORE, GRADE5_36_EXPANDED

# Load external data
with open('scratch/yccd_all_grades.json', 'r', encoding='utf-8') as f:
    yccd_all = json.load(f)

with open('scratch/skills_46.json', 'r', encoding='utf-8') as f:
    skills_all = json.load(f)

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Color palettes
NAVY_HEADER = "1F4E78"
WHITE_TEXT = "FFFFFF"
YELLOW_CORE = "FFF2CC"
BLUE_ASSESS = "D9E1F2"
GRAY_LIGHT = "F2F2F2"
BORDER_GRAY = "BFBFBF"

font_title = Font(name="Arial", size=14, bold=True, color="1F4E78")
font_subtitle = Font(name="Arial", size=10, italic=True, color="595959")
font_header = Font(name="Arial", size=10, bold=True, color=WHITE_TEXT)
font_bold = Font(name="Arial", size=9, bold=True)
font_regular = Font(name="Arial", size=9)
font_italic = Font(name="Arial", size=8, italic=True, color="595959")

fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
fill_core = PatternFill(start_color=YELLOW_CORE, end_color=YELLOW_CORE, fill_type="solid")
fill_assess = PatternFill(start_color=BLUE_ASSESS, end_color=BLUE_ASSESS, fill_type="solid")
fill_gray = PatternFill(start_color=GRAY_LIGHT, end_color=GRAY_LIGHT, fill_type="solid")

thin_side = Side(border_style="thin", color=BORDER_GRAY)
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

# -------------------------------------------------------------
# 1. SHEET TỔNG QUAN KIẾN TRÚC
# -------------------------------------------------------------
ws_arch = wb.create_sheet(title="Tong_Quan_Kien_Truc")
ws_arch.views.sheetView[0].showGridLines = True

ws_arch.cell(row=1, column=1, value="BẢNG TỔNG QUAN KIẾN TRÚC CHƯƠNG TRÌNH GIÁO DỤC AI CẤP TIỂU HỌC (LỚP 1 – 5)").font = font_title
ws_arch.cell(row=2, column=1, value="Căn cứ Quyết định số 2422/QĐ-BGDĐT, Công văn số 5588/BGDĐT-GDPT và Bảng Ma trận 46 Kỹ năng v2.0").font = font_subtitle

arch_headers = ["Khối lớp", "Thời lượng / tiết", "Khung thời lượng", "Số lượng YCCĐ", "Cơ cấu tiết theo mạch", "Đặc trưng nhận thức & Sư phạm", "Phương thức đánh giá (Mục VI)"]
for col_idx, h in enumerate(arch_headers, 1):
    cell = ws_arch.cell(row=4, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_header
    cell.border = border_cell

arch_data = [
    ["LỚP 1", "30–35 phút", "12 Core / 36 Toàn diện", "21 YCCĐ (17 Cốt lõi + 4 Mở rộng)", "Mạch A (8t) + B (6t) + C (6t) + D (11t) + Đánh giá (5t)", "Trực quan – Cảm xúc Người vs Máy – Thẻ tranh & Đóng vai Unplugged", "Đánh giá quá trình bằng quan sát, nhận xét; hồ sơ tranh vẽ/mô hình robot giấy"],
    ["LỚP 2", "30–35 phút", "12 Core / 36 Toàn diện", "29 YCCĐ (17 Cốt lõi + 12 Mở rộng)", "Mạch A (9t) + B (7t) + C (6t) + D (9t) + Đánh giá (5t)", "Trực quan – AI Gia đình – Nhận diện dữ liệu & Thiên lệch ban đầu", "Bảng kiểm hành vi (Checklist), quan sát thao tác phân loại thẻ dữ liệu"],
    ["LỚP 3", "35 phút", "12 Core / 36 Toàn diện", "32 YCCĐ (21 Cốt lõi + 11 Mở rộng)", "Mạch A (11t) + B (5t) + C (6t) + D (9t) + Đánh giá (5t)", "Tư duy logic – AI Trường học – Khám phá Dữ liệu & Học máy phòng máy", "Đánh giá sản phẩm dự án Teachable Machine, nhật ký kiểm chứng thông tin"],
    ["LỚP 4", "35 phút", "12 Core / 36 Toàn diện", "17 YCCĐ (10 Cốt lõi + 7 Mở rộng)", "Mạch A (7t) + B (5t) + C (6t) + D (13t) + Đánh giá (5t)", "Khám phá Khoa học – Ứng dụng AI Đời sống & Sáng tạo số – 4 bước học máy", "Hồ sơ học tập số (Digital Portfolio), sản phẩm AutoDraw, rubric đánh giá nhóm"],
    ["LỚP 5", "35 phút", "12 Core / 36 Toàn diện", "24 YCCĐ (15 Cốt lõi + 9 Mở rộng)", "Mạch A (7t) + B (5t) + C (6t) + D (13t) + Đánh giá (5t)", "Thiết kế Hệ thống – Đạo đức & Trách nhiệm giải trình – Dự án AI vì cộng đồng", "Báo cáo dự án học máy hoàn chỉnh (Model Card), thuyết trình giải trình trách nhiệm"]
]

for row_idx, r_vals in enumerate(arch_data, 5):
    for col_idx, val in enumerate(r_vals, 1):
        cell = ws_arch.cell(row=row_idx, column=col_idx, value=val)
        cell.font = font_bold if col_idx in [1, 3] else font_regular
        cell.alignment = align_center if col_idx in [1, 2, 3] else align_left
        cell.border = border_cell
        if row_idx % 2 == 1:
            cell.fill = fill_gray

# Set column widths
ws_arch.column_dimensions["A"].width = 12
ws_arch.column_dimensions["B"].width = 16
ws_arch.column_dimensions["C"].width = 22
ws_arch.column_dimensions["D"].width = 26
ws_arch.column_dimensions["E"].width = 30
ws_arch.column_dimensions["F"].width = 35
ws_arch.column_dimensions["G"].width = 40

# -------------------------------------------------------------
# HÀM TẠO SHEET PPCT CHO TỪNG LỚP
# -------------------------------------------------------------
def build_grade_sheet(ws, title, subtitle, headers, data, is_36=False):
    ws.views.sheetView[0].showGridLines = True
    ws.cell(row=1, column=1, value=title).font = font_title
    ws.cell(row=2, column=1, value=subtitle).font = font_subtitle
    
    for c_idx, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=c_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_header
        c.border = border_cell
        
    for r_idx, row in enumerate(data, 5):
        c1 = ws.cell(row=r_idx, column=1, value=row["tiet"])
        c2 = ws.cell(row=r_idx, column=2, value=row["chude"])
        c3 = ws.cell(row=r_idx, column=3, value=row["ten_bai"])
        c4 = ws.cell(row=r_idx, column=4, value=row["noidung"])
        c5 = ws.cell(row=r_idx, column=5, value=row["phanloai"])
        c6 = ws.cell(row=r_idx, column=6, value=row["hoclieu"])
        c7 = ws.cell(row=r_idx, column=7, value=row["ma_yccd"])
        
        c1.alignment = align_center
        c2.alignment = align_left
        c3.alignment = align_left
        c4.alignment = align_left
        c5.alignment = align_center
        c6.alignment = align_left
        c7.alignment = align_center
        
        c1.font = font_bold
        c2.font = font_regular
        c3.font = font_bold
        c4.font = font_regular
        c5.font = font_bold
        c6.font = font_regular
        c7.font = font_bold
        
        for c in [c1, c2, c3, c4, c5, c6, c7]:
            c.border = border_cell
            if row["phanloai"] == "Cốt lõi":
                c.fill = fill_core
            elif row["phanloai"] == "Đánh giá quá trình":
                c.fill = fill_assess

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 18

headers_ppct = ["Tiết", "Mạch / Chủ đề", "Tên bài học", "Nội dung trọng tâm & Kỹ năng tích hợp", "Phân loại", "Thiết bị & Học liệu", "Mã YCCĐ"]

# -------------------------------------------------------------
# 2. SHEETS CHO 5 KHỐI LỚP (12 CORE VÀ 36 TOÀN DIỆN)
# -------------------------------------------------------------
grades_raw = {
    1: {"12": GRADE1_12_CORE, "36": GRADE1_36_EXPANDED},
    2: {"12": GRADE2_12_CORE, "36": GRADE2_36_EXPANDED},
    3: {"12": GRADE3_12_CORE, "36": GRADE3_36_EXPANDED},
    4: {"12": GRADE4_12_CORE, "36": GRADE4_36_EXPANDED},
    5: {"12": GRADE5_12_CORE, "36": GRADE5_36_EXPANDED}
}

for g in range(1, 6):
    cfg = GRADES_CONFIG[g]
    
    # 12 CORE SHEET
    ws_12 = wb.create_sheet(title=f"Lop_{g}_12_Tiet")
    t12 = f"KẾ HOẠCH BÀI DẠY CHUYÊN ĐỀ GIÁO DỤC AI BẮT BUỘC - LỚP {g} (12 TIẾT CỐT LÕI)"
    sub12 = f"Bao phủ 100% {cfg['core_count']} YCCĐ Cốt lõi của Lớp {g} theo QĐ 2422 | Thời lượng: {cfg['duration']} | Chuẩn CV 5588"
    build_grade_sheet(ws_12, t12, sub12, headers_ppct, grades_raw[g]["12"], is_36=False)
    
    # 36 EXPANDED SHEET
    ws_36 = wb.create_sheet(title=f"Lop_{g}_36_Tiet")
    t36 = f"KẾ HOẠCH GIÁO DỤC TOÀN DIỆN MÔN TRÍ TUỆ NHÂN TẠO (AI) - LỚP {g} (36 TIẾT/NĂM HỌC)"
    sub36 = f"Đúng 12 tiết Cốt lõi (tô vàng) + 24 tiết Mở rộng/Đánh giá (nền trắng/xanh) | Bao phủ 100% YCCĐ Cốt lõi & Mở rộng | Đánh giá quá trình Mục VI QĐ 2422"
    build_grade_sheet(ws_36, t36, sub36, headers_ppct, grades_raw[g]["36"], is_36=True)

# -------------------------------------------------------------
# 3. SHEET BẢNG MA TRẬN 46 KỸ NĂNG TIỂU HỌC
# -------------------------------------------------------------
ws_skills = wb.create_sheet(title="Ma_Tran_46_Ky_Nang_Tieu_Hoc")
ws_skills.views.sheetView[0].showGridLines = True
ws_skills.cell(row=1, column=1, value="BẢNG MA TRẬN PHÂN BỔ 46 KỸ NĂNG GIÁO DỤC AI CẤP TIỂU HỌC (LỚP 1 – 5) — BẢN v2.0").font = font_title
ws_skills.cell(row=2, column=1, value="Đối chiếu phân tầng nhận thức từ Mức 1 (Quan sát) -> Mức 2 (Thực hành có hướng dẫn) -> Mức 3 (Thực hành độc lập)").font = font_subtitle

skill_headers = ["STT", "Tên kỹ năng", "Mạch NL", "Lớp 1 (30–35p)", "Lớp 2 (30–35p)", "Lớp 3 (35p)", "Lớp 4 (35p)", "Lớp 5 (35p)"]
for c_idx, h in enumerate(skill_headers, 1):
    c = ws_skills.cell(row=4, column=c_idx, value=h)
    c.font = font_header
    c.fill = fill_header
    c.alignment = align_header
    c.border = border_cell

for r_idx, sk in enumerate(skills_all, 5):
    c1 = ws_skills.cell(row=r_idx, column=1, value=sk["stt"])
    c2 = ws_skills.cell(row=r_idx, column=2, value=sk["name"])
    c3 = ws_skills.cell(row=r_idx, column=3, value=sk["mach"])
    c4 = ws_skills.cell(row=r_idx, column=4, value=sk["l1"].replace('<br>', '\n'))
    c5 = ws_skills.cell(row=r_idx, column=5, value=sk["l2"].replace('<br>', '\n'))
    c6 = ws_skills.cell(row=r_idx, column=6, value=sk["l3"].replace('<br>', '\n'))
    c7 = ws_skills.cell(row=r_idx, column=7, value=sk["l4"].replace('<br>', '\n'))
    c8 = ws_skills.cell(row=r_idx, column=8, value=sk["l5"].replace('<br>', '\n'))
    
    c1.alignment = align_center
    c2.alignment = align_left
    c3.alignment = align_center
    for c in [c4, c5, c6, c7, c8]:
        c.alignment = align_left
        
    c1.font = font_bold
    c2.font = font_bold
    c3.font = font_regular
    for c in [c4, c5, c6, c7, c8]:
        c.font = font_regular
        
    for c in [c1, c2, c3, c4, c5, c6, c7, c8]:
        c.border = border_cell
        if r_idx % 2 == 1:
            c.fill = fill_gray

ws_skills.column_dimensions["A"].width = 6
ws_skills.column_dimensions["B"].width = 28
ws_skills.column_dimensions["C"].width = 10
ws_skills.column_dimensions["D"].width = 35
ws_skills.column_dimensions["E"].width = 35
ws_skills.column_dimensions["F"].width = 35
ws_skills.column_dimensions["G"].width = 35
ws_skills.column_dimensions["H"].width = 35

# -------------------------------------------------------------
# 4. SHEET TRA CỨU 123 YCCĐ CHUẨN BỘ GD&ĐT
# -------------------------------------------------------------
ws_lookup = wb.create_sheet(title="Tra_Cuu_123_YCCD_Chuan")
ws_lookup.views.sheetView[0].showGridLines = True
ws_lookup.cell(row=1, column=1, value="BẢNG TRA CỨU 123 YÊU CẦU CẦN ĐẠT (YCCĐ) AI CẤP TIỂU HỌC THEO QĐ 2422/QĐ-BGDĐT").font = font_title
ws_lookup.cell(row=2, column=1, value="Bao phủ 100% 80 YCCĐ Cốt lõi và 43 YCCĐ Mở rộng từ Lớp 1 đến Lớp 5 | Single Source of Truth").font = font_subtitle

lookup_headers = ["STT", "Khối lớp", "Mã YCCĐ", "Phân loại", "Mạch & Chủ đề", "Nội dung Yêu cầu cần đạt nguyên văn", "Đối sánh Benchmark Quốc tế"]
for c_idx, h in enumerate(lookup_headers, 1):
    c = ws_lookup.cell(row=4, column=c_idx, value=h)
    c.font = font_header
    c.fill = fill_header
    c.alignment = align_header
    c.border = border_cell

curr_row = 5
item_stt = 1
for g in range(1, 6):
    g_data = yccd_all[str(g)]
    # Core first
    for code, info in g_data["core"].items():
        ws_lookup.cell(row=curr_row, column=1, value=item_stt).alignment = align_center
        ws_lookup.cell(row=curr_row, column=2, value=f"Lớp {g}").alignment = align_center
        c_code = ws_lookup.cell(row=curr_row, column=3, value=code)
        c_code.alignment = align_center
        c_code.font = font_bold
        c_type = ws_lookup.cell(row=curr_row, column=4, value="Cốt lõi")
        c_type.alignment = align_center
        c_type.font = font_bold
        ws_lookup.cell(row=curr_row, column=5, value=info["mach"]).alignment = align_left
        ws_lookup.cell(row=curr_row, column=6, value=info["yccd"]).alignment = align_left
        ws_lookup.cell(row=curr_row, column=7, value=info["intl"]).alignment = align_left
        
        for c in range(1, 8):
            cell = ws_lookup.cell(row=curr_row, column=c)
            cell.border = border_cell
            cell.fill = fill_core
            cell.font = font_regular if c not in [1, 2, 3, 4] else font_bold
        curr_row += 1
        item_stt += 1
        
    # MR next
    for code, info in g_data["mr"].items():
        ws_lookup.cell(row=curr_row, column=1, value=item_stt).alignment = align_center
        ws_lookup.cell(row=curr_row, column=2, value=f"Lớp {g}").alignment = align_center
        c_code = ws_lookup.cell(row=curr_row, column=3, value=code)
        c_code.alignment = align_center
        c_code.font = font_bold
        c_type = ws_lookup.cell(row=curr_row, column=4, value="Mở rộng")
        c_type.alignment = align_center
        c_type.font = font_bold
        ws_lookup.cell(row=curr_row, column=5, value=info["mach"]).alignment = align_left
        ws_lookup.cell(row=curr_row, column=6, value=info["yccd"]).alignment = align_left
        ws_lookup.cell(row=curr_row, column=7, value=info["intl"]).alignment = align_left
        
        for c in range(1, 8):
            cell = ws_lookup.cell(row=curr_row, column=c)
            cell.border = border_cell
            cell.font = font_regular if c not in [1, 2, 3, 4] else font_bold
        curr_row += 1
        item_stt += 1

ws_lookup.column_dimensions["A"].width = 6
ws_lookup.column_dimensions["B"].width = 10
ws_lookup.column_dimensions["C"].width = 14
ws_lookup.column_dimensions["D"].width = 12
ws_lookup.column_dimensions["E"].width = 30
ws_lookup.column_dimensions["F"].width = 50
ws_lookup.column_dimensions["G"].width = 40

# Save Workbook
excel_out = '03_San_Pham_Chinh_Thuc/Phan_Phoi_Chuong_Trinh_AI_Tieu_Hoc_Lop1_5_QD2422.xlsx'
wb.save(excel_out)
print(f"Generated Excel successfully at: {excel_out}")
