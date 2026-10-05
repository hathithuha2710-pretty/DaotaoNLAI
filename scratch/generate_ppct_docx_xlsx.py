import os
import sys
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

md_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP3_Phan_Phoi_Chuong_Trinh_AI_Tieu_Hoc_Lop1_5_Moi.md"
docx_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP3_Phan_Phoi_Chuong_Trinh_AI_Tieu_Hoc_Lop1_5_Moi.docx"
xlsx_path = "03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP3_Phan_Phoi_Chuong_Trinh_AI_Tieu_Hoc_Lop1_5_Moi.xlsx"

with open(md_path, "r", encoding="utf-8") as f:
    text = f.read()

# -------------------------------------------------------------
# 1. GENERATE DOCX (A4 Portrait)
# -------------------------------------------------------------
doc = Document()
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79)     # 2.0 cm
    section.bottom_margin = Inches(0.79)  # 2.0 cm
    section.left_margin = Inches(0.98)    # 2.5 cm
    section.right_margin = Inches(0.79)   # 2.0 cm

COLOR_PRIMARY = RGBColor(15, 76, 129)     # Classic Navy Blue
COLOR_SECONDARY = RGBColor(40, 116, 166)
COLOR_TEXT = RGBColor(33, 33, 33)
HEX_PRIMARY_BG = "0F4C81"
HEX_ZEBRA_BG = "F7F9FB"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY_BG}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

# Main Title
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.space_before = Pt(0)
tp.paragraph_format.space_after = Pt(2)
rt = tp.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — CHƯƠNG TRÌNH GIÁO DỤC TRÍ TUỆ NHÂN TẠO (AI)\nKẾ HOẠCH BÀI DẠY VÀ PHÂN PHỐI CHƯƠNG TRÌNH CẤP TIỂU HỌC")
rt.font.name = "Times New Roman"
rt.font.size = Pt(13)
rt.font.bold = True
rt.font.color.rgb = COLOR_PRIMARY

sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sp.paragraph_format.space_before = Pt(0)
sp.paragraph_format.space_after = Pt(8)
rs = sp.add_run("BẢN CHUẨN HOÁ TIỂU HỌC 6–11 TUỔI THEO QUYẾT ĐỊNH 2422/QĐ-BGDĐT VÀ THAM CHIẾU QUỐC TẾ (SP3)")
rs.font.name = "Times New Roman"
rs.font.size = Pt(10.5)
rs.font.bold = True
rs.font.color.rgb = COLOR_SECONDARY

# Callout / Preamble
cp = doc.add_paragraph()
cp.paragraph_format.space_before = Pt(0)
cp.paragraph_format.space_after = Pt(10)
cp.paragraph_format.line_spacing = 1.15
rc = cp.add_run(
    "• Căn cứ: Quyết định số 2422/QĐ-BGDĐT ngày 18/8/2026 và Công văn số 5588/BGDĐT-GDPT ngày 19/8/2026 của Bộ GD&ĐT.\n"
    "• Cấu trúc: 36 tiết/lớp/năm học (12 tiết Cốt lõi + 24 tiết Mở rộng/Đánh giá) phủ kín 100% 123 YCCĐ của Bộ GD&ĐT.\n"
    "• Rào chắn sư phạm: Khoanh vùng nghiêm ngặt lứa tuổi 6–11; loại bỏ hoàn toàn các khái niệm và công cụ THCS."
)
rc.font.name = "Times New Roman"
rc.font.size = Pt(9.5)
rc.font.italic = True
rc.font.color.rgb = COLOR_TEXT

# Column widths: 0.35, 1.4, 2.5, 0.65, 0.95, 0.65 = 6.5 in
col_widths = [Inches(0.35), Inches(1.4), Inches(2.5), Inches(0.65), Inches(0.95), Inches(0.65)]
headers = ["Tiết", "Tên bài học / Chuyên đề", "Nội dung trọng tâm & Kỹ năng tích hợp", "Phân loại", "Thiết bị & Học liệu", "Mã YCCĐ"]

# Split text by sections
# Look for sections: Grade 1, 2, 3, 4, 5
grade_blocks = re.split(r'## PHẦN [I|V|X]+\. KẾ HOẠCH BÀI DẠY VÀ PHÂN PHỐI CHƯƠNG TRÌNH AI - LỚP (\d)', text)

excel_data_by_grade = {}

for b_idx in range(1, len(grade_blocks), 2):
    g_num = int(grade_blocks[b_idx])
    b_content = grade_blocks[b_idx + 1]

    # Section Title
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True
    r_h1 = h1.add_run(f"KẾ HOẠCH BÀI DẠY VÀ PHÂN PHỐI CHƯƠNG TRÌNH AI - LỚP {g_num}")
    r_h1.font.name = "Times New Roman"
    r_h1.font.size = Pt(12)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    # Subsections: 12 Core vs 36 Full
    sub_parts = re.split(r'### (\d\..+)', b_content)

    grade_core_rows = []
    grade_full_rows = []

    for s_idx in range(1, len(sub_parts), 2):
        sub_title = sub_parts[s_idx].strip()
        sub_body = sub_parts[s_idx + 1]

        h2 = doc.add_paragraph()
        h2.paragraph_format.space_before = Pt(10)
        h2.paragraph_format.space_after = Pt(4)
        h2.paragraph_format.keep_with_next = True
        r_h2 = h2.add_run(sub_title)
        r_h2.font.name = "Times New Roman"
        r_h2.font.size = Pt(10.5)
        r_h2.font.bold = True
        r_h2.font.color.rgb = COLOR_SECONDARY

        # Parse table rows in sub_body
        lines = sub_body.strip().split('\n')
        table_rows = []
        for line in lines:
            if re.match(r'\|\s*\*\*\d+\*\*', line):
                parts = [p.strip() for p in line.split('|')[1:-1]]
                if len(parts) >= 6:
                    tiet = parts[0].replace('**', '')
                    ten = parts[1].replace('**', '')
                    noidung = parts[2]
                    phanloai = parts[3].replace('`', '')
                    hoclieu = parts[4]
                    mayccd = parts[5].replace('**', '')
                    table_rows.append((tiet, ten, noidung, phanloai, hoclieu, mayccd))

        if not table_rows:
            continue

        if "12 tiết" in sub_title or "Bắt buộc" in sub_title:
            grade_core_rows = table_rows
        else:
            grade_full_rows = table_rows

        # Render Table in DOCX
        tbl = doc.add_table(rows=len(table_rows) + 1, cols=6)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)

        # Header Row
        hr = tbl.rows[0]
        make_row_header(hr)
        make_row_cant_split(hr)
        for c_idx, cell in enumerate(hr.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, HEX_PRIMARY_BG)
            set_cell_margins(cell, top=70, bottom=70, left=60, right=60)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 3, 5] else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(headers[c_idx])
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, r_data in enumerate(table_rows):
            row = tbl.rows[r_idx + 1]
            make_row_cant_split(row)
            is_zebra = (r_idx % 2 == 1)
            bg = HEX_ZEBRA_BG if is_zebra else "FFFFFF"

            for c_idx, cell in enumerate(row.cells):
                cell.width = col_widths[c_idx]
                if bg != "FFFFFF":
                    set_cell_background(cell, bg)
                set_cell_margins(cell, top=50, bottom=50, left=50, right=50)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1

                val = r_data[c_idx]
                if c_idx in [0, 3, 5]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(val)
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.5)
                    run.font.bold = (c_idx == 0 or c_idx == 5)
                    run.font.color.rgb = COLOR_TEXT
                elif c_idx == 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(val)
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.5)
                    run.font.bold = True
                    run.font.color.rgb = COLOR_TEXT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(val)
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = COLOR_TEXT

        sp_tbl = doc.add_paragraph()
        sp_tbl.paragraph_format.space_before = Pt(0)
        sp_tbl.paragraph_format.space_after = Pt(6)

    excel_data_by_grade[g_num] = {
        "core": grade_core_rows,
        "full": grade_full_rows
    }

doc.save(docx_path)
print(f"Generated DOCX: {docx_path} successfully!")

# -------------------------------------------------------------
# 2. GENERATE XLSX (Excel)
# -------------------------------------------------------------
wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

font_h = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
fill_h = PatternFill(start_color="0F4C81", end_color="0F4C81", fill_type="solid")
font_sub = Font(name="Times New Roman", size=11, bold=True, color="0F4C81")
fill_sub = PatternFill(start_color="EAF2F8", end_color="EAF2F8", fill_type="solid")
font_c = Font(name="Times New Roman", size=9)
font_cb = Font(name="Times New Roman", size=9, bold=True)
fill_z = PatternFill(start_color="F7F9FB", end_color="F7F9FB", fill_type="solid")

thin_b = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

align_c = Alignment(horizontal="center", vertical="top", wrap_text=True)
align_l = Alignment(horizontal="left", vertical="top", wrap_text=True)

excel_headers = ["Tiết", "Tên bài học / Chuyên đề", "Nội dung trọng tâm & Kỹ năng tích hợp", "Phân loại", "Thiết bị & Học liệu", "Mã YCCĐ"]

# Create a sheet for each grade
for g_num in range(1, 6):
    ws = wb.create_sheet(title=f"PPCT Lớp {g_num}")

    # Title
    ws.merge_cells("A1:F1")
    ws["A1"] = f"KẾ HOẠCH BÀI DẠY VÀ PHÂN PHỐI CHƯƠNG TRÌNH AI - LỚP {g_num} (36 TIẾT)"
    ws["A1"].font = Font(name="Times New Roman", size=13, bold=True, color="0F4C81")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:F2")
    ws["A2"] = "Chuẩn hoá Quyết định 2422/QĐ-BGDĐT và Công văn 5588/BGDĐT-GDPT"
    ws["A2"].font = Font(name="Times New Roman", size=10, italic=True, color="555555")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    r_cur = 4
    for col_idx, h_text in enumerate(excel_headers, 1):
        cell = ws.cell(row=r_cur, column=col_idx, value=h_text)
        cell.font = font_h
        cell.fill = fill_h
        cell.alignment = align_c
        cell.border = thin_b
    ws.row_dimensions[r_cur].height = 24
    r_cur += 1

    g_rows = excel_data_by_grade.get(g_num, {}).get("full", [])
    for r_idx, r_data in enumerate(g_rows):
        is_zebra = (r_idx % 2 == 1)
        r_fill = fill_z if is_zebra else None

        for col_idx, val in enumerate(r_data, 1):
            cell = ws.cell(row=r_cur, column=col_idx, value=int(val) if col_idx == 1 and val.isdigit() else val)
            cell.font = font_cb if col_idx in [1, 2, 4, 6] else font_c
            if r_fill:
                cell.fill = r_fill
            cell.alignment = align_c if col_idx in [1, 4, 6] else align_l
            cell.border = thin_b

        ws.row_dimensions[r_cur].height = 55
        r_cur += 1

    # Widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 45
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 16

# Also create an overview sheet: Sheet 1 (Khung Cốt Lõi 12 tiết cả 5 lớp)
ws_core = wb.create_sheet(title="Khung 12 tiết Cốt lõi (1-5)", index=0)
ws_core.merge_cells("A1:G1")
ws_core["A1"] = "TỔNG HỢP CHUYÊN ĐỀ 12 TIẾT CỐT LÕI BẮT BUỘC (LỚP 1 ĐẾN LỚP 5)"
ws_core["A1"].font = Font(name="Times New Roman", size=13, bold=True, color="0F4C81")
ws_core["A1"].alignment = Alignment(horizontal="center", vertical="center")

core_headers = ["Lớp", "Tiết", "Tên bài học cốt lõi", "Nội dung trọng tâm & Kỹ năng tích hợp", "Phân loại", "Thiết bị & Học liệu", "Mã YCCĐ"]
r_cur = 3
for col_idx, h_text in enumerate(core_headers, 1):
    cell = ws_core.cell(row=r_cur, column=col_idx, value=h_text)
    cell.font = font_h
    cell.fill = fill_h
    cell.alignment = align_c
    cell.border = thin_b
ws_core.row_dimensions[r_cur].height = 24
r_cur += 1

for g_num in range(1, 6):
    c_rows = excel_data_by_grade.get(g_num, {}).get("core", [])
    for r_idx, r_data in enumerate(c_rows):
        is_zebra = (r_idx % 2 == 1)
        r_fill = fill_z if is_zebra else None

        row_vals = [f"Lớp {g_num}"] + list(r_data)
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws_core.cell(row=r_cur, column=col_idx, value=int(val) if col_idx == 2 and val.isdigit() else val)
            cell.font = font_cb if col_idx in [1, 2, 3, 5, 7] else font_c
            if r_fill:
                cell.fill = r_fill
            cell.alignment = align_c if col_idx in [1, 2, 5, 7] else align_l
            cell.border = thin_b

        ws_core.row_dimensions[r_cur].height = 50
        r_cur += 1

ws_core.column_dimensions['A'].width = 10
ws_core.column_dimensions['B'].width = 8
ws_core.column_dimensions['C'].width = 30
ws_core.column_dimensions['D'].width = 45
ws_core.column_dimensions['E'].width = 12
ws_core.column_dimensions['F'].width = 25
ws_core.column_dimensions['G'].width = 16

wb.save(xlsx_path)
print(f"Generated XLSX: {xlsx_path} successfully!")
