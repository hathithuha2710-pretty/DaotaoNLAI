import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('03_San_Pham_Chinh_Thuc/YCCD_Lop_1_5_Moi/SP3_Phan_Phoi_Chuong_Trinh_AI_Tieu_Hoc_Lop1_5_Moi.docx')
print(f"Total tables in docx: {len(doc.tables)}")

for idx, tbl in enumerate(doc.tables):
    rows = tbl.rows
    # check first cell of each row after header
    tiets = []
    for r in rows[1:]:
        c0 = r.cells[0].text.strip()
        tiets.append(c0)
    print(f"Table {idx+1}: {len(rows)-1} data rows -> {tiets}")
